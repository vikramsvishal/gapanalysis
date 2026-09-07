#!/usr/bin/env python3
"""NW Hardware Catalog Matcher

GUI inputs (exactly two):
  1. IS NW Catalog file (.xlsx, .xlsm, or .csv)
  2. NW CMDB report (.xlsx, .xlsm, or .csv)

The supplied NW Catalog Match Excel template is embedded in this script. The user
is never asked to select it. The output is always an .xlsx workbook based on that
template.
"""
from __future__ import annotations

import base64
import csv
import io
import re
import traceback
from dataclasses import dataclass
from pathlib import Path
from tempfile import NamedTemporaryFile
from typing import Any, Iterable

from openpyxl import load_workbook

APP_TITLE = "NW Hardware Catalog Matcher"
TEMPLATE_B64 = "UEsDBBQABgAIAAAAIQBi7p1oXgEAAJAEAAATAAgCW0NvbnRlbnRfVHlwZXNdLnhtbCCiBAIooAACAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAACslMtOwzAQRfdI/EPkLUrcskAINe2CxxIqUT7AxJPGqmNbnmlp/56J+xBCoRVqN7ESz9x7MvHNaLJubbaCiMa7UgyLgcjAVV4bNy/Fx+wlvxcZknJaWe+gFBtAMRlfX41mmwCYcbfDUjRE4UFKrBpoFRY+gOOd2sdWEd/GuQyqWqg5yNvB4E5W3hE4yqnTEOPRE9RqaSl7XvPjLUkEiyJ73BZ2XqVQIVhTKWJSuXL6l0u+cyi4M9VgYwLeMIaQvQ7dzt8Gu743Hk00GrKpivSqWsaQayu/fFx8er8ojov0UPq6NhVoXy1bnkCBIYLS2ABQa4u0Fq0ybs99xD8Vo0zL8MIg3fsl4RMcxN8bZLqej5BkThgibSzgpceeRE85NyqCfqfIybg4wE/tYxx8bqbRB+QERfj/FPYR6brzwEIQycAhJH2H7eDI6Tt77NDlW4Pu8ZbpfzL+BgAA//8DAFBLAwQUAAYACAAAACEAtVUwI/QAAABMAgAACwAIAl9yZWxzLy5yZWxzIKIEAiigAAIAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAKySTU/DMAyG70j8h8j31d2QEEJLd0FIuyFUfoBJ3A+1jaMkG92/JxwQVBqDA0d/vX78ytvdPI3qyCH24jSsixIUOyO2d62Gl/pxdQcqJnKWRnGs4cQRdtX11faZR0p5KHa9jyqruKihS8nfI0bT8USxEM8uVxoJE6UchhY9mYFaxk1Z3mL4rgHVQlPtrYawtzeg6pPPm3/XlqbpDT+IOUzs0pkVyHNiZ9mufMhsIfX5GlVTaDlpsGKecjoieV9kbMDzRJu/E/18LU6cyFIiNBL4Ms9HxyWg9X9atDTxy515xDcJw6vI8MmCix+o3gEAAP//AwBQSwMEFAAGAAgAAAAhAAMMuYa7AwAAfQkAAA8AAAB4bC93b3JrYm9vay54bWysVm1vozgQ/n7S/QfEd4rNSyCodMWrrlK7qtJseidVWrngFKuAc8Y0qar97zsmJGmb0ynXvSixsT08fsbzzDjnXzZNrT1T0THehjo+Q7pG24KXrH0M9W/z3PB1rZOkLUnNWxrqL7TTv1z8/tv5mounB86fNABou1CvpFwFptkVFW1Id8ZXtIWVJRcNkTAUj2a3EpSUXUWpbGrTQmhiNoS1+hYhEKdg8OWSFTTlRd/QVm5BBK2JBPpdxVbdDq0pToFriHjqV0bBmxVAPLCayZcBVNeaIrh8bLkgDzW4vcGuthHwncAPI2is3U6wdLRVwwrBO76UZwBtbkkf+Y+RifG7I9gcn8FpSI4p6DNTMdyzEpNPsprssSYHMIx+GQ2DtAatBHB4n0Rz99ws/eJ8yWq62EpXI6vVV9KoSNW6VpNOZiWTtAx1D4Z8Td9NiH4V96yGVRtZlq+bF3s53witpEvS13IOQt7Bh7qFLBshZQnCiGpJRUskTXgrQYejX7+quQE7qTgoXJvRv3smKCQW6At8hZYUAXnoboistF7UoZ4E9986cP9+wZ4EaRasq0h9n/J1W3PIs/s3AiXH2fAfJEoK5bcJjm/JbZ8/HgJwFMFOhjdSaPB8mV5BKG7JMwQGwl+OeXsJJ+9/f40TK7adNDJSK3ENB1uuMbV830DIzfLUsxI7z36AF2ISFJz0shqDrTBD3YHIHi1dk81uBaOgZ+Vh/1c0fgAfoQ/Nbu2H8lSVtQWj6+4gCzXUNnesLfk61A2sxPzyfrgeFu9YKSsopFPHApPt3B+UPVbAGCMlIVU7FLNQf02tGKcumhoon3qGkyS5EWE/MqIM+V4c506aTgZG5htKQwEFakOvtYPob1VRxVCpVa9OF55FoPYQlyUeord7rSB1ASJX3WA4xciaKgu6kVedHHrQFwN62EGRh6aOgTIb4uNPLcN3bMtInNTKXC9Ls9hV8VEXQPB/lMFB5sHuZlEsKyLkXJDiCe6jGV3GpAMlbR0Cvm/Jxq4fIxsoOjnOQUxTZMTxxDHcNLddD6dJ5uYHssr95SeLkG8Ob1Mie0hQlZvDOFBtPs7uJ5fbiTFO75IumKXq3Me3/83wFryv6YnG+eJEw+Tr9fz6RNurbP79Lj/VOLqO0+h0+2g2i/6aZ3/utjD/8UDNDwFPsTNFdhYZtp04huPlnuHnyDVsx3MS14kzjLxDwOt18fy5eFuOuVNk8vZ/wliMVPwVeDD+idI6KscluDWG1BuIK/pDfu3RLn4CAAD//wMAUEsDBBQABgAIAAAAIQCBPpSX8wAAALoCAAAaAAgBeGwvX3JlbHMvd29ya2Jvb2sueG1sLnJlbHMgogQBKKAAAQAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAACsUk1LxDAQvQv+hzB3m3YVEdl0LyLsVesPCMm0KdsmITN+9N8bKrpdWNZLLwNvhnnvzcd29zUO4gMT9cErqIoSBHoTbO87BW/N880DCGLtrR6CRwUTEuzq66vtCw6acxO5PpLILJ4UOOb4KCUZh6OmIkT0udKGNGrOMHUyanPQHcpNWd7LtOSA+oRT7K2CtLe3IJopZuX/uUPb9gafgnkf0fMZCUk8DXkA0ejUISv4wUX2CPK8/GZNec5rwaP6DOUcq0seqjU9fIZ0IIfIRx9/KZJz5aKZu1Xv4XRC+8opv9vyLMv072bkycfV3wAAAP//AwBQSwMEFAAGAAgAAAAhAMSmhy+PAgAA1wYAABgAAAB4bC93b3Jrc2hlZXRzL3NoZWV0MS54bWyclWtvmzAUhr9P2n+w/D3hEkISBFRp06xd161ru313jEmsAma2c6mm/fcdw0ILZFJVKUg2efyei88rwrNDnqEdk4qLIsLO0MaIFVQkvFhH+MfjcjDFSGlSJCQTBYvwM1P4LP74IdwL+aQ2jGkECoWK8EbrMrAsRTcsJ2ooSlbAP6mQOdGwlWtLlZKRpDqUZ5Zr276VE17gWiGQb9EQacopWwi6zVmhaxHJMqIhf7XhpTqq5fQtcjmRT9tyQEVegsSKZ1w/V6IY5TS4XhdCklUGdR8cj1B0kPBz4Rkdw1Tve5FyTqVQItVDULbqnPvlz6yZRWij1K//TTKOZ0m24+YCX6Tc96XkjBst90Vs9E4xvxEz7ZLBlicR/j11bM9ferPB+fJyPvDmF5PB3HfPB/bldOIuFwvfHtt/cBwmHG7YVIUkSyM8d4JHB1txWM3PT8726tUaabJ6YBmjmkEMByMznishngx4Da9sUFQVYBQJ1XzHLliWRfjGTPivKgYsIYDVRHi9PkZbVgN9J1HCUrLN9L3YXzG+3mgI6w3HUKiZlCB5XjBFYUQh9HA0bhJfEE3iUIo9guuGPFVJjHmcwAXnGZGR/3+ROKTm2Nycq04Dr6CwXWyH1g6ypf+I8z7htImLPuG2iUWfGLWJyz7htYllnxi3iU99wm8TV31i0iauawImtunHtE187hOzNnHTJ5xOU7+cQDpdvT2BdNr6tUbAEU2yTqev304gncbenUA6nf1eI97rQJ3W3p9AOr19OBGo01ywpRnHVkUv3bVg1o8+qoe/JGt2S+SaFwplLK0cMsFI1iayh7DWojS+mYCdVkJrkR93G/iAMLAAeAqjVAh93BjfNp+k+C8AAAD//wMAUEsDBBQABgAIAAAAIQD2YLRBuAcAABEiAAATAAAAeGwvdGhlbWUvdGhlbWUxLnhtbOxazY8btxW/B8j/QMxd1szoe2E50Kc39u564ZVd5EhJlIZeznBAUrsrFAEK59RLgQJp0UuB3nooigZogAa55I8xYCNN/4g8ckaa4YqKvf5AkmJ3LzPU7z3+5r3HxzePc/eTq5ihCyIk5UnXC+74HiLJjM9psux6TybjSttDUuFkjhlPSNdbE+l9cu/jj+7iAxWRmCCQT+QB7nqRUulBtSpnMIzlHZ6SBH5bcBFjBbdiWZ0LfAl6Y1YNfb9ZjTFNPJTgGNQ+WizojKCJVund2ygfMbhNlNQDMybOtGpiSRjs/DzQCLmWAybQBWZdD+aZ88sJuVIeYlgq+KHr+ebPq967W8UHuRBTe2RLcmPzl8vlAvPz0MwpltPtpP4obNeDrX4DYGoXN2rr/60+A8CzGTxpxqWsM2g0/XaYY0ug7NKhu9MKaja+pL+2wznoNPth3dJvQJn++u4zjjujYcPCG1CGb+zge37Y79QsvAFl+OYOvj7qtcKRhTegiNHkfBfdbLXbzRy9hSw4O3TCO82m3xrm8AIF0bCNLj3FgidqX6zF+BkXYwBoIMOKJkitU7LAM4jiXqq4REMqU4bXHkpxwiUM+2EQQOjV/XD7byyODwguSWtewETuDGk+SM4ETVXXewBavRLk5TffvHj+9Yvn/3nxxRcvnv8LHdFlpDJVltwhTpZluR/+/sf//fV36L///tsPX/7JjZdl/Kt//v7Vt9/9lHpYaoUpXv75q1dff/XyL3/4/h9fOrT3BJ6W4RMaE4lOyCV6zGN4QGMKmz+ZiptJTCJMLQkcgW6H6pGKLODJGjMXrk9sEz4VkGVcwPurZxbXs0isFHXM/DCKLeAx56zPhdMAD/VcJQtPVsnSPblYlXGPMb5wzT3AieXg0SqF9EpdKgcRsWieMpwovCQJUUj/xs8JcTzdZ5Radj2mM8ElXyj0GUV9TJ0mmdCpFUiF0CGNwS9rF0FwtWWb46eoz5nrqYfkwkbCssDMQX5CmGXG+3ilcOxSOcExKxv8CKvIRfJsLWZl3Egq8PSSMI5GcyKlS+aRgOctOf0hhsTmdPsxW8c2Uih67tJ5hDkvI4f8fBDhOHVypklUxn4qzyFEMTrlygU/5vYK0ffgB5zsdfdTSix3vz4RPIEEV6ZUBIj+ZSUcvrxPuL0e12yBiSvL9ERsZdeeoM7o6K+WVmgfEcLwJZ4Tgp586mDQ56ll84L0gwiyyiFxBdYDbMeqvk+IhDJJ1zW7KfKISitkz8iS7+FzvL6WeNY4ibHYp/kEvG6F7lTAYnRQeMRm52XgCYXyD+LFaZRHEnSUgnu0T+tphK29S99Ld7yuheW/N1ljsC6f3XRdggy5sQwk9je2zQQza4IiYCaYoiNXugURy/2FiN5XjdjKKbewF23hBiiMrHonpsnrip8TLAS//Hlqnw9W9bgVv0u9sy+vHF6rcvbhfoW1zRCvklMC28lu4rotbW5LG+//vrTZt5ZvC5rbgua2oHG9gn2QgqaoYaC8KVo9pvET7+37LChjZ2rNyJE0rR8JrzXzMQyanpRpTG77gGkEl/p5YAILtxTYyCDB1W+ois4inEJ/KDBdzKXMVS8lSrmEtpEZNv1Uck23aT6t4mM+z9qdpr/kZyaUWBXjfgMaT9k4tKpUhm628kHNb0PdsF2aVuuGgJa9CYnSZDaJmoNEazP4GhK6c/Z+WHQcLNpa/cZVO6YAaluvwHs3grf1rteoZ4ygIwc1+lz7KXP1xrvaOe/V0/uMycoRAK3FXU93NNe9j6efLgu1N/C0RcI4JQsrm4TxlSnwZARvw3l0lvvuPxVwN/V1p3CpRU+bYrMaChqt9ofwtU4i13IDS8qZgiXoEtZ4CIvOQzOcdr0F9I3hMk4heKR+98JsCYcvMyWyFf82qSUVUg2xjDKLm6yT+SemigjEaNz19PNvw4ElJolk5DqwdH+p5EK94H5p5MDrtpfJYkFmquz30oi2dHYLKT5LFs5fjfjbg7UkX4G7z6L5JZqylXiMIcQarUB7d04lHB8EmavnFM7DtpmsiL9rO1Oe/a1DriIfY5ZGON9Sytk8g5sNZUvH3G1tULrLnxkMumvC6VLvsO+87b5+r9aWK/bHTrFpWmlFb5vubPrhdvkSq2IXtVhluft6zu1skh0EqnObePe9v0StmMyiphnv5mGdtPNRm9p7rAhKu09zj922m4TTEm+79YPc9ajVO8SmsDSBbw7Oy2fbfPoMkscQThFXLDvtZgncmdIyPRXGt1M+X+eXTGaJJvO5LkqzVP6YLBCdX3W90FU55ofHeTXAEkCbmhdW2FbQWe3Zgnqzy0WzBbsVzsrYa/WqLbyV2ByzboVNa9FFW11tTtR1rW5m1g7LntqkYWMpuNq1IrTJBYbSOTvMzXIv5JkrlVfacIVWgna93/qNXn0QNgYVv90YVeq1ul9pN3q1Sq/RqAWjRuAP++HnQE9FcdDIvnwYw2kQW+ffP5jxnW8g4s2B150Zj6vcfONQNd4330AE4f5vIMCRQCscBfWwFw4qg2HQrNTDYbPSbtV6lUHYHIY92LSb497nHrow4KA/HI7HjbDSHACu7vcalV6/Nqg026N+OA5G9aEP4Hz7uYK3GJ1zc1vApeF170cAAAD//wMAUEsDBBQABgAIAAAAIQBKog6ycwMAAD8KAAANAAAAeGwvc3R5bGVzLnhtbMRW227bOBB9X6D/QPBd0cWSKxmSijiOgALdYIFkgX2lJcohyotA0Yncxf57h7rY8jbtuukWfRI5HB2emTlDMn3XCY6eqG6Zkhn2rzyMqCxVxeQuw38+FE6MUWuIrAhXkmb4QFv8Ln/zW9qaA6f3j5QaBBCyzfCjMc3KddvykQrSXqmGSliplRbEwFTv3LbRlFSt/UlwN/C8pSsIk3hAWInyEhBB9Md945RKNMSwLePMHHosjES5er+TSpMtB6qdH5ISdf5SB6jT0ya99Yt9BCu1alVtrgDXVXXNSvol3cRNXFKekAD5dUh+5HrBWeydfiVS6Gr6xGz5cJ7WSpoWlWovTYYXQNSmYPVRqmdZ2CWo8OiVp+0n9EQ4WHzs5mmpuNLIQOkgc71FEkEHj+vGqBbdEa3Vs/WtiWD8MKwF1tCXfHQWDApgja4lM1DK0631+lUbJj81wD7OFgJlnB9zH9o0gyFPQaSGalnABI3jh0MDSZbQT0Oeer//8N5pcvCD6PIfWsVZZVnsbvrS6t02w0WRBBsv8izMdlxgsqIdrTK8DHv0GWFbxEvI/XuvUUYhRoZZJXpXb5Mkif1lHMdJuPDDsJfNxOBCd/clZj1ByP5W6QpOsUn7VuaDKU85rQ3Eq9nu0X6Namz0yhjo9DytGNkpSbhV7PTHOADYknJ+b0+6v+oz7K5Gci8KYd5D4uDMtFqfhpCxcTjgDROLP0cbsGewEVD+fljU1Uf8r/0dAL+XSR3/RqRp+MGeEbb7h9k1Zzsp6GDKUzJN0bMmzQPtelcbVld/nflsbxjOE/LNvW27/M9MIKwpC3AyfjcTCPQsvz6oe6r6xUHe7cWW6qK/Ek/hnaf9B0LvBQaSmun2TLVH/SF7uGf4zhLhcE2MGkLbPePQsi8oFjCr7tQD/Qli7D3bd8dxF2iFitZkz83DcTHDp/HvtGJ7Afkavf5gT8r0EBk+jT/YVvWX9jgCmX1o4RqBL9prluG/b9dvk81tETixt46dcEEjJ4nWGycKb9abTZF4gXfzz+y2/4G7vn+cQN39cNVyeBHoMdiR/P3JluHZZKDfH6ZAe849CZbedeR7TrHwfCdcktiJl4vIKSI/2CzD9W1URDPu0SvfBJ7r+8PrwpKPVoYJypmcajVVaG6FIsH0G0G4UyXc08sv/wwAAP//AwBQSwMEFAAGAAgAAAAhAKbaQVZZAQAAHwMAABQAAAB4bC9zaGFyZWRTdHJpbmdzLnhtbHySTU/DMAyG70j8hyh3lrEDQqjthDoQEwyQBuJYea3bRmqSkjhA/z0p25BYO47x+zh+/RHNv1TDPtA6aXTMzydTzlDnppC6ivnry+3ZJWeOQBfQGI0x79DxeXJ6EjlHLORqF/OaqL0SwuU1KnAT06IOSmmsAgpPWwnXWoTC1YikGjGbTi+EAqk5y43XFPNZKOu1fPeY/gaSyMkkouTaOST2bKUC27F77CJBSSR6cQuky8kjKByEjS5l5S1QaI0tCdWAaMC5QRAIK2MHVdZ+kx+RHmSJLO3yBtmaoBo4OdR7ivyg8gq0LyEnb9EeulqZAhu2XIzGR9vfZmivNsPf1mglNGxcTIHCriv2n50wiR+mBlt8gsWMuna4gB3EfinVd3HYwt1bdqOL7KnM1r5tjaVsEeZ8nOqnOYrsjYcDLjJTZk0PFiN/7cEd53Zl/6IiHHjyDQAA//8DAFBLAwQUAAYACAAAACEAh4nlv0QBAABrAgAAEQAIAWRvY1Byb3BzL2NvcmUueG1sIKIEASigAAEAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAjJJRT8MgFIXfTfwPDe8ttDNTSdslavbkEhPnNL4RuNvICiWAdvv30nar1fngI5xzP865IZ/tVRV9gnWy1gVKE4Ii0LwWUm8K9LKcxzcocp5pwapaQ4EO4NCsvLzIuaG8tvBkawPWS3BRIGlHuSnQ1ntDMXZ8C4q5JDh0ENe1VcyHo91gw/iObQBnhEyxAs8E8wy3wNgMRHRECj4gzYetOoDgGCpQoL3DaZLib68Hq9yfA50ycirpDyZ0OsYdswXvxcG9d3IwNk2TNJMuRsif4rfF43NXNZa63RUHVOaCU26B+dqWK7mzTEUr6basyvFIabdYMecXYeFrCeLu8Nt8bgjkrkiPBxGFaLQvclJeJ/cPyzkqM5JNY3Ibk3RJrukVodnte/v+j/k2an+hjin+T0xpRkbEE6DM8dn3KL8AAAD//wMAUEsDBBQABgAIAAAAIQBsBndilQEAACADAAAQAAgBZG9jUHJvcHMvYXBwLnhtbCCiBAEooAABAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAJyS32/TMBDH35H4HyK/r04HmlDleIIONCQmKrUbz8a5NNYc27q7RS1/PU6idinwxNv9+Orrj+9O3R46X/SA5GKoxHJRigKCjbUL+0o87r5cfRAFsQm18TFAJY5A4la/faM2GBMgO6AiWwSqRMucVlKSbaEztMjtkDtNxM5wTnEvY9M4C3fRvnQQWF6X5Y2EA0Ooob5KZ0MxOa56/l/TOtqBj552x5SBtfqYknfWcP6lfnAWI8WGi88HC17JeVNlui3YF3R81KWS81RtrfGwzsa6MZ5AydeCugczDG1jHJJWPa96sByxIPcrj+1aFD8NwYBTid6gM4Ez1iCbkjH2iRj1j4jP1AIwKZkFU3EM59p57N7r5SjIwaVwMJhAcuMScefYA31vNgb5H8TLOfHIMPFOONuBb3pzzjd+Ob/0h/c6dsmEo/766aFYR0wRx00oeWqoby4802PaxTvDcJruZVFtW4NQ54Wcp38uqPs8WPSDybo1YQ/1SfN3Y7iFp+ng9fJmUb4r85pnNSVfT1v/BgAA//8DAFBLAQItABQABgAIAAAAIQBi7p1oXgEAAJAEAAATAAAAAAAAAAAAAAAAAAAAAABbQ29udGVudF9UeXBlc10ueG1sUEsBAi0AFAAGAAgAAAAhALVVMCP0AAAATAIAAAsAAAAAAAAAAAAAAAAAlwMAAF9yZWxzLy5yZWxzUEsBAi0AFAAGAAgAAAAhAAMMuYa7AwAAfQkAAA8AAAAAAAAAAAAAAAAAvAYAAHhsL3dvcmtib29rLnhtbFBLAQItABQABgAIAAAAIQCBPpSX8wAAALoCAAAaAAAAAAAAAAAAAAAAAKQKAAB4bC9fcmVscy93b3JrYm9vay54bWwucmVsc1BLAQItABQABgAIAAAAIQDEpocvjwIAANcGAAAYAAAAAAAAAAAAAAAAANcMAAB4bC93b3Jrc2hlZXRzL3NoZWV0MS54bWxQSwECLQAUAAYACAAAACEA9mC0QbgHAAARIgAAEwAAAAAAAAAAAAAAAACcDwAAeGwvdGhlbWUvdGhlbWUxLnhtbFBLAQItABQABgAIAAAAIQBKog6ycwMAAD8KAAANAAAAAAAAAAAAAAAAAIUXAAB4bC9zdHlsZXMueG1sUEsBAi0AFAAGAAgAAAAhAKbaQVZZAQAAHwMAABQAAAAAAAAAAAAAAAAAIxsAAHhsL3NoYXJlZFN0cmluZ3MueG1sUEsBAi0AFAAGAAgAAAAhAIeJ5b9EAQAAawIAABEAAAAAAAAAAAAAAAAArhwAAGRvY1Byb3BzL2NvcmUueG1sUEsBAi0AFAAGAAgAAAAhAGwGd2KVAQAAIAMAABAAAAAAAAAAAAAAAAAAKR8AAGRvY1Byb3BzL2FwcC54bWxQSwUGAAAAAAoACgCAAgAA9CEAAAAA"


def text(value: Any) -> str:
    return "" if value is None else str(value).strip()


def header_key(value: Any) -> str:
    """Normalize headers so spaces, underscores and case do not matter."""
    return re.sub(r"[^a-z0-9]+", "", text(value).lower())


def compact(value: Any) -> str:
    """Normalize model values by removing separators and punctuation."""
    return re.sub(r"[^a-z0-9]+", "", text(value).lower())


def model_tokens(value: Any) -> list[str]:
    """Return compound and split alphanumeric model fragments."""
    raw = text(value).lower()
    compound = re.findall(r"[a-z0-9]+", raw)
    split = re.findall(r"[a-z]+|[0-9]+", raw)
    result: list[str] = []
    for token in compound + split:
        if len(token) >= 2 and token not in result:
            result.append(token)
    return result


@dataclass(frozen=True)
class TableData:
    headers: list[str]
    rows: list[list[Any]]
    source_sheet: str

    @property
    def columns(self) -> dict[str, int]:
        result: dict[str, int] = {}
        for index, name in enumerate(self.headers):
            key = header_key(name)
            if key and key not in result:
                result[key] = index
        return result


@dataclass(frozen=True)
class CatalogRecord:
    hardware_type: str
    hardware_model: str
    manufacturer: str
    source_row: int
    type_compact: str
    model_compact: str
    type_tokens: tuple[str, ...]
    model_tokens: tuple[str, ...]


@dataclass(frozen=True)
class MatchResult:
    record: CatalogRecord | None
    method: str
    score: int = 0
    ambiguous: bool = False


def _decode_csv(path: Path) -> str:
    raw = path.read_bytes()
    for encoding in ("utf-8-sig", "utf-8", "cp1252", "latin-1"):
        try:
            return raw.decode(encoding)
        except UnicodeDecodeError:
            continue
    raise ValueError(f"Unable to decode CSV file: {path.name}")


def _read_csv(path: Path, expected_headers: Iterable[str]) -> TableData:
    content = _decode_csv(path)
    sample = content[:65536]
    try:
        dialect = csv.Sniffer().sniff(sample, delimiters=",;\t|")
    except csv.Error:
        dialect = csv.excel
    rows = list(csv.reader(io.StringIO(content), dialect))
    if not rows:
        raise ValueError(f"{path.name} is empty.")

    expected = {header_key(h) for h in expected_headers}
    best_row, best_score = 0, -1
    for i, row in enumerate(rows[:20]):
        score = len({header_key(v) for v in row if text(v)} & expected)
        if score > best_score:
            best_row, best_score = i, score
    if best_score <= 0:
        raise ValueError(f"Expected columns were not found in {path.name}.")

    headers = [text(v) for v in rows[best_row]]
    width = len(headers)
    data_rows = [(row + [""] * width)[:width] for row in rows[best_row + 1:]]
    return TableData(headers, data_rows, "CSV")


def _read_excel(path: Path, expected_headers: Iterable[str]) -> TableData:
    wb = load_workbook(path, data_only=True, read_only=True)
    try:
        expected = {header_key(h) for h in expected_headers}
        best_ws = None
        best_header_row = 1
        best_score = -1
        for ws in wb.worksheets:
            for row_number in range(1, min(ws.max_row, 20) + 1):
                values = [cell.value for cell in ws[row_number]]
                score = len({header_key(v) for v in values if text(v)} & expected)
                if score > best_score:
                    best_ws, best_header_row, best_score = ws, row_number, score
        if best_ws is None or best_score <= 0:
            raise ValueError(f"Expected columns were not found in {path.name}.")

        headers = [text(cell.value) for cell in best_ws[best_header_row]]
        width = len(headers)
        data_rows: list[list[Any]] = []
        for values in best_ws.iter_rows(
            min_row=best_header_row + 1,
            max_col=width,
            values_only=True,
        ):
            data_rows.append(list(values))
        return TableData(headers, data_rows, best_ws.title)
    finally:
        wb.close()


def read_table(path: Path, expected_headers: Iterable[str]) -> TableData:
    suffix = path.suffix.lower()
    if suffix == ".csv":
        return _read_csv(path, expected_headers)
    if suffix in {".xlsx", ".xlsm"}:
        return _read_excel(path, expected_headers)
    raise ValueError(
        f"Unsupported file type for {path.name}. Use .xlsx, .xlsm, or .csv."
    )


def require_columns(table: TableData, required: Iterable[str], label: str) -> None:
    missing = [name for name in required if header_key(name) not in table.columns]
    if missing:
        raise ValueError(f"{label} is missing required column(s): {', '.join(missing)}")


def build_catalog(table: TableData) -> tuple[list[CatalogRecord], dict[str, list[CatalogRecord]]]:
    require_columns(table, ["hardware_type", "hardware_model", "manufacturer"], "IS NW Catalog")
    cols = table.columns
    records: list[CatalogRecord] = []
    exact_index: dict[str, list[CatalogRecord]] = {}

    for excel_row, row in enumerate(table.rows, start=2):
        hw_type = text(row[cols[header_key("hardware_type")]])
        hw_model = text(row[cols[header_key("hardware_model")]])
        manufacturer = text(row[cols[header_key("manufacturer")]])
        if not hw_type and not hw_model:
            continue
        record = CatalogRecord(
            hardware_type=hw_type,
            hardware_model=hw_model,
            manufacturer=manufacturer,
            source_row=excel_row,
            type_compact=compact(hw_type),
            model_compact=compact(hw_model),
            type_tokens=tuple(model_tokens(hw_type)),
            model_tokens=tuple(model_tokens(hw_model)),
        )
        records.append(record)
        if record.model_compact:
            exact_index.setdefault(record.model_compact, []).append(record)

    if not records:
        raise ValueError("The IS NW Catalog contains no usable catalog records.")
    return records, exact_index


def fallback_score(cmdb_model: str, record: CatalogRecord) -> int:
    """Rank a candidate only when hardware_type provides genuine evidence."""
    query_compact = compact(cmdb_model)
    query_tokens = model_tokens(cmdb_model)
    if not query_compact or not record.type_compact:
        return 0

    score = 0
    if record.type_compact == query_compact:
        score += 900
    elif record.type_compact in query_compact:
        score += 700 + min(len(record.type_compact), 100)
    elif query_compact in record.type_compact:
        score += 550 + min(len(query_compact), 100)
    else:
        shared_type = set(query_tokens) & set(record.type_tokens)
        if not shared_type:
            return 0
        score += 120 * len(shared_type)

    # Use remaining fragments such as K9 to distinguish multiple type matches.
    shared_model = set(query_tokens) & set(record.model_tokens)
    score += 250 * len(shared_model)

    if (
        record.model_compact
        and len(record.model_compact) >= 2
        and record.model_compact in query_compact
    ):
        score += 350

    for token in query_tokens:
        if token in record.model_compact:
            score += 80
            if query_compact.endswith(token) and record.model_compact.endswith(token):
                score += 100
    return score


def match_model(
    cmdb_model: Any,
    catalog: list[CatalogRecord],
    exact_index: dict[str, list[CatalogRecord]],
) -> MatchResult:
    raw = text(cmdb_model)
    key = compact(raw)
    if not key:
        return MatchResult(None, "BLANK_MODEL")

    exact = exact_index.get(key, [])
    if exact:
        outputs = {
            (r.hardware_type.lower(), r.hardware_model.lower(), r.manufacturer.lower())
            for r in exact
        }
        if len(outputs) == 1:
            return MatchResult(exact[0], "EXACT_HARDWARE_MODEL", 10000)
        return MatchResult(None, "AMBIGUOUS_EXACT_MODEL", 10000, True)

    candidates = [(fallback_score(raw, record), record) for record in catalog]
    candidates = [(score, record) for score, record in candidates if score > 0]
    if not candidates:
        return MatchResult(None, "NO_MATCH")

    candidates.sort(key=lambda item: (-item[0], item[1].source_row))
    top_score = candidates[0][0]
    top_records = [record for score, record in candidates if score == top_score]
    outputs = {
        (r.hardware_type.lower(), r.hardware_model.lower(), r.manufacturer.lower())
        for r in top_records
    }
    if len(outputs) > 1:
        return MatchResult(None, "AMBIGUOUS_HARDWARE_TYPE", top_score, True)
    return MatchResult(top_records[0], "HARDWARE_TYPE_FALLBACK", top_score)


def _write_embedded_template(path: Path) -> None:
    path.write_bytes(base64.b64decode(TEMPLATE_B64.encode("ascii")))


def create_output(catalog_path: Path, cmdb_path: Path, output_path: Path) -> dict[str, int]:
    if catalog_path.resolve() == cmdb_path.resolve():
        raise ValueError("The catalog and CMDB report must be different files.")
    if output_path.suffix.lower() != ".xlsx":
        output_path = output_path.with_suffix(".xlsx")
    if output_path.resolve() in {catalog_path.resolve(), cmdb_path.resolve()}:
        raise ValueError("The output file must be different from both input files.")

    catalog_table = read_table(
        catalog_path, ["hardware_type", "hardware_model", "manufacturer"]
    )
    cmdb_table = read_table(cmdb_path, ["Model number"])
    require_columns(cmdb_table, ["Model number"], "NW CMDB report")
    catalog, exact_index = build_catalog(catalog_table)

    output_path.parent.mkdir(parents=True, exist_ok=True)
    with NamedTemporaryFile(suffix=".xlsx", delete=False) as temp_file:
        temp_path = Path(temp_file.name)
    try:
        _write_embedded_template(temp_path)
        wb = load_workbook(temp_path)
        try:
            ws = wb.active
            template_headers = [text(cell.value) for cell in ws[1]]
            out_cols = {
                header_key(name): index
                for index, name in enumerate(template_headers, start=1)
                if header_key(name)
            }
            required_output = [
                "Model number",
                "Catalog Manufacturer",
                "catalog hardware_type",
                "Catalog hardware_model",
            ]
            missing = [name for name in required_output if header_key(name) not in out_cols]
            if missing:
                raise ValueError(
                    "Embedded output template is missing column(s): " + ", ".join(missing)
                )

            # Clear any residual template data without changing the header layout.
            if ws.max_row > 1:
                ws.delete_rows(2, ws.max_row - 1)

            cmdb_cols = cmdb_table.columns
            shared_columns = [
                (source_index, out_cols[key])
                for key, source_index in cmdb_cols.items()
                if key in out_cols
            ]
            model_index = cmdb_cols[header_key("Model number")]
            stats = {
                "processed": 0,
                "exact": 0,
                "fallback": 0,
                "unmatched": 0,
                "ambiguous": 0,
                "blank_model": 0,
            }

            output_row = 2
            for source_row in cmdb_table.rows:
                if not any(text(value) for value in source_row):
                    continue
                stats["processed"] += 1

                for source_index, destination_column in shared_columns:
                    ws.cell(output_row, destination_column).value = source_row[source_index]

                result = match_model(source_row[model_index], catalog, exact_index)
                if result.record is not None:
                    ws.cell(
                        output_row, out_cols[header_key("Catalog Manufacturer")]
                    ).value = result.record.manufacturer
                    ws.cell(
                        output_row, out_cols[header_key("catalog hardware_type")]
                    ).value = result.record.hardware_type
                    ws.cell(
                        output_row, out_cols[header_key("Catalog hardware_model")]
                    ).value = result.record.hardware_model
                    if result.method == "EXACT_HARDWARE_MODEL":
                        stats["exact"] += 1
                    else:
                        stats["fallback"] += 1
                else:
                    stats["unmatched"] += 1
                    if result.ambiguous:
                        stats["ambiguous"] += 1
                    if result.method == "BLANK_MODEL":
                        stats["blank_model"] += 1
                output_row += 1

            wb.save(output_path)
        finally:
            wb.close()
    finally:
        temp_path.unlink(missing_ok=True)
    return stats


def launch_gui() -> None:
    import tkinter as tk
    from tkinter import filedialog, messagebox, ttk

    class MatcherGUI(tk.Tk):
        def __init__(self) -> None:
            super().__init__()
            self.title(APP_TITLE)
            self.geometry("860x320")
            self.minsize(720, 300)
            self.catalog_var = tk.StringVar()
            self.cmdb_var = tk.StringVar()
            self.status_var = tk.StringVar(
                value="Select the two input files, then click Run Catalog Match."
            )
            self._build()

        @staticmethod
        def file_types():
            return [
                ("Supported files", "*.xlsx *.xlsm *.csv"),
                ("Excel workbooks", "*.xlsx *.xlsm"),
                ("CSV files", "*.csv"),
                ("All files", "*.*"),
            ]

        def _build(self) -> None:
            frame = ttk.Frame(self, padding=20)
            frame.pack(fill="both", expand=True)
            frame.columnconfigure(1, weight=1)

            ttk.Label(
                frame, text=APP_TITLE, font=("Segoe UI", 16, "bold")
            ).grid(row=0, column=0, columnspan=3, sticky="w", pady=(0, 20))

            ttk.Label(frame, text="1. IS NW Catalog file").grid(
                row=1, column=0, sticky="w", padx=(0, 12), pady=8
            )
            ttk.Entry(frame, textvariable=self.catalog_var).grid(
                row=1, column=1, sticky="ew", pady=8
            )
            ttk.Button(frame, text="Browse", command=self.pick_catalog).grid(
                row=1, column=2, padx=(12, 0), pady=8
            )

            ttk.Label(frame, text="2. NW CMDB report").grid(
                row=2, column=0, sticky="w", padx=(0, 12), pady=8
            )
            ttk.Entry(frame, textvariable=self.cmdb_var).grid(
                row=2, column=1, sticky="ew", pady=8
            )
            ttk.Button(frame, text="Browse", command=self.pick_cmdb).grid(
                row=2, column=2, padx=(12, 0), pady=8
            )

            self.progress = ttk.Progressbar(frame, mode="indeterminate")
            self.progress.grid(row=3, column=0, columnspan=3, sticky="ew", pady=(18, 8))
            ttk.Label(frame, textvariable=self.status_var, wraplength=800).grid(
                row=4, column=0, columnspan=3, sticky="w"
            )
            ttk.Button(
                frame,
                text="Run Catalog Match",
                command=self.run_catalog_match,
                width=24,
            ).grid(row=5, column=0, columnspan=3, pady=(20, 0))

        def pick_catalog(self) -> None:
            selected = filedialog.askopenfilename(
                title="Select IS NW Catalog file", filetypes=self.file_types()
            )
            if selected:
                self.catalog_var.set(selected)

        def pick_cmdb(self) -> None:
            selected = filedialog.askopenfilename(
                title="Select NW CMDB report", filetypes=self.file_types()
            )
            if selected:
                self.cmdb_var.set(selected)

        def run_catalog_match(self) -> None:
            catalog_path = Path(self.catalog_var.get().strip())
            cmdb_path = Path(self.cmdb_var.get().strip())
            if not catalog_path.is_file() or not cmdb_path.is_file():
                messagebox.showerror(
                    APP_TITLE, "Please select both valid input files."
                )
                return

            default_name = f"{cmdb_path.stem}_NW_Catalog_Match.xlsx"
            output_name = filedialog.asksaveasfilename(
                title="Save NW Catalog Match output",
                initialdir=str(cmdb_path.parent),
                initialfile=default_name,
                defaultextension=".xlsx",
                filetypes=[("Excel workbook", "*.xlsx")],
            )
            if not output_name:
                return

            self.progress.start(10)
            self.status_var.set("Matching NW hardware catalog records...")
            self.update_idletasks()
            try:
                stats = create_output(
                    catalog_path, cmdb_path, Path(output_name)
                )
                summary = (
                    f"Completed | Processed: {stats['processed']} | "
                    f"Exact: {stats['exact']} | Fallback: {stats['fallback']} | "
                    f"Unmatched: {stats['unmatched']} | "
                    f"Ambiguous: {stats['ambiguous']} | "
                    f"Blank model: {stats['blank_model']}"
                )
                self.status_var.set(summary)
                messagebox.showinfo(
                    APP_TITLE, summary + f"\n\nOutput saved to:\n{output_name}"
                )
            except Exception as exc:
                self.status_var.set("Catalog match failed.")
                messagebox.showerror(
                    APP_TITLE,
                    f"{exc}\n\nTechnical details:\n{traceback.format_exc(limit=6)}",
                )
            finally:
                self.progress.stop()

    MatcherGUI().mainloop()


if __name__ == "__main__":
    launch_gui()
