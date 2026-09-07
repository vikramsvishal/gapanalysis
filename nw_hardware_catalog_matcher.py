#!/usr/bin/env python3
"""Network Hardware Catalog Matcher.

Inputs selected in the GUI:
1) IS NW hardware catalog workbook
2) NW CMDB report workbook
3) NW Catalog Match output template workbook

The program preserves the template layout, copies same-named CMDB columns into it,
and populates Catalog hardware fields using deterministic exact and fallback matching.
"""
from __future__ import annotations

import re
import shutil
import traceback
from dataclasses import dataclass
from pathlib import Path
from typing import Any, Iterable
import tkinter as tk
from tkinter import filedialog, messagebox, ttk

from openpyxl import load_workbook

APP_TITLE = "NW Hardware Catalog Matcher"
DEFAULT_TEMPLATE_NAME = "NW HW Catalog Match File.xlsx"


def clean_text(value: Any) -> str:
    if value is None:
        return ""
    return str(value).strip()


def header_key(value: Any) -> str:
    """Case/spacing/punctuation-insensitive header key."""
    return re.sub(r"[^a-z0-9]+", "", clean_text(value).lower())


def compact(value: Any) -> str:
    """Comparison form with punctuation removed."""
    return re.sub(r"[^a-z0-9]+", "", clean_text(value).lower())


def tokens(value: Any) -> list[str]:
    """Alphanumeric components, also splitting letter/number boundaries."""
    text = clean_text(value).lower()
    chunks = re.findall(r"[a-z0-9]+", text)
    split_chunks = re.findall(r"[a-z]+|[0-9]+", text)
    result: list[str] = []
    for item in chunks + split_chunks:
        if item not in result:
            result.append(item)
    return result


def meaningful_tokens(value: Any) -> list[str]:
    """Discard tiny fragments that cause noisy partial matches."""
    return [t for t in tokens(value) if len(t) >= 2]


def choose_sheet(workbook, required_headers: Iterable[str]):
    required = {header_key(h) for h in required_headers}
    best_ws = None
    best_score = -1
    best_header_row = 1
    # Allow title rows by checking the first 10 rows.
    for ws in workbook.worksheets:
        for row_idx in range(1, min(ws.max_row, 10) + 1):
            keys = {header_key(c.value) for c in ws[row_idx] if clean_text(c.value)}
            score = len(keys & required)
            if score > best_score:
                best_ws, best_score, best_header_row = ws, score, row_idx
    if best_ws is None or best_score == 0:
        raise ValueError("No worksheet with the expected headers was found.")
    return best_ws, best_header_row


def column_map(ws, header_row: int) -> dict[str, int]:
    result: dict[str, int] = {}
    for cell in ws[header_row]:
        key = header_key(cell.value)
        if key and key not in result:
            result[key] = cell.column
    return result


def require_columns(mapping: dict[str, int], required: Iterable[str], source_name: str) -> None:
    missing = [name for name in required if header_key(name) not in mapping]
    if missing:
        raise ValueError(f"{source_name} is missing required column(s): {', '.join(missing)}")


@dataclass(frozen=True)
class CatalogRecord:
    hardware_type: str
    hardware_model: str
    manufacturer: str
    row_number: int
    type_compact: str
    model_compact: str
    type_tokens: tuple[str, ...]
    model_tokens: tuple[str, ...]


@dataclass(frozen=True)
class MatchResult:
    record: CatalogRecord | None
    method: str
    score: int = 0
    ambiguous: bool = False


def load_catalog(ws, header_row: int) -> tuple[list[CatalogRecord], dict[str, list[CatalogRecord]]]:
    cmap = column_map(ws, header_row)
    require_columns(cmap, ["hardware_type", "hardware_model", "manufacturer"], "IS NW Catalog")
    records: list[CatalogRecord] = []
    exact_index: dict[str, list[CatalogRecord]] = {}
    for row in range(header_row + 1, ws.max_row + 1):
        hw_type = clean_text(ws.cell(row, cmap[header_key("hardware_type")]).value)
        hw_model = clean_text(ws.cell(row, cmap[header_key("hardware_model")]).value)
        manufacturer = clean_text(ws.cell(row, cmap[header_key("manufacturer")]).value)
        if not hw_type and not hw_model:
            continue
        rec = CatalogRecord(
            hardware_type=hw_type,
            hardware_model=hw_model,
            manufacturer=manufacturer,
            row_number=row,
            type_compact=compact(hw_type),
            model_compact=compact(hw_model),
            type_tokens=tuple(meaningful_tokens(hw_type)),
            model_tokens=tuple(meaningful_tokens(hw_model)),
        )
        records.append(rec)
        if rec.model_compact:
            exact_index.setdefault(rec.model_compact, []).append(rec)
    if not records:
        raise ValueError("IS NW Catalog contains no usable data rows.")
    return records, exact_index


def score_fallback(cmdb_model: str, rec: CatalogRecord) -> int:
    """Score hardware_type/model evidence without fuzzy-edit-distance guessing."""
    query_compact = compact(cmdb_model)
    query_tokens = meaningful_tokens(cmdb_model)
    if not query_compact or not rec.type_compact:
        return 0

    score = 0
    # The fallback must be grounded in hardware_type.
    if rec.type_compact == query_compact:
        score += 900
    elif rec.type_compact in query_compact:
        score += 700 + min(len(rec.type_compact), 100)
    elif query_compact in rec.type_compact:
        score += 550 + min(len(query_compact), 100)
    else:
        shared_type = set(query_tokens) & set(rec.type_tokens)
        if not shared_type:
            return 0
        score += 120 * len(shared_type)

    # Model fragments such as K9 break ties between multiple type matches.
    qset = set(query_tokens)
    model_shared = qset & set(rec.model_tokens)
    score += 250 * len(model_shared)

    # Reward any catalog model compact fragment explicitly present in CMDB model.
    if rec.model_compact and len(rec.model_compact) >= 2 and rec.model_compact in query_compact:
        score += 350

    # Reward query tokens that appear in the catalog model, especially suffixes.
    for token in query_tokens:
        if token in rec.model_compact:
            score += 80
            if query_compact.endswith(token) and rec.model_compact.endswith(token):
                score += 100
    return score


def match_model(cmdb_model: Any, records: list[CatalogRecord], exact_index: dict[str, list[CatalogRecord]]) -> MatchResult:
    raw = clean_text(cmdb_model)
    key = compact(raw)
    if not key:
        return MatchResult(None, "blank_model")

    exact = exact_index.get(key, [])
    if len(exact) == 1:
        return MatchResult(exact[0], "exact_hardware_model", 10000)
    if len(exact) > 1:
        # Accept duplicate rows only when their business output is identical.
        outputs = {(r.hardware_model.lower(), r.manufacturer.lower(), r.hardware_type.lower()) for r in exact}
        if len(outputs) == 1:
            return MatchResult(exact[0], "exact_hardware_model_duplicate_identical", 10000)
        return MatchResult(None, "ambiguous_exact_hardware_model", 10000, True)

    scored = [(score_fallback(raw, rec), rec) for rec in records]
    scored = [(score, rec) for score, rec in scored if score > 0]
    if not scored:
        return MatchResult(None, "no_match")
    scored.sort(key=lambda x: (-x[0], x[1].row_number))
    top_score = scored[0][0]
    top = [rec for score, rec in scored if score == top_score]

    # Only select a tied result if all tied records produce the same output.
    outputs = {(r.hardware_model.lower(), r.manufacturer.lower(), r.hardware_type.lower()) for r in top}
    if len(outputs) > 1:
        return MatchResult(None, "ambiguous_fallback", top_score, True)
    return MatchResult(top[0], "fallback_hardware_type", top_score)


def first_empty_data_row(ws, header_row: int) -> int:
    """Templates with only headers start at the next row; existing rows are overwritten."""
    return header_row + 1


def run_match(catalog_path: Path, cmdb_path: Path, template_path: Path, output_path: Path) -> dict[str, int]:
    if output_path.resolve() in {catalog_path.resolve(), cmdb_path.resolve(), template_path.resolve()}:
        raise ValueError("Output file must be different from all input files.")
    output_path.parent.mkdir(parents=True, exist_ok=True)
    shutil.copy2(template_path, output_path)

    cat_wb = load_workbook(catalog_path, data_only=True, read_only=True)
    cmdb_wb = load_workbook(cmdb_path, data_only=True, read_only=True)
    out_wb = load_workbook(output_path)
    try:
        cat_ws, cat_hr = choose_sheet(cat_wb, ["hardware_type", "hardware_model", "manufacturer"])
        cmdb_ws, cmdb_hr = choose_sheet(cmdb_wb, ["Model number"])
        out_ws, out_hr = choose_sheet(out_wb, ["Model number", "Catalog Manufacturer", "Catalog hardware_model"])

        records, exact_index = load_catalog(cat_ws, cat_hr)
        cmdb_cols = column_map(cmdb_ws, cmdb_hr)
        out_cols = column_map(out_ws, out_hr)
        require_columns(cmdb_cols, ["Model number"], "NW CMDB report")
        require_columns(out_cols, ["Model number", "Catalog Manufacturer", "Catalog hardware_model"], "NW Catalog Match template")

        # Clear template data values while preserving row styles and workbook layout.
        for row in range(out_hr + 1, out_ws.max_row + 1):
            for col in range(1, out_ws.max_column + 1):
                out_ws.cell(row, col).value = None

        stats = {"processed": 0, "exact": 0, "fallback": 0, "unmatched": 0, "ambiguous": 0, "blank_model": 0}
        out_row = first_empty_data_row(out_ws, out_hr)
        model_col = cmdb_cols[header_key("Model number")]

        # Map all same-named CMDB columns to the template, as requested.
        common = [(cmdb_col, out_cols[key]) for key, cmdb_col in cmdb_cols.items() if key in out_cols]

        for src_row in range(cmdb_hr + 1, cmdb_ws.max_row + 1):
            # Skip completely blank CMDB rows.
            if all(clean_text(cmdb_ws.cell(src_row, c).value) == "" for c in cmdb_cols.values()):
                continue
            stats["processed"] += 1
            for src_col, dst_col in common:
                out_ws.cell(out_row, dst_col).value = cmdb_ws.cell(src_row, src_col).value

            result = match_model(cmdb_ws.cell(src_row, model_col).value, records, exact_index)
            if result.record:
                out_ws.cell(out_row, out_cols[header_key("Catalog Manufacturer")]).value = result.record.manufacturer
                if header_key("catalog hardware_type") in out_cols:
                    out_ws.cell(out_row, out_cols[header_key("catalog hardware_type")]).value = result.record.hardware_type
                out_ws.cell(out_row, out_cols[header_key("Catalog hardware_model")]).value = result.record.hardware_model
                if result.method.startswith("exact"):
                    stats["exact"] += 1
                else:
                    stats["fallback"] += 1
            else:
                stats["unmatched"] += 1
                if result.ambiguous:
                    stats["ambiguous"] += 1
                if result.method == "blank_model":
                    stats["blank_model"] += 1
            out_row += 1

        # Reapply filter to the actual populated range when the template has one.
        if out_ws.auto_filter.ref:
            from openpyxl.utils import get_column_letter
            out_ws.auto_filter.ref = f"A{out_hr}:{get_column_letter(out_ws.max_column)}{max(out_hr, out_row - 1)}"
        out_wb.save(output_path)
        return stats
    finally:
        cat_wb.close()
        cmdb_wb.close()
        out_wb.close()


class MatcherGUI(tk.Tk):
    def __init__(self) -> None:
        super().__init__()
        self.title(APP_TITLE)
        self.geometry("900x390")
        self.minsize(760, 360)
        self.catalog_var = tk.StringVar()
        self.cmdb_var = tk.StringVar()
        default_template = Path(__file__).resolve().with_name(DEFAULT_TEMPLATE_NAME)
        self.template_var = tk.StringVar(value=str(default_template) if default_template.exists() else "")
        self.output_var = tk.StringVar()
        self.status_var = tk.StringVar(value="Select the files, then click Run Match.")
        self._build()

    def _build(self) -> None:
        frame = ttk.Frame(self, padding=18)
        frame.pack(fill="both", expand=True)
        ttk.Label(frame, text=APP_TITLE, font=("Segoe UI", 16, "bold")).grid(row=0, column=0, columnspan=3, sticky="w", pady=(0, 16))
        rows = [
            ("1. IS NW Catalog file", self.catalog_var, self._pick_catalog),
            ("2. NW CMDB report", self.cmdb_var, self._pick_cmdb),
            ("3. NW Catalog Match template", self.template_var, self._pick_template),
            ("4. Output file", self.output_var, self._pick_output),
        ]
        for idx, (label, variable, command) in enumerate(rows, start=1):
            ttk.Label(frame, text=label).grid(row=idx, column=0, sticky="w", padx=(0, 10), pady=7)
            ttk.Entry(frame, textvariable=variable).grid(row=idx, column=1, sticky="ew", pady=7)
            ttk.Button(frame, text="Browse", command=command, width=12).grid(row=idx, column=2, padx=(10, 0), pady=7)
        frame.columnconfigure(1, weight=1)
        self.progress = ttk.Progressbar(frame, mode="indeterminate")
        self.progress.grid(row=5, column=0, columnspan=3, sticky="ew", pady=(18, 8))
        ttk.Label(frame, textvariable=self.status_var, wraplength=820).grid(row=6, column=0, columnspan=3, sticky="w")
        ttk.Button(frame, text="Run Match", command=self._run, width=18).grid(row=7, column=0, columnspan=3, pady=(18, 0))

    @staticmethod
    def _excel_types():
        return [("Excel workbooks", "*.xlsx *.xlsm"), ("All files", "*.*")]

    def _pick_catalog(self):
        p = filedialog.askopenfilename(title="Select IS NW Catalog file", filetypes=self._excel_types())
        if p: self.catalog_var.set(p)

    def _pick_cmdb(self):
        p = filedialog.askopenfilename(title="Select NW CMDB report", filetypes=self._excel_types())
        if p: self.cmdb_var.set(p)

    def _pick_template(self):
        p = filedialog.askopenfilename(title="Select NW Catalog Match template", filetypes=self._excel_types())
        if p: self.template_var.set(p)

    def _pick_output(self):
        p = filedialog.asksaveasfilename(title="Save NW Catalog Match output", defaultextension=".xlsx", filetypes=[("Excel workbook", "*.xlsx")])
        if p: self.output_var.set(p)

    def _run(self):
        fields = [self.catalog_var.get(), self.cmdb_var.get(), self.template_var.get()]
        if not all(fields):
            messagebox.showerror(APP_TITLE, "Please select the catalog, CMDB report, and template files.")
            return
        if not self.output_var.get():
            default = Path(self.cmdb_var.get()).with_name("NW_Catalog_Match_Output.xlsx")
            self.output_var.set(str(default))
        for p in fields:
            if not Path(p).is_file():
                messagebox.showerror(APP_TITLE, f"File not found:\n{p}")
                return
        self.progress.start(10)
        self.status_var.set("Matching catalog records...")
        self.update_idletasks()
        try:
            stats = run_match(Path(fields[0]), Path(fields[1]), Path(fields[2]), Path(self.output_var.get()))
            summary = (
                f"Completed. Processed: {stats['processed']} | Exact: {stats['exact']} | "
                f"Fallback: {stats['fallback']} | Unmatched: {stats['unmatched']} | "
                f"Ambiguous: {stats['ambiguous']} | Blank model: {stats['blank_model']}"
            )
            self.status_var.set(summary)
            messagebox.showinfo(APP_TITLE, summary + f"\n\nOutput:\n{self.output_var.get()}")
        except Exception as exc:
            self.status_var.set("Failed. Review the error message.")
            messagebox.showerror(APP_TITLE, f"{exc}\n\nTechnical details:\n{traceback.format_exc(limit=5)}")
        finally:
            self.progress.stop()


if __name__ == "__main__":
    MatcherGUI().mainloop()
