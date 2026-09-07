"""WPP Enterprise CMDB and Inventory Services Gap Analysis V5 - Complete standalone app.
Requires pandas and openpyxl. Optional: pyxlsb.
"""
from __future__ import annotations
import os,re,ipaddress,traceback,threading,sqlite3,json,hashlib
from pathlib import Path
from datetime import datetime,timezone
import pandas as pd
import tkinter as tk
from tkinter import ttk,filedialog,messagebox
from openpyxl import load_workbook
from openpyxl.styles import Font,PatternFill,Alignment,Border,Side
from openpyxl.utils import get_column_letter

APP="WPP Enterprise CMDB and Inventory Services Gap Analysis V5"
STAMP=datetime.now().strftime("%d-%b-%Y")
SLOTS=["IS Report","NW CMDB Report","Server CMDB Report","Inventory OS Catalog Report","NW Bulk Load Template","Inventory HW Catalog Report","Server Baseline Report","Server Hardware Catalog","Network Hardware Catalog","Field Mapping File"]
TYPES=[("Supported files","*.csv *.xlsx *.xlsb"),("CSV","*.csv"),("Excel","*.xlsx"),("Excel Binary","*.xlsb")]
META={"mandatory","optional","required","recommended","conditional","reference"}

def clean(v): return "" if pd.isna(v) else re.sub(r"\s+"," ",str(v).strip())
def hk(v): return re.sub(r"[\s_.-]+","",clean(v).lower())
def pk(v): return re.sub(r"[^a-z0-9]","",clean(v).lower())
def vk(v):
 s=clean(v).lower();return re.sub(r"\d+",lambda m:str(int(m.group())),s) if s else ""
def sk(v): return re.sub(r"\s+","",clean(v).lower())
def life(v): return {"operational":"operational","endoflife":"end of life","eol":"end of life","production":"production","installed":"installed","deploy":"deploy","design":"design"}.get(pk(v),clean(v).lower())
def read_file(path,meta=False):
 e=Path(path).suffix.lower()
 if e=='.csv': d=pd.read_csv(path,dtype=str,keep_default_na=False,encoding_errors='replace')
 elif e=='.xlsx': d=pd.read_excel(path,dtype=str,keep_default_na=False,engine='openpyxl')
 elif e=='.xlsb': d=pd.read_excel(path,dtype=str,keep_default_na=False,engine='pyxlsb')
 else: raise ValueError('Unsupported file type: '+e)
 d.columns=[clean(c) for c in d.columns];m=None
 if len(d) and any(pk(x) in META for x in d.iloc[0].tolist() if clean(x)): m=d.iloc[0].map(clean);d=d.iloc[1:].reset_index(drop=True)
 for c in d.columns:d[c]=d[c].map(clean)
 return (d,m) if meta else d
def col(d,*names,required=True):
 x={hk(c):c for c in d.columns}
 for n in names:
  if hk(n) in x:return x[hk(n)]
 if required:raise ValueError('Required column not found: '+names[0])
 return None
def split_fw(v):
 s=clean(v);ms=list(re.finditer(r"(?<!\d)(\d+(?:\.\d+)+(?:\([^)]+\))*[A-Za-z0-9_.()\-]*)",s))
 if not ms:return s,''
 m=ms[-1];n=clean(s[:m.start()]);return (n,clean(m.group(1))) if n else (s,'')
def norm_host(v):return clean(v).replace('/','_')
def valid_fqdn(v):
 s=clean(v).replace('/','_').lower().strip('.');return s if pk(s) not in {'','domainnamenotconfigured','domainnamenotapplicable','notapplicable','na','workgroup'} and '.' in s and ' ' not in s else ''
def norm_ip(v):
 s=clean(v)
 if pk(s) in {'stack','standby','na','notapplicable','notconfigured'}:return '','Placeholder removed',''
 found=[]
 for x in re.findall(r'(?<!\d)(?:\d{1,3}\.){3}\d{1,3}(?!\d)',s):
  try:ipaddress.IPv4Address(x);found.append(x)
  except ValueError:pass
 if not found:return '','No valid IPv4 found - review required' if s else '',''
 if len(found)>1:return found[0],'Multiple IP found, kept the first one, review required',', '.join(found[1:])
 return found[0],'' if found[0]==s else 'IP normalized',''
def write_book(path,sheets):
 with pd.ExcelWriter(path,engine='openpyxl') as w:
  for n,d in sheets.items():(d if d is not None and len(d) else pd.DataFrame(columns=['No records'])).to_excel(w,index=False,sheet_name=re.sub(r'[\[\]:*?/\\]','-',n)[:31])
 wb=load_workbook(path);navy=PatternFill('solid',fgColor='1F4E78');white=Font(color='FFFFFF',bold=True);thin=Side(style='thin',color='D9E2F3');bd=Border(left=thin,right=thin,top=thin,bottom=thin)
 for ws in wb.worksheets:
  ws.freeze_panes='A2';ws.sheet_view.showGridLines=False
  if ws.max_row:ws.auto_filter.ref=ws.dimensions
  for c in ws[1]:c.fill=navy;c.font=white;c.alignment=Alignment(horizontal='center',vertical='center',wrap_text=True);c.border=bd
  for row in ws.iter_rows(min_row=2):
   for c in row:c.border=bd;c.alignment=Alignment(vertical='top',wrap_text=True)
  for i,cells in enumerate(ws.columns,1):ws.column_dimensions[get_column_letter(i)].width=min(max(max(len('' if c.value is None else str(c.value)) for c in cells)+2,12),55)
  if ws.title=='Summary':
   for r in range(2,ws.max_row+1):
    metric=str(ws.cell(r,1).value or '')
    if metric=='Load To IS records':ws.cell(r,1).fill=PatternFill('solid',fgColor='E2F0D9')
    if metric=='Retire From IS records':ws.cell(r,1).fill=PatternFill('solid',fgColor='F4CCCC')
 wb.save(path)
def cat_index(cat):
 pc=col(cat,'software_product_name');vc=col(cat,'software_version');pr=col(cat,'software_provider',required=False)
 exact={};base={}
 for _,r in cat.iterrows():
  p={'name':clean(r[pc]),'version':clean(r[vc]),'provider':clean(r[pr]) if pr else ''};exact.setdefault(pk(p['name'])+'|'+vk(p['version']),p);base.setdefault(pk(p['name'])+'|'+re.sub(r'(?<=\d)[a-z]+$','',vk(p['version'])),p)
 return exact,base
def reconcile_nw(cm,isr,cat,progress=lambda m,p:None):
 progress('Resolving NW fields',10)
 names=['Configuration Item','Class','Life Cycle Stage','Life Cycle Stage Status','Manufacturer','Model ID','Model.Name','Model number','Serial number','IP Address','Firmware version','Fully qualified domain name']
 mp={n:col(cm,n,'Life Cycle Stag' if n=='Life Cycle Stage' else n,required=n!='Fully qualified domain name') for n in names};out=pd.DataFrame({n:cm[c].map(clean) if c else '' for n,c in mp.items()});q=out['Firmware version'].map(split_fw).apply(pd.Series);q.columns=['Parsed Firmware Name','Parsed Firmware Version'];out=pd.concat([out,q],axis=1)
 ex,ba=cat_index(cat);hc=col(isr,'hostname','host_name');sc=col(isr,'os_parent_serial_number');lc=col(isr,'os_lifecycle_status');oid=col(isr,'os_opaque_id','OS Opaque ID',required=False);onc=col(isr,'operating_system_name');ovc=col(isr,'operating_system_version');opc=col(isr,'operating_system_provider',required=False)
 hm={};sm={}
 for i,r in isr.iterrows():
  if clean(r[hc]).lower():hm.setdefault(clean(r[hc]).lower(),i)
  if sk(r[sc]):sm.setdefault(sk(r[sc]),i)
 rows=[];progress('Reconciling NW CMDB and IS',45)
 for _,r in out.iterrows():
  ci=clean(r['Configuration Item']);idx=next((hm[x] for x in [ci.lower(),ci.replace('/','_').lower(),ci.replace('/','-').lower()] if x in hm),None);mode='Hostname Match' if idx is not None else ''
  if idx is None and sk(r['Serial number']) in sm:idx=sm[sk(r['Serial number'])];mode='Serial Match Different Hostname'
  present=idx is not None;ir=isr.iloc[idx] if present else None;k=pk(r['Parsed Firmware Name'])+'|'+vk(r['Parsed Firmware Version']);hit=ex.get(k);status='Match'
  if not hit:hit=ba.get(pk(r['Parsed Firmware Name'])+'|'+re.sub(r'(?<=\d)[a-z]+$','',vk(r['Parsed Firmware Version'])));status='Closest Match Proposed' if hit else 'No Match';hit=hit or {}
  osmatch=present and hit and pk(ir[onc])==pk(hit.get('name')) and vk(ir[ovc])==vk(hit.get('version'));acts=[]
  if status=='No Match':acts.append('Catalogue update required')
  if present and hit and not osmatch:acts.append('Inventory OS update required')
  if life(r['Life Cycle Stage'])=='operational' and not present:acts.append('Load To IS')
  if life(r['Life Cycle Stage'])=='end of life' and present and life(ir[lc]) in {'production','installed'}:acts.append('Retire From IS')
  if mode.startswith('Serial'):acts.append('Hostname review required')
  d=r.to_dict();d.update({'Present in IS ?':'Yes - '+mode if present else 'No','Host name in IS':clean(ir[hc]) if present else '','os_opaque_id':clean(ir[oid]) if present and oid else '','os_parent_serial_number':clean(ir[sc]) if present else '','Lifecycle Status in IS':clean(ir[lc]) if present else '','Current operating_system_provider':clean(ir[opc]) if present and opc else '','Current operating_system_name':clean(ir[onc]) if present else '','Current operating_system_version':clean(ir[ovc]) if present else '','Catalogue Match?':status,'Required operating_system_provider':hit.get('provider',''),'Required operating_system_name':hit.get('name',''),'Required operating_system_version':hit.get('version',''),'Action':'; '.join(acts) or 'No Action'});rows.append(d)
 progress('NW reconciliation complete',100);return pd.DataFrame(rows)
def summary(rec):
 return pd.DataFrame([('NW CMDB records analyzed',len(rec)),('Records present in IS',int(rec['Present in IS ?'].astype(str).str.startswith('Yes').sum())),('Records not present in IS',int((rec['Present in IS ?']=='No').sum())),('Catalogue exact matches',int((rec['Catalogue Match?']=='Match').sum())),('Catalogue closest matches proposed',int((rec['Catalogue Match?']=='Closest Match Proposed').sum())),('Catalogue updates required',int(rec['Action'].str.contains('Catalogue update required',na=False).sum())),('Inventory OS updates required',int(rec['Action'].str.contains('Inventory OS update required',na=False).sum())),('Load To IS records',int(rec['Action'].str.contains('Load To IS',na=False).sum())),('Retire From IS records',int(rec['Action'].str.contains('Retire From IS',na=False).sum()))],columns=['Metric','Value'])

# Flag learning repository: dynamic Mandatory Y/N pattern counts by cohort.
class Flags:
 def __init__(self,path):self.c=sqlite3.connect(path);self.c.executescript('CREATE TABLE IF NOT EXISTS versions(hash TEXT PRIMARY KEY,created TEXT,source TEXT);CREATE TABLE IF NOT EXISTS obs(field TEXT,value TEXT,cohort TEXT);');self.c.commit()
 def ingest(self,path):
  h=hashlib.sha256(Path(path).read_bytes()).hexdigest()
  if self.c.execute('SELECT 1 FROM versions WHERE hash=?',(h,)).fetchone():return
  d,m=read_file(path,True);m=m if m is not None else pd.Series('',index=d.columns);flags=[c for c in d if clean(m.get(c,'')).lower()=='mandatory' and {clean(x).upper() for x in d[c] if clean(x)}.issubset({'Y','N'}) and any(clean(x) for x in d[c])]
  for _,r in d.iterrows():
   cohort='|'.join([pk(r.get(col,'')) for col in ['operating_system_subcategory','parent_system_type','image_purpose','os_lifecycle_status','operating_system_name']])
   for f in flags:
    v=clean(r[f]).upper()
    if v in {'Y','N'}:self.c.execute('INSERT INTO obs VALUES(?,?,?)',(hk(f),v,cohort))
  self.c.execute('INSERT INTO versions VALUES(?,?,?)',(h,datetime.now(timezone.utc).isoformat(),Path(path).name));self.c.commit()
 def predict(self,field,cohort):
  # progressively broaden by truncating dimensions
  parts=cohort.split('|')
  for n in range(len(parts),0,-1):
   prefix='|'.join(parts[:n]);vals=[x[0] for x in self.c.execute('SELECT value FROM obs WHERE field=? AND cohort LIKE ?',(hk(field),prefix+'%')).fetchall()]
   if vals:
    y=vals.count('Y');z=vals.count('N');support=y+z;agr=max(y,z)/support;val='Y' if y>=z else 'N'
    return val,agr,support,'Auto-prefill' if agr>=.9 and support>=20 else 'Suggest for review' if agr>=.75 else 'No proposal'
  return '',0,0,'No evidence'
 def close(self):self.c.close()
def bulk_load(rec,ispath,template,outdir,fmt,askfq,askcat,repo,progress=lambda m,p:None):
 cand=rec[(rec['Present in IS ?']=='No')&(rec['Life Cycle Stage'].map(life)=='operational')&rec.Action.str.contains('Load To IS',na=False)].copy();dupkey=cand['Configuration Item'].map(pk)+'|'+cand['Life Cycle Stage'].map(life)+'|'+cand['Serial number'].map(sk);bad=dupkey.duplicated(False);dups=cand[bad];cand=cand[~bad]
 un=cand[cand['Catalogue Match?']=='No Match'];catpath=''
 if len(un):
  if askcat(len(un)):catpath=os.path.join(outdir,f'Catalogue Update Required - {STAMP}.xlsx');write_book(catpath,{'Catalogue Update Required':un})
  cand=cand[cand['Catalogue Match?']!='No Match']
 fq=cand['Fully qualified domain name'].map(valid_fqdn);miss=fq.eq('');fqex=pd.DataFrame()
 if miss.any():
  if askfq(int(miss.sum())):fq.loc[miss]=cand.loc[miss,'Configuration Item'].map(lambda x:norm_host(x).lower()+'.nw.wpp.net')
  else:fqex=cand[miss];cand=cand[~miss];fq=fq[~miss]
 td,meta=read_file(template,True);headers=list(td.columns);extras=['CMDB Hostname','CMDB Serial Number','CMDB Lifecycle Stage'];rows=[];iprows=[]
 fl=Flags(repo);fl.ingest(ispath);mandatory_flags=[h for h in headers if meta is not None and clean(meta.get(h,'')).lower()=='mandatory'];analysis=[]
 for ix,r in cand.iterrows():
  ip,msg,more=norm_ip(r['IP Address']);vals={'hostname':norm_host(r['Configuration Item']),'fully_qualified_hostname':fq.loc[ix].lower(),'os_parent_serial_number':r['Serial number'],'os_lifecycle_status':'Production','parent_system_type':'Appliance' if pk(r['Class']) in {'wirelessaccesspoint','interfacecard'} else 'Network','operating_system_provider':r['Required operating_system_provider'],'operating_system_name':r['Required operating_system_name'],'operating_system_version':r['Required operating_system_version'],'ip_address':ip};row={h:'' for h in headers};keys={hk(k):v for k,v in vals.items()}
  for h in headers:
   if hk(h) in keys:row[h]=keys[hk(h)]
  cohort='|'.join([pk('Network'),pk(row.get(next((h for h in headers if hk(h)==hk('parent_system_type')),''),'')),pk(r['Class']),pk('Production'),pk(r['Required operating_system_name'])])
  for h in mandatory_flags:
   val,agr,sup,decision=fl.predict(h,cohort)
   if clean(row[h]).upper() not in {'Y','N'} and decision=='Auto-prefill':row[h]=val;applied='Y'
   else:applied='N'
   if val:analysis.append({'CMDB Hostname':r['Configuration Item'],'Field':h,'Proposed Value':val,'Agreement':agr,'Support Count':sup,'Decision':decision,'Applied':applied})
  rows.append({extras[0]:r['Configuration Item'],extras[1]:r['Serial number'],extras[2]:r['Life Cycle Stage'],**row});iprows.append({'CMDB Hostname':r['Configuration Item'],'Original IP':r['IP Address'],'Selected IP':ip,'Additional IPs':more,'Message':msg})
 fl.close();data=pd.DataFrame(rows,columns=extras+headers);md=['Reference']*3+([clean(meta[h]) for h in headers] if meta is not None else ['']*len(headers));final=pd.concat([pd.DataFrame([md],columns=extras+headers),data],ignore_index=True);ext='xlsx' if fmt=='xlsx' else 'csv';loadpath=os.path.join(outdir,f'NW Bulk Load - {STAMP}.{ext}');final.to_excel(loadpath,index=False) if ext=='xlsx' else final.to_csv(loadpath,index=False);valpath=os.path.join(outdir,f'NW Bulk Load Validation - {STAMP}.xlsx');write_book(valpath,{'Flag Intelligence':pd.DataFrame(analysis),'IP Validation':pd.DataFrame(iprows),'CMDB Correction Required':dups,'Missing FQDN Excluded':fqex});return [x for x in [loadpath,valpath,catpath] if x]

class Scroll(ttk.Frame):
 def __init__(self,p):super().__init__(p);self.c=tk.Canvas(self,bg='#F4F6F8',highlightthickness=0);self.s=ttk.Scrollbar(self,command=self.c.yview);self.f=ttk.Frame(self.c,padding=18);self.w=self.c.create_window((0,0),window=self.f,anchor='nw');self.f.bind('<Configure>',lambda e:self.c.configure(scrollregion=self.c.bbox('all')));self.c.bind('<Configure>',lambda e:self.c.itemconfigure(self.w,width=e.width));self.c.configure(yscrollcommand=self.s.set);self.c.pack(side='left',fill='both',expand=True);self.s.pack(side='right',fill='y')
class App(tk.Tk):
 def __init__(self):super().__init__();self.title(APP);self.geometry('1450x850');self.paths={s:'' for s in SLOTS};self.labels={};self.results={};self.busy=False;self.repo=str(Path(__file__).with_name('v5_flag_intelligence.sqlite'));self.build()
 def build(self):
  sf=Scroll(self);sf.pack(fill='both',expand=True);root=sf.f;ttk.Label(root,text=APP,font=('Segoe UI',19,'bold'),foreground='#1F4E78').pack(anchor='w');ttk.Label(root,text='Complete Network and Server reconciliation, bulk-load, delta update and Priority 1 flag intelligence').pack(anchor='w',pady=(0,12));pan=ttk.PanedWindow(root,orient='horizontal');pan.pack(fill='both',expand=True);left=ttk.Frame(pan);right=ttk.Frame(pan);pan.add(left,weight=3);pan.add(right,weight=4);lf=ttk.LabelFrame(left,text='File Loader',padding=10);lf.pack(fill='x')
  for i,s in enumerate(SLOTS):ttk.Label(lf,text=s,width=29).grid(row=i,column=0,sticky='w',pady=3);ttk.Button(lf,text='Browse',command=lambda x=s:self.browse(x)).grid(row=i,column=1,padx=3);ttk.Button(lf,text='Clear',command=lambda x=s:self.clear(x)).grid(row=i,column=2,padx=3);l=tk.Label(lf,text='Not loaded',fg='#9C0006',anchor='w');l.grid(row=i,column=3,sticky='ew');self.labels[s]=l
  lf.columnconfigure(3,weight=1);af=ttk.LabelFrame(left,text='Dedicated Actions',padding=10);af.pack(fill='x',pady=10)
  for t,c in [('Reconcile NW CMDB and IS',self.nw),('Reconcile Server CMDB and IS',self.server),('Create NW Bulk Load file',self.load),('Create NW Bulk Update file',self.update)]:ttk.Button(af,text=t,command=c,width=38).pack(anchor='w',pady=4)
  ctl=ttk.Frame(left);ctl.pack(fill='x');ttk.Button(ctl,text='Clear All Files',command=self.clearall).pack(side='left');ttk.Button(ctl,text='Close App',command=self.destroy).pack(side='right');self.pb=ttk.Progressbar(left,maximum=100);self.pb.pack(fill='x',pady=(12,3));self.pt=ttk.Label(left,text='Progress: 0% - Ready');self.pt.pack(anchor='w');rf=ttk.LabelFrame(right,text='Activity Log',padding=8);rf.pack(fill='both',expand=True);self.log=tk.Text(rf,wrap='word',font=('Consolas',10));ys=ttk.Scrollbar(rf,command=self.log.yview);self.log.configure(yscrollcommand=ys.set);self.log.pack(side='left',fill='both',expand=True);ys.pack(side='right',fill='y');self.say('Application ready.')
 def say(self,x):
  if threading.current_thread() is not threading.main_thread():self.after(0,lambda:self.say(x));return
  self.log.insert('end',f'[{datetime.now():%H:%M:%S}] {x}\n');self.log.see('end')
 def prog(self,x,n):
  if threading.current_thread() is not threading.main_thread():self.after(0,lambda:self.prog(x,n));return
  self.pb['value']=n;self.pt.configure(text=f'Progress: {n}% - {x}');self.say(x)
 def browse(self,s):
  p=filedialog.askopenfilename(title='Select '+s,filetypes=TYPES)
  if p:self.paths[s]=p;self.labels[s].configure(text=Path(p).name,fg='#006100');self.say('Loaded '+s+': '+p)
 def clear(self,s):self.paths[s]='';self.labels[s].configure(text='Not loaded',fg='#9C0006')
 def clearall(self):
  for s in SLOTS:self.clear(s)
 def require(self,*s):
  m=[x for x in s if not self.paths[x]]
  if m:messagebox.showerror('Missing files','Load:\n'+'\n'.join(m));return False
  return True
 def run(self,fn):
  if self.busy:return
  self.busy=True
  def work():
   try:fn()
   except Exception as exc:
    err=exc;tb=traceback.format_exc();self.after(0,lambda e=err,t=tb:self.fail(e,t))
   finally:self.busy=False
  threading.Thread(target=work,daemon=True).start()
 def fail(self,e,t):self.say(t);self.prog('Failed',0);messagebox.showerror('Operation failed',str(e))
 def syncyes(self,title,text):
  a=[];ev=threading.Event();self.after(0,lambda:(a.append(messagebox.askyesno(title,text)),ev.set()));ev.wait();return a[0]
 def nw(self):
  if not self.require('NW CMDB Report','IS Report','Inventory OS Catalog Report'):return
  out=filedialog.askdirectory(title='Select output folder')
  if not out:return
  def task():
   r=reconcile_nw(read_file(self.paths['NW CMDB Report']),read_file(self.paths['IS Report']),read_file(self.paths['Inventory OS Catalog Report']),self.prog);self.results['nw']=r;p=os.path.join(out,f'NW CMDB to IS Reconciliation V5 - {STAMP}.xlsx');write_book(p,{'Summary':summary(r),'Reconciliation':r,'Load To IS':r[r.Action.str.contains('Load To IS',na=False)],'Retire From IS':r[r.Action.str.contains('Retire From IS',na=False)],'Inventory OS Update Required':r[r.Action.str.contains('Inventory OS update required',na=False)],'Catalogue Update Required':r[r.Action.str.contains('Catalogue update required',na=False)]});self.say('Created: '+p);self.after(0,lambda:messagebox.showinfo('Complete',p))
  self.run(task)
 def ensure_nw(self):
  if 'nw' in self.results:return self.results['nw']
  if not self.require('NW CMDB Report','IS Report','Inventory OS Catalog Report'):return None
  self.prog('Running NW reconciliation required for bulk-load creation',5);r=reconcile_nw(read_file(self.paths['NW CMDB Report']),read_file(self.paths['IS Report']),read_file(self.paths['Inventory OS Catalog Report']),self.prog);self.results['nw']=r;return r
 def load(self):
  if not self.require('IS Report','NW CMDB Report','Inventory OS Catalog Report'):return
  if not self.paths['NW Bulk Load Template']:
   self.browse('NW Bulk Load Template')
   if not self.paths['NW Bulk Load Template']:messagebox.showwarning('Template required','Bulk-load creation was cancelled because no template was selected.');return
  out=filedialog.askdirectory(title='Select output folder')
  if not out:return
  fmt='xlsx' if messagebox.askyesno('Output format','Create XLSX? Select No for CSV.') else 'csv'
  def task():
   r=self.ensure_nw();
   if r is None:return
   paths=bulk_load(r,self.paths['IS Report'],self.paths['NW Bulk Load Template'],out,fmt,lambda n:self.syncyes('Missing FQDN',f'{n} devices have no valid FQDN. Generate <hostname>.nw.wpp.net? Select No to exclude and continue.'),lambda n:self.syncyes('Catalogue Match',f'{n} devices have no catalogue match. Create Catalogue Update Required file?'),self.repo,self.prog)
   for p in paths:self.say('Created: '+p)
   self.after(0,lambda:messagebox.showinfo('Complete','\n'.join(paths)))
  self.run(task)
 def server(self):
  if not self.require('Server CMDB Report','IS Report','Inventory OS Catalog Report'):return
  out=filedialog.askdirectory(title='Select output folder')
  if not out:return
  def task():
   cm=read_file(self.paths['Server CMDB Report']);isr=read_file(self.paths['IS Report']);cat=read_file(self.paths['Inventory OS Catalog Report']);sg=col(cm,'Support Group',required=False);mg=col(cm,'Managed by Group',required=False);scope=pd.Series(False,index=cm.index)
   if sg:scope|=cm[sg].map(clean).str.upper().str.startswith('W-KYN')
   if mg:scope|=cm[mg].map(clean).str.upper().str.startswith('W-KYN')
   scoped=cm[scope].copy();outside=cm[~scope].copy();namec=col(scoped,'name','Server Name');hostc=col(scoped,'host_name','Host Name','hostname',required=False);ipc=col(scoped,'ip_address','IP Address');classc=col(scoped,'sys_class_name','Class');catc=col(scoped,'category');subc=col(scoped,'subcategory');lifec=col(scoped,'life_cycle_stage','Life Cycle Stage');lsc=col(scoped,'life_cycle_stage_status','Life Cycle Stage Status');osc=col(scoped,'os','Operating System');osvc=col(scoped,'os_version','OS Version')
   b=pd.DataFrame({'name':scoped[namec],'host_name':scoped[hostc] if hostc else scoped[namec],'ip_address':scoped[ipc],'sys_class_name':scoped[classc],'category':scoped[catc],'subcategory':scoped[subc],'life_cycle_stage':scoped[lifec],'life_cycle_stage_status':scoped[lsc],'os':scoped[osc],'os_version':scoped[osvc]});b['_name']=b.name.map(pk);b['_life']=b.life_cycle_stage.map(life);dups=b[b.duplicated(['_name','_life'],False)].copy();b=b.drop_duplicates(['_name','_life']);active=set(b.loc[b._life=='operational','_name']);hist=b[(b._life=='end of life')&b._name.isin(active)].copy();b=b[~((b._life=='end of life')&b._name.isin(active))]
   temp=pd.DataFrame({'Configuration Item':b.name,'Class':b.sys_class_name,'Life Cycle Stage':b.life_cycle_stage,'Life Cycle Stage Status':b.life_cycle_stage_status,'Manufacturer':'','Model ID':'','Model.Name':'','Model number':'','Serial number':'','IP Address':b.ip_address,'Firmware version':b.os+' '+b.os_version,'Fully qualified domain name':b.host_name});r=reconcile_nw(temp,isr,cat,self.prog);result=pd.concat([b.drop(columns=['_name','_life']).reset_index(drop=True),r.drop(columns=[c for c in temp if c in r],errors='ignore').reset_index(drop=True)],axis=1);self.results['server']=result;p=os.path.join(out,f'Server CMDB to IS Reconciliation V5 - {STAMP}.xlsx');write_book(p,{'Reconciliation':result,'Load To IS':result[result.Action.str.contains('Load To IS',na=False)],'Retire From IS':result[result.Action.str.contains('Retire From IS',na=False)],'Out of Kyndryl Scope':outside,'Historical EOL Suppressed':hist,'CMDB Duplicate Review':dups});self.say('Created: '+p);self.after(0,lambda:messagebox.showinfo('Complete',p))
  self.run(task)
 def update(self):
  r=self.results.get('nw')
  if r is None:
   p=filedialog.askopenfilename(title='Select NW reconciliation workbook',filetypes=[('Excel','*.xlsx')])
   if not p:return
   r=pd.read_excel(p,sheet_name='Reconciliation');self.results['nw']=r
  pairs={'operating_system_provider':('Current operating_system_provider','Required operating_system_provider',pk),'operating_system_name':('Current operating_system_name','Required operating_system_name',pk),'operating_system_version':('Current operating_system_version','Required operating_system_version',vk)};eligible=r[r['Present in IS ?'].astype(str).str.startswith('Yes') & r['os_opaque_id'].astype(str).ne('')].copy();available={}
  for f,(a,b,norm) in pairs.items():
   if a in eligible and b in eligible:
    count=int((eligible[a].map(norm)!=eligible[b].map(norm)).sum())
    if count:available[f]=(a,b,norm,count)
  if not available:messagebox.showinfo('No updates','No fields have update candidates.');return
  w=tk.Toplevel(self);w.title('Select NW update fields');vars={}
  for f,(_,_,_,count) in available.items():vars[f]=tk.BooleanVar(value=True);ttk.Checkbutton(w,text=f'{f} ({count} candidates)',variable=vars[f]).pack(anchor='w',padx=15,pady=4)
  def create():
   chosen=[f for f,v in vars.items() if v.get()]
   if not chosen:return
   path=filedialog.asksaveasfilename(defaultextension='.csv',filetypes=[('CSV','*.csv'),('Excel','*.xlsx')],initialfile=f'NW Bulk Update - {STAMP}.csv')
   if not path:return
   mask=pd.Series(False,index=eligible.index)
   for f in chosen:
    a,b,norm,_=available[f];mask|=eligible[a].map(norm)!=eligible[b].map(norm)
   src=eligible[mask];out=pd.DataFrame({'os_opaque_id':src.os_opaque_id})
   for f in chosen:out[f]=src[available[f][1]]
   out.to_excel(path,index=False) if path.lower().endswith('.xlsx') else out.to_csv(path,index=False);w.destroy();self.say('Created: '+path);messagebox.showinfo('Complete',path)
  ttk.Button(w,text='Create Update File',command=create).pack(pady=10)
if __name__=='__main__':App().mainloop()
