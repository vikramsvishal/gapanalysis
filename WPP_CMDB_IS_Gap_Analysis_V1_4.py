"""WPP Enterprise CMDB and Inventory Services Gap Analysis V1.4

Standalone Tkinter application. Requires pandas and openpyxl. Optional: pyxlsb.
The application is deliberately layered: engines below have no Tkinter dependency and
can be imported by an agent, FastAPI service, Azure Function, or test runner.
"""
from __future__ import annotations
import os, re, json, sqlite3, hashlib, ipaddress, traceback, threading
from pathlib import Path
from datetime import datetime, date
from typing import Any, Dict, List, Optional, Tuple
from dataclasses import dataclass
from collections import defaultdict
import pandas as pd
from openpyxl import load_workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter
import tkinter as tk
from tkinter import ttk, filedialog, messagebox

APP_NAME = "WPP Enterprise CMDB and Inventory Services Gap Analysis V1.4"
APP_DIR = Path(__file__).resolve().parent
REGISTRY_PATH = APP_DIR / "v1_4_resource_registry.json"
LEARNING_DB = APP_DIR / "v1_4_flag_intelligence.sqlite"
DATE_LABEL = datetime.now().strftime("%d-%b-%Y")
DISPLAY_DATE = datetime.now().strftime("%d.%b.%Y").upper()
FILE_TYPES = [("Supported files", "*.csv *.xlsx *.xlsm *.xlsb"), ("CSV", "*.csv"), ("Excel", "*.xlsx *.xlsm"), ("Excel Binary", "*.xlsb"), ("All files", "*.*")]

PERSISTENT_RESOURCES = {
    "Bulk Load Template": "bulk_load_template",
    "IS Product Catalog Operating System": "catalog_os",
    "IS Product Catalog Network": "catalog_network",
    "IS Product Catalog Server": "catalog_server",
    "Field List and Mapping": "field_mapping",
}
SESSION_RESOURCES = ["IS Operating System Report", "NW CMDB Report", "Server CMDB Report", "IS Network Category Report", "IS Server Category Report"]
ALL_RESOURCES = SESSION_RESOURCES + list(PERSISTENT_RESOURCES)

# Microsoft-published lifecycle dates captured from Windows Server release information.
# V5.2 treats versions as supported until extended support end. Refresh through code/config
# when Microsoft publishes changed lifecycle information.
WINDOWS_SERVER_SUPPORT = {
    "2025": ("10.0.26100", date(2034, 11, 14)),
    "2022": ("10.0.20348", date(2031, 10, 14)),
    "2019": ("10.0.17763", date(2029, 1, 9)),
    "2016": ("10.0.14393", date(2027, 1, 12)),
    "2012 r2": ("6.3.9600", date(2023, 10, 10)),
    "2012": ("6.2.9200", date(2023, 10, 10)),
    "2008 r2": ("6.1.7601", date(2020, 1, 14)),
}

# ---------------------------- reusable core utilities ----------------------------
def clean(v: Any) -> str:
    return "" if pd.isna(v) else re.sub(r"\s+", " ", str(v).strip())

def hkey(v: Any) -> str:
    return re.sub(r"[\s_.-]+", "", clean(v).lower())

def pkey(v: Any) -> str:
    return re.sub(r"[^a-z0-9]", "", clean(v).lower())

def vkey(v: Any) -> str:
    s = clean(v).lower()
    return re.sub(r"\d+", lambda m: str(int(m.group())), s) if s else ""

def lifecycle(v: Any) -> str:
    return {"operational":"operational", "endoflife":"end of life", "eol":"end of life",
            "production":"production", "installed":"installed", "deploy":"deploy", "design":"design"}.get(pkey(v), clean(v).lower())

def serial_key(v: Any) -> str:
    return re.sub(r"[^a-z0-9]", "", clean(v).lower())

def normalize_hostname(v: Any) -> str:
    return clean(v).replace("/", "_")

def normalize_fqdn(v: Any) -> str:
    return clean(v).replace("/", "_").lower().strip(".")

def valid_fqdn(v: Any) -> str:
    s = normalize_fqdn(v)
    invalid = {"", "domainnamenotconfigured", "domainnamenotapplicable", "notapplicable", "na", "workgroup", "workgroupserver"}
    return s if pkey(s) not in invalid and "." in s and " " not in s else ""

def normalize_ipv4(v: Any) -> Tuple[str, str, str]:
    raw = clean(v)
    if pkey(raw) in {"stack", "standby", "na", "notapplicable", "notconfigured", "none"}:
        return "", "Placeholder removed", ""
    found = []
    for token in re.findall(r"(?<!\d)(?:\d{1,3}\.){3}\d{1,3}(?!\d)", raw):
        try:
            ipaddress.IPv4Address(token); found.append(token)
        except ValueError:
            pass
    if not found:
        return "", ("No valid IPv4 found - review required" if raw else ""), ""
    if len(found) > 1:
        return found[0], "Multiple IP found, kept the first one, review required", ", ".join(found[1:])
    return found[0], ("IP normalized" if found[0] != raw else ""), ""

def split_product_version(v: Any) -> Tuple[str, str]:
    raw=clean(v)
    # Prefer a trailing dotted/build version, but also support a trailing four-digit
    # server release such as Microsoft Windows Server 2022.
    match=re.search(r"(?<!\d)(\d{4}|\d+(?:\.\d+)+(?:\([^)]+\))*[A-Za-z0-9_.()\-]*)\s*$",raw)
    if not match:return raw,""
    name=clean(raw[:match.start()]);version=clean(match.group(1))
    return (name,version) if name else (raw,"")



def read_tabular(path: str, keep_metadata: bool=False):
    ext = Path(path).suffix.lower()
    if ext == ".csv":
        df = pd.read_csv(path, dtype=str, keep_default_na=False, encoding_errors="replace")
    elif ext in {".xlsx", ".xlsm"}:
        df = pd.read_excel(path, dtype=str, keep_default_na=False, engine="openpyxl")
    elif ext == ".xlsb":
        df = pd.read_excel(path, dtype=str, keep_default_na=False, engine="pyxlsb")
    else:
        raise ValueError(f"Unsupported file type: {ext}")
    df.columns = [clean(c) for c in df.columns]
    metadata = None
    if len(df):
        terms = {pkey(x) for x in df.iloc[0].tolist() if clean(x)}
        if terms & {"mandatory", "optional", "reference", "required", "recommended"}:
            metadata = df.iloc[0].map(clean)
            df = df.iloc[1:].reset_index(drop=True)
    for c in df.columns:
        df[c] = df[c].map(clean)
    return (df, metadata) if keep_metadata else df

def find_col(df: pd.DataFrame, *aliases: str, required=True) -> Optional[str]:
    index = {hkey(c): c for c in df.columns}
    for alias in aliases:
        if hkey(alias) in index:
            return index[hkey(alias)]
    if required:
        raise ValueError("Required field not found. Expected one of: " + ", ".join(aliases))
    return None

def next_versioned_path(folder: str, stem: str, extension: str) -> str:
    """Return the next writable daily version without overwriting an existing file."""
    target_dir = Path(folder).expanduser().resolve()
    target_dir.mkdir(parents=True, exist_ok=True)
    if not target_dir.is_dir():
        raise NotADirectoryError(f"Output location is not a folder: {target_dir}")

    # Fail early with a useful message if Windows/OneDrive denies folder writes.
    probe = target_dir / f".__wpp_write_test_{datetime.now():%Y%m%d%H%M%S%f}.tmp"
    try:
        with open(probe, "x", encoding="utf-8") as handle:
            handle.write("write-test")
    except PermissionError as exc:
        raise PermissionError(
            f"The selected output folder is not writable: {target_dir}. "
            "Choose a writable folder such as Desktop or Documents, or review Windows folder protection."
        ) from exc
    finally:
        try:
            probe.unlink(missing_ok=True)
        except OSError:
            pass

    ext = extension.lstrip(".")
    date_label = datetime.now().strftime("%d-%b-%Y")
    for version in range(1, 10000):
        candidate = target_dir / f"{stem} - {date_label} - V{version}.{ext}"
        if candidate.exists():
            continue
        # Reserve and immediately release the name. This verifies the exact filename
        # can be created and skips a locked or protected version when necessary.
        try:
            with open(candidate, "x+b"):
                pass
            candidate.unlink()
            return str(candidate)
        except FileExistsError:
            continue
        except PermissionError:
            continue
    raise RuntimeError(f"Unable to allocate a writable output filename in: {target_dir}")
























_OUTPUT_PROGRESS_CALLBACK = None

def set_output_progress(callback=None) -> None:
    global _OUTPUT_PROGRESS_CALLBACK
    _OUTPUT_PROGRESS_CALLBACK = callback

def _output_progress(message: str, percentage: int) -> None:
    callback = _OUTPUT_PROGRESS_CALLBACK
    if callback:
        callback(message, percentage)

def _next_retry_path(path: str) -> str:
    candidate = Path(path)
    match = re.match(r"^(.*) - V(\d+)(\.[^.]+)$", candidate.name, re.I)
    if not match:
        return str(candidate.with_name(candidate.stem + " - V2" + candidate.suffix))
    stem, version, suffix = match.groups()
    return str(candidate.with_name(f"{stem} - V{int(version)+1}{suffix}"))

def _fast_format_workbook(workbook) -> None:
    """Apply presentation formatting without iterating through every data cell."""
    navy = PatternFill("solid", fgColor="1F4E78")
    metadata_fill = PatternFill("solid", fgColor="D9EAF7")
    load_fill = PatternFill("solid", fgColor="E2F0D9")
    retire_fill = PatternFill("solid", fgColor="F4CCCC")
    retired_fill = PatternFill("solid", fgColor="D9D9D9")
    white_bold = Font(name="Aptos", size=10, color="FFFFFF", bold=True)
    metadata_font = Font(name="Aptos", size=9, color="1F1F1F", italic=True)
    thin = Side(style="thin", color="D9E2F3")
    header_border = Border(left=thin, right=thin, top=thin, bottom=thin)

    for worksheet in workbook.worksheets:
        worksheet.sheet_view.showGridLines = False
        worksheet.sheet_view.zoomScale = 90
        worksheet.freeze_panes = "A2"
        worksheet.row_dimensions[1].height = 34
        for cell in worksheet[1]:
            cell.fill = navy; cell.font = white_bold; cell.border = header_border
            cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)

        metadata_row = False
        if worksheet.max_row >= 2:
            values = [pkey(cell.value) for cell in worksheet[2] if clean(cell.value)]
            recognized = sum(v in {"mandatory","optional","required","reference","recommended"} for v in values)
            metadata_row = bool(values) and recognized >= max(1, len(values)//2)
        if metadata_row:
            worksheet.freeze_panes = "A3"
            for cell in worksheet[2]:
                cell.fill = metadata_fill; cell.font = metadata_font; cell.border = header_border
                cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)

        if worksheet.title.lower() == "summary":
            metric_col = next((c for c in range(1, worksheet.max_column+1) if pkey(worksheet.cell(1,c).value)=="metric"), None)
            if metric_col:
                for row_index in range(2, worksheet.max_row+1):
                    metric = pkey(worksheet.cell(row_index, metric_col).value)
                    if "loadtoisrecords" in metric: worksheet.cell(row_index,metric_col).fill = load_fill
                    elif "retirefromisrecords" in metric: worksheet.cell(row_index,metric_col).fill = retire_fill
                    elif "retireddevicesnoactionrequired" in metric: worksheet.cell(row_index,metric_col).fill = retired_fill

        # Width calculation is intentionally sampled for large enterprise reports.
        sample_end = min(worksheet.max_row, 100)
        for column_index in range(1, worksheet.max_column+1):
            longest = len(clean(worksheet.cell(1,column_index).value))
            for row_index in range(2, sample_end+1):
                value = worksheet.cell(row_index,column_index).value
                if value is not None: longest=max(longest,min(len(str(value)),60))
            worksheet.column_dimensions[get_column_letter(column_index)].width=min(max(longest+2,12),45)
        if worksheet.max_row and worksheet.max_column: worksheet.auto_filter.ref=worksheet.dimensions
        worksheet.page_setup.orientation="landscape";worksheet.page_setup.fitToWidth=1;worksheet.page_setup.fitToHeight=0
        worksheet.sheet_properties.pageSetUpPr.fitToPage=True;worksheet.print_title_rows="1:1"

def _publish_temp_workbook(temp_path: Path, requested_path: Path) -> str:
    if not temp_path.exists() or temp_path.stat().st_size < 1000:
        raise IOError(f"Workbook creation did not produce a valid file: {temp_path}")
    _output_progress("Validating generated Excel workbook", 96)
    check=load_workbook(temp_path,read_only=True,data_only=False);check.close()
    final_path=requested_path
    for _ in range(100):
        if final_path.exists(): final_path=Path(_next_retry_path(str(final_path)));continue
        try:
            os.replace(temp_path,final_path)
            return str(final_path)
        except PermissionError:
            final_path=Path(_next_retry_path(str(final_path)))
    raise PermissionError(f"Unable to publish a versioned workbook in {requested_path.parent}")

def _atomic_write_sheets(requested_path: str, sheets: Dict[str,pd.DataFrame]) -> str:
    requested=Path(requested_path).resolve();requested.parent.mkdir(parents=True,exist_ok=True)
    workdir=requested.parent/".wpp_work";workdir.mkdir(parents=True,exist_ok=True)
    final=requested
    while final.exists():final=Path(_next_retry_path(str(final)))
    temp=workdir/f"{hashlib.sha256(str(final).encode()).hexdigest()[:12]}_{os.getpid()}_{threading.get_ident()}.part.xlsx"
    try:
        count=max(len(sheets),1);_output_progress("Preparing Excel output workbook",72)
        with pd.ExcelWriter(temp,engine="openpyxl") as writer:
            for index,(name,dataframe) in enumerate(sheets.items(),1):
                _output_progress(f"Writing worksheet {index} of {count}: {name}",72+int((index-1)/count*18))
                safe=re.sub(r"[\[\]:*?/\\]","-",str(name))[:31] or "Sheet1";data=dataframe if dataframe is not None and len(dataframe) else pd.DataFrame(columns=["No records"]);data.to_excel(writer,index=False,sheet_name=safe)
            _output_progress("Applying Excel formatting",92);_fast_format_workbook(writer.book)
        if not temp.exists() or temp.stat().st_size<1000:raise IOError(f"Workbook creation did not produce a valid file: {temp}")
        _output_progress("Validating generated Excel workbook",96);check=load_workbook(temp,read_only=True,data_only=False)
        expected={re.sub(r"[\[\]:*?/\\]","-",str(x))[:31] or "Sheet1" for x in sheets}
        if not expected.issubset(set(check.sheetnames)):raise IOError("Generated workbook is missing expected worksheets")
        check.close()
        for _ in range(100):
            if final.exists():final=Path(_next_retry_path(str(final)));continue
            try:os.replace(temp,final);_output_progress(f"Excel output saved: {final.name}",100);return str(final)
            except (PermissionError,FileExistsError):final=Path(_next_retry_path(str(final)))
        raise PermissionError(f"Unable to publish a versioned workbook in {requested.parent}")
    finally:
        try:temp.unlink(missing_ok=True)
        except OSError:pass
        try:
            if workdir.exists() and not any(workdir.iterdir()):workdir.rmdir()
        except OSError:pass


def write_workbook(path: str, sheets: Dict[str,pd.DataFrame]) -> str:
    return _atomic_write_sheets(path,sheets)

def _write_formatted_load(dataframe: pd.DataFrame,path: str) -> str:
    return _atomic_write_sheets(path,{"Load":dataframe})


# -------------------------- persistent resource manager --------------------------
class ResourceRegistry:
    def __init__(self, path=REGISTRY_PATH):
        self.path = Path(path); self.data = self._load()
    def _load(self):
        if not self.path.exists(): return {"resources": {}}
        try:
            data=json.loads(self.path.read_text(encoding="utf-8"))
            return data if isinstance(data,dict) and isinstance(data.get("resources"),dict) else {"resources": {}}
        except Exception: return {"resources": {}}
    def save(self):
        self.path.write_text(json.dumps(self.data, indent=2), encoding="utf-8")
    def register(self, resource_id: str, path: str):
        p=Path(path).resolve()
        if not p.exists(): raise FileNotFoundError(path)
        digest=hashlib.sha256(p.read_bytes()).hexdigest();now=datetime.now().isoformat(timespec="seconds")
        previous=self.data["resources"].get(resource_id,{})
        loaded_on=previous.get("loaded_on",now) if previous.get("sha256")==digest else now
        self.data["resources"][resource_id]={"path":str(p),"file_name":p.name,"loaded_on":loaded_on,"last_used":previous.get("last_used",""),"sha256":digest}
        self.save()
    def get(self, resource_id: str):
        item=self.data["resources"].get(resource_id)
        if item and Path(item.get("path","")).exists(): return item
        return None
    def remove(self,resource_id: str):
        self.data["resources"].pop(resource_id,None);self.save()
    def mark_used(self, resource_id: str):
        item=self.data["resources"].get(resource_id)
        if item:item["last_used"]=datetime.now().isoformat(timespec="seconds");self.save()



# ------------------------- enterprise field intelligence -------------------------
class FieldIntelligence:
    SOURCES = ["IS Operating System Report","NW CMDB Report","Server CMDB Report","IS Server Report","IS Network Report","IS Product Catalog Operating System","IS Product Catalog Network","IS Product Catalog Server"]
    DEFAULT = {
        "fully_qualified_hostname":{"IS Operating System Report":"fully_qualified_hostname","NW CMDB Report":"Fully qualified domain name","Server CMDB Report":"Fully qualified domain name"},
        "hostname":{"IS Operating System Report":"hostname","NW CMDB Report":"Configuration Item","Server CMDB Report":"Name"},
        "domain_name":{"IS Operating System Report":"domain_name","Server CMDB Report":"Domain"},
        "os_lifecycle_status":{"IS Operating System Report":"os_lifecycle_status","NW CMDB Report":"Life Cycle Stage","Server CMDB Report":"Life Cycle Stage"},
        "os_parent_serial_number":{"IS Operating System Report":"os_parent_serial_number","NW CMDB Report":"Serial number","Server CMDB Report":"Serial number"},
        "os_parent_manufacturer":{"IS Operating System Report":"os_parent_manufacturer","NW CMDB Report":"Manufacturer","Server CMDB Report":"Manufacturer"},
        "os_parent_opaque_id":{"IS Operating System Report":"os_parent_opaque_id","IS Server Report":"server_opaque_id","IS Network Report":"network_opaque_id"},
        "operating_system_name":{"IS Operating System Report":"operating_system_name","Server CMDB Report":"Operating System","IS Product Catalog Operating System":"software_product_name"},
        "operating_system_version":{"IS Operating System Report":"operating_system_version","Server CMDB Report":"OS Version","IS Product Catalog Operating System":"software_version"},
        "operating_system_provider":{"IS Operating System Report":"operating_system_provider","IS Product Catalog Operating System":"software_provider"},
        "ip_address":{"IS Operating System Report":"ip_address","NW CMDB Report":"IP Address","Server CMDB Report":"IP Address"},
        "class":{"NW CMDB Report":"Class","Server CMDB Report":"Class"},
        "category":{"NW CMDB Report":"Category","Server CMDB Report":"Category"},
        "subcategory":{"NW CMDB Report":"Subcategory","Server CMDB Report":"Subcategory","IS Product Catalog Operating System":"subcategory"},
    }
    def __init__(self): self.mapping=json.loads(json.dumps(self.DEFAULT)); self.field_lists={}
    def _source_name(self,value):
        return re.sub(r"\s+"," ",clean(value)).strip()
    def load_workbook(self,path: str):
        book=pd.ExcelFile(path,engine="openpyxl")
        sheets={self._source_name(s):s for s in book.sheet_names}
        if "Field List" in sheets:
            raw=pd.read_excel(path,sheet_name=sheets["Field List"],dtype=str,keep_default_na=False,engine="openpyxl")
            self.field_lists={self._source_name(c):[clean(x) for x in raw[c] if clean(x)] for c in raw.columns}
        if "Field Mapping" in sheets:
            raw=pd.read_excel(path,sheet_name=sheets["Field Mapping"],dtype=str,keep_default_na=False,engine="openpyxl")
            original_cols=list(raw.columns);normalized={c:self._source_name(c) for c in original_cols}
            if original_cols:
                logical_source=original_cols[0]
                for _,r in raw.iterrows():
                    logical=clean(r.get(logical_source,""))
                    if not logical: continue
                    self.mapping.setdefault(logical,{})
                    for original in original_cols:
                        value=clean(r.get(original,""))
                        if value:self.mapping[logical][normalized[original]]=value
    def field(self,logical: str,source: str,df: pd.DataFrame,required=True):
        aliases=[]
        mapped=self.mapping.get(logical,{}).get(source)
        if mapped:aliases.append(mapped)
        aliases.extend([logical,logical.replace("_"," ")])
        return find_col(df,*aliases,required=required)




# -------------------- learned OS catalogue identity intelligence --------------------
# Curated mappings supplied from verified NW-CMDB-to-Inventory-OS-catalogue review.
# Exact curated mappings take precedence over generic normalization.
CURATED_FIRMWARE_CATALOG_MAP = {
    pkey("NSX-T 3.2.1"): ("vmware nsx contoroller", "3.2.0.1", "broadcom"),
    pkey("NSX 22.1.5"): ("vmware nsx advanced load balancer", "22.1", "broadcom"),
    pkey("NSX 4.1.0"): ("vmware nsx edge", "4.1.0", "broadcom"),
    pkey("NSX 4.2.1.4"): ("vmware nsx", "4.2.1.4", "broadcom"),
    pkey("MX OS MX 18.211.5"): ("cisco meraki", "mx 18.2.11", "cisco"),
}

def normalize_firmware_catalog_identity(firmware: Any) -> Tuple[str, str, str, str]:
    """Return product, version, provider and method for NW firmware catalogue matching.

    The rules are intentionally conservative. Verified exact mappings are evaluated
    first. Pattern rules then normalize known Cisco, Meraki and Infoblox families.
    Remaining values fall back to the standard product/version parser.
    """
    raw = clean(firmware)
    if not raw:
        return "", "", "", "BLANK FIRMWARE"
    curated = CURATED_FIRMWARE_CATALOG_MAP.get(pkey(raw))
    if curated:
        return curated[0], curated[1], curated[2], "CURATED VERIFIED MAPPING"

    # Meraki firmware families use one catalogue product with a family prefix in version.
    match = re.match(r"^(MR|MX|MS)\s+OS(?:\s+\1)?\s+(.+?)\s*$", raw, re.I)
    if match:
        family = match.group(1).lower()
        return "cisco meraki", f"{family} {clean(match.group(2)).lower()}", "cisco", "MERAKI FAMILY NORMALIZATION"

    # Infoblox catalogue stores the release and omits the post-release build identifier.
    match = re.match(r"^Infoblox\s+NIOS\s+([0-9]+(?:\.[0-9]+)+)(?:-[A-Za-z0-9._-]+)?\s*$", raw, re.I)
    if match:
        return "nios", match.group(1), "infoblox", "INFOBLOX BUILD NORMALIZATION"

    families = [
        (r"^Cisco\s+IOS-XE\s+(.+)$", "cisco ios-xe"),
        (r"^Cisco\s+IOS\s+(.+)$", "cisco ios"),
        (r"^Cisco\s+FXOS\s+(.+)$", "cisco fx-os"),
        (r"^Cisco\s+ASA\s+(.+)$", "cisco asa"),
        (r"^Cisco\s+NXOS\s+(.+)$", "cisco nx-os"),
        (r"^Cisco\s+ISE\s+(.+)$", "cisco ise"),
        (r"^Cisco\s+AireOS\s+(.+)$", "cisco aireos"),
    ]
    for pattern, product in families:
        match = re.match(pattern, raw, re.I)
        if not match:
            continue
        version = clean(match.group(1)).lower()
        # Inventory catalogue ignores the ESW maintenance label in this verified case.
        version = re.sub(r"\s+ESW\d+\s*$", "", version, flags=re.I)
        # 03.11.12.E and similar dotted alpha suffixes are catalogued without the dot.
        version = re.sub(r"\.([a-z])$", r"\1", version, flags=re.I)
        return product, version, "cisco", "NETWORK FAMILY NORMALIZATION"

    product, version = split_product_version(raw)
    return product, version, "", "STANDARD PARSER"


# --------------------------- catalogue/reconciliation ---------------------------
def catalog_index(cat: pd.DataFrame, fi: FieldIntelligence):
    pc=fi.field("operating_system_name","IS Product Catalog Operating System",cat);vc=fi.field("operating_system_version","IS Product Catalog Operating System",cat);pr=fi.field("operating_system_provider","IS Product Catalog Operating System",cat,False)
    exact={};base={}
    for _,r in cat.iterrows():
        p={"name":clean(r[pc]),"version":clean(r[vc]),"provider":clean(r[pr]) if pr else ""}
        exact.setdefault(pkey(p["name"])+"|"+vkey(p["version"]),p)
        base.setdefault(pkey(p["name"])+"|"+re.sub(r"(?<=\d)[a-z]+$","",vkey(p["version"])),p)
    return exact,base

def reconcile_nw(cm,isr,cat,fi,progress=lambda m,p:None):
    progress("Resolving NW fields through Field Intelligence",10)
    ci=fi.field("hostname","NW CMDB Report",cm); cls=fi.field("class","NW CMDB Report",cm); life_c=fi.field("os_lifecycle_status","NW CMDB Report",cm); ip=fi.field("ip_address","NW CMDB Report",cm); fq=fi.field("fully_qualified_hostname","NW CMDB Report",cm,False)
    serial=fi.field("os_parent_serial_number","NW CMDB Report",cm); manuf=fi.field("os_parent_manufacturer","NW CMDB Report",cm,False); firmware=find_col(cm,"Firmware version","Firmware Version")
    def opt(*a):
        c=find_col(cm,*a,required=False);return cm[c] if c else ""
    out=pd.DataFrame({"Configuration Item":cm[ci],"Class":cm[cls],"Life Cycle Stage":cm[life_c],"Life Cycle Stage Status":opt("Life Cycle Stage Status"),"Manufacturer":cm[manuf] if manuf else "","Model ID":opt("Model ID"),"Model.Name":opt("Model.Name","Model Name"),"Model number":opt("Model number"),"Serial number":cm[serial],"IP Address":cm[ip],"Firmware version":cm[firmware],"Fully qualified domain name":cm[fq] if fq else ""})
    parsed=out["Firmware version"].map(normalize_firmware_catalog_identity).apply(pd.Series);parsed.columns=["Parsed Firmware Name","Parsed Firmware Version","Parsed Firmware Provider","Firmware Mapping Method"];out=pd.concat([out,parsed],axis=1)
    exact,base=catalog_index(cat,fi);ih=fi.field("hostname","IS Operating System Report",isr);isn=fi.field("os_parent_serial_number","IS Operating System Report",isr);il=fi.field("os_lifecycle_status","IS Operating System Report",isr);io=find_col(isr,"os_opaque_id","OS Opaque ID",required=False);ion=fi.field("operating_system_name","IS Operating System Report",isr);iov=fi.field("operating_system_version","IS Operating System Report",isr);iop=fi.field("operating_system_provider","IS Operating System Report",isr,False)
    hm={};sm={}
    for i,r in isr.iterrows():
        if clean(r[ih]):hm.setdefault(clean(r[ih]).lower(),i)
        if serial_key(r[isn]):sm.setdefault(serial_key(r[isn]),i)
    rows=[];total=max(len(out),1);progress("Reconciling NW CMDB and IS",25)
    for pos,(_,r) in enumerate(out.iterrows(),1):
        if pos==1 or pos==total or pos%500==0:progress(f"Reconciling NW records: {pos} of {len(out)}",25+int(pos/total*40))
        name=clean(r["Configuration Item"]);idx=next((hm[x] for x in [name.lower(),name.replace("/","_").lower(),name.replace("/","-").lower()] if x in hm),None);mode="Hostname Match" if idx is not None else ""
        if idx is None and serial_key(r["Serial number"]) in sm:idx=sm[serial_key(r["Serial number"])];mode="Serial Match Different Hostname"
        present=idx is not None;ir=isr.iloc[idx] if present else None;k=pkey(r["Parsed Firmware Name"])+"|"+vkey(r["Parsed Firmware Version"]);hit=exact.get(k);status="Match"
        if not hit:hit=base.get(pkey(r["Parsed Firmware Name"])+"|"+re.sub(r"(?<=\d)[a-z]+$","",vkey(r["Parsed Firmware Version"])));status="Closest Match Proposed" if hit else "No Match"
        hit=hit or {};osmatch=bool(present and hit and pkey(ir[ion])==pkey(hit.get("name")) and vkey(ir[iov])==vkey(hit.get("version")))
        life=lifecycle(r["Life Cycle Stage"]);actions=[];classification="ACTIVE OR REVIEW";reason=""
        if life=="end of life" and not present:
            actions=["Retired Device - No Action Required"];classification="RETIRED DEVICE - NOT PRESENT IN IS";reason="CMDB DEVICE IS END OF LIFE AND NO ACTIVE IS OS RECORD EXISTS"
        else:
            if status=="No Match":actions.append("Catalogue update required")
            if present and hit and not osmatch:actions.append("Inventory OS update required")
            if life=="operational" and not present:actions.append("Load To IS")
            if life=="end of life" and present and lifecycle(ir[il]) in {"production","installed"}:actions.append("Retire From IS")
        d=r.to_dict();d.update({"Present in IS ?":"Yes - "+mode if present else "No","Host name in IS":clean(ir[ih]) if present else "","os_opaque_id":clean(ir[io]) if present and io else "","os_parent_serial_number":clean(ir[isn]) if present else "","Lifecycle Status in IS":clean(ir[il]) if present else "","Current operating_system_provider":clean(ir[iop]) if present and iop else "","Current operating_system_name":clean(ir[ion]) if present else "","Current operating_system_version":clean(ir[iov]) if present else "","Catalogue Match?":status,"Required operating_system_provider":hit.get("provider",""),"Required operating_system_name":hit.get("name",""),"Required operating_system_version":hit.get("version",""),"Device Classification":classification,"Action Reason":reason,"Action":"; ".join(actions) or "No Action"});rows.append(d)
    progress("NW reconciliation complete",68);return pd.DataFrame(rows)


def nw_summary(r):
    return pd.DataFrame([("NW CMDB records analyzed",len(r)),("Records present in IS",int(r["Present in IS ?"].astype(str).str.startswith("Yes").sum())),("Records not present in IS",int((r["Present in IS ?"]=="No").sum())),("Catalogue exact matches",int((r["Catalogue Match?"]=="Match").sum())),("Catalogue closest matches proposed",int((r["Catalogue Match?"]=="Closest Match Proposed").sum())),("Catalogue updates required",int(r["Action"].str.contains("Catalogue update required",na=False).sum())),("Inventory OS updates required",int(r["Action"].str.contains("Inventory OS update required",na=False).sum())),("Load To IS records",int(r["Action"].str.contains("Load To IS",na=False).sum())),("Retire From IS records",int(r["Action"].str.contains("Retire From IS",na=False).sum())),("Retired devices - no action required",int(r["Action"].eq("Retired Device - No Action Required").sum()))],columns=["Metric","Value"])


def reconcile_server(cm,isr,cat,fi,progress=lambda m,p:None):
    progress("Resolving Server fields through Field Intelligence",10);support=find_col(cm,"Support Group",required=False);managed=find_col(cm,"Managed by Group","Managed By Group",required=False)
    if not support and not managed:raise ValueError("Server CMDB requires Support Group or Managed By Group for W-KYN scoping")
    scope=pd.Series(False,index=cm.index)
    if support:scope|=cm[support].map(clean).str.upper().str.startswith("W-KYN")
    if managed:scope|=cm[managed].map(clean).str.upper().str.startswith("W-KYN")
    scoped=cm[scope].copy().reset_index(drop=True);outside=cm[~scope].copy().reset_index(drop=True)
    if scoped.empty:return pd.DataFrame(),outside,pd.DataFrame(),pd.DataFrame()
    name=fi.field("hostname","Server CMDB Report",scoped);life=fi.field("os_lifecycle_status","Server CMDB Report",scoped);ip=fi.field("ip_address","Server CMDB Report",scoped);serial=fi.field("os_parent_serial_number","Server CMDB Report",scoped,False);osn=fi.field("operating_system_name","Server CMDB Report",scoped);osv=fi.field("operating_system_version","Server CMDB Report",scoped);fq=fi.field("fully_qualified_hostname","Server CMDB Report",scoped,False);cls=fi.field("class","Server CMDB Report",scoped,False)
    def col(*a):c=find_col(scoped,*a,required=False);return scoped[c] if c else ""
    base=pd.DataFrame({"Configuration Item":scoped[name],"Class":scoped[cls] if cls else "Server","Life Cycle Stage":scoped[life],"Life Cycle Stage Status":col("Life Cycle Stage Status"),"Manufacturer":col("Manufacturer"),"Model ID":col("Model ID"),"Model.Name":"","Model number":"","Serial number":scoped[serial] if serial else "","IP Address":scoped[ip],"Firmware version":scoped[osn].map(clean)+" "+scoped[osv].map(clean),"Fully qualified domain name":scoped[fq] if fq else scoped[name],"category":col("Category"),"subcategory":col("Subcategory"),"Attribute Field Value":col("Attribute Field Value","Attributes")})
    base["_name_key"]=base["Configuration Item"].map(pkey);base["_life_key"]=base["Life Cycle Stage"].map(lifecycle);duplicates=base[base.duplicated(["_name_key","_life_key"],False)].copy();base=base.drop_duplicates(["_name_key","_life_key"],keep="first");active=set(base.loc[base["_life_key"]=="operational","_name_key"]);historical=base[(base["_life_key"]=="end of life")&base["_name_key"].isin(active)].copy();eligible=base[~((base["_life_key"]=="end of life")&base["_name_key"].isin(active))].copy();context=eligible[["Configuration Item","category","subcategory","Attribute Field Value"]].reset_index(drop=True);standardized=eligible.drop(columns=["_name_key","_life_key","category","subcategory","Attribute Field Value"])
    def sp(m,p):progress({"Resolving NW fields through Field Intelligence":"Preparing Server reconciliation fields","Reconciling NW CMDB and IS":"Reconciling Server CMDB and IS","NW reconciliation complete":"Server reconciliation complete"}.get(m,m.replace("NW","Server")),p)
    result=reconcile_nw(standardized,isr,cat,fi,sp)
    for c in context.columns[1:]:result[c]=context[c].values
    progress("Server reconciliation complete",68);return result,outside,duplicates,historical



def server_summary(r,outside,duplicates,historical):
    vals=[("Server CMDB records analyzed",len(r)),("Records present in IS",int(r["Present in IS ?"].astype(str).str.startswith("Yes").sum()) if not r.empty else 0),("Records not present in IS",int((r["Present in IS ?"]=="No").sum()) if not r.empty else 0),("Load To IS records",int(r["Action"].str.contains("Load To IS",na=False).sum()) if not r.empty else 0),("Retire From IS records",int(r["Action"].str.contains("Retire From IS",na=False).sum()) if not r.empty else 0),("Retired devices - no action required",int(r["Action"].eq("Retired Device - No Action Required").sum()) if not r.empty else 0),("Out of W-KYN scope",len(outside)),("CMDB duplicate records",len(duplicates)),("Historical EOL suppressed",len(historical))];return pd.DataFrame(vals,columns=["Metric","Value"])


# ------------------------------- policy intelligence -------------------------------
def windows_supported(name,version):
    text=(clean(name)+" "+clean(version)).lower()
    for label,(build,end) in WINDOWS_SERVER_SUPPORT.items():
        if label in text or (build and build in text): return date.today() <= end,label,end
    return False,"Unknown",None

def is_server_context(context):
    text=" ".join(clean(context.get(k,"")) for k in ["source_type","subcategory","parent_system_type","image_purpose","os_name","class","category"]).lower()
    return any(x in text for x in ["server","windows","linux","rhel","ubuntu","aix","solaris","esx","hypervisor"])

def apply_policy(row: Dict[str,str],context: Dict[str,str]):
    applied=[]
    def put(field,value,rule): row[field]=str(value);applied.append({"Field":field,"Value":str(value),"Rule":rule})
    put("is_disaster_recovery_contracted","N","P001 Global")
    put("is_designated_for_compliance_reporting","Y","P002 Global")
    put("bac_id","WPP Europe","P003 Global")
    put("bam_id","BAM000415E","P004 Global")
    sub=clean(context.get("subcategory","")).lower();parent=clean(context.get("parent_system_type","")).lower();osname=clean(context.get("os_name","")).lower();source=clean(context.get("source_type","")).lower()
    supported,release,end=windows_supported(context.get("os_name",""),context.get("os_version",""))
    if source=="network" or parent in {"network","appliance"} or "esx" in osname:
        put("is_anti_virus_required","N","AV002/AV003/AV004")
    elif sub=="windows" and supported:
        put("is_anti_virus_required","Y","AV001 Supported Windows")
        put("is_patching_required","Y","PATCH001 Supported Windows")
        put("is_health_check_required","Y","HC001 Supported Windows")
    if row.get("is_health_check_required")=="Y":
        put("health_check_interval","12","HC002")
        put("health_check_unit","month","HC002")
    if is_server_context(context):
        put("is_employment_verification_required","Y","QEV001")
        put("employment_verification_interval","3","QEV001")
        put("employment_verification_unit","month","QEV001")
        put("is_continued_business_need_required","Y","CBN001")
        sox="sox" in clean(context.get("attribute_field_value","")).lower()
        put("continued_business_need_interval","3" if sox else "12","CBN002 SOX" if sox else "CBN001 Annual")
        put("continued_business_need_unit","month","CBN001/CBN002")
    return applied

# ----------------------------- learned flag intelligence -----------------------------
class FlagLearning:
    COHORT_FIELDS=["operating_system_subcategory","parent_system_type","image_purpose","os_lifecycle_status","operating_system_name","business_criticality_device_grouping","os_inventory_item_owned_by","os_inventory_item_managed_by"]
    PLACEHOLDERS={"jane doe","yyy@kyndryl.com","testgroup","testsupportgroup","analisy ongoing","analysis ongoing"}
    def __init__(self,path=LEARNING_DB):
        self.conn=sqlite3.connect(path); self.conn.executescript("CREATE TABLE IF NOT EXISTS reports(hash TEXT PRIMARY KEY,loaded TEXT,name TEXT);CREATE TABLE IF NOT EXISTS observations(field TEXT,value TEXT,priority TEXT,cohort TEXT);CREATE TABLE IF NOT EXISTS anomalies(report_hash TEXT,os_opaque_id TEXT,hostname TEXT,field TEXT,value TEXT,type TEXT);"); self.conn.commit()
    def ingest(self,path):
        digest=hashlib.sha256(Path(path).read_bytes()).hexdigest()
        if self.conn.execute("SELECT 1 FROM reports WHERE hash=?",(digest,)).fetchone(): return digest
        df,meta=read_tabular(path,True); meta=meta if meta is not None else pd.Series("",index=df.columns); oid=find_col(df,"os_opaque_id","OS Opaque ID",required=False); host=find_col(df,"hostname","host_name",required=False)
        duplicates=set(df.loc[df[oid].duplicated(False),oid])-{""} if oid else set(); flags=[]
        for c in df.columns:
            values={clean(x).upper() for x in df[c] if clean(x)}
            if values and values.issubset({"Y","N"}): flags.append((c,"Priority 1" if clean(meta.get(c,"" )).lower()=="mandatory" else "Priority 2"))
        for _,r in df.iterrows():
            rid=clean(r.get(oid,"")) if oid else ""; hostname=clean(r.get(host,"")) if host else ""
            if not rid:self.conn.execute("INSERT INTO anomalies VALUES(?,?,?,?,?,?)",(digest,rid,hostname,"os_opaque_id","","Missing unique identifier"))
            if rid in duplicates:self.conn.execute("INSERT INTO anomalies VALUES(?,?,?,?,?,?)",(digest,rid,hostname,"os_opaque_id",rid,"Duplicate os_opaque_id"))
            cohort="|".join(pkey(r.get(x,"")) for x in self.COHORT_FIELDS)
            for field,priority in flags:
                raw=clean(r[field]); value=raw.upper()
                if raw.lower() in self.PLACEHOLDERS:self.conn.execute("INSERT INTO anomalies VALUES(?,?,?,?,?,?)",(digest,rid,hostname,field,raw,"Placeholder value"));continue
                if value in {"Y","N"}:self.conn.execute("INSERT INTO observations VALUES(?,?,?,?)",(hkey(field),value,priority,cohort))
        self.conn.execute("INSERT INTO reports VALUES(?,?,?)",(digest,datetime.now().isoformat(timespec="seconds"),Path(path).name));self.conn.commit();return digest
    def predict(self,field,priority,cohort):
        parts=cohort.split("|")
        for n in range(len(parts),0,-1):
            prefix="|".join(parts[:n]); vals=[x[0] for x in self.conn.execute("SELECT value FROM observations WHERE field=? AND priority=? AND cohort LIKE ?",(hkey(field),priority,prefix+"%")).fetchall()]
            if vals:
                y=vals.count("Y");z=vals.count("N");support=y+z;agreement=max(y,z)/support
                return (("Y" if y>=z else "N"),agreement,support,"Auto-prefill",n) if agreement>=.75 else ("",agreement,support,"No proposal",n)
        return "",0,0,"No evidence",0
    def anomalies(self,digest): return pd.read_sql_query("SELECT os_opaque_id,hostname,field,value,type FROM anomalies WHERE report_hash=?",self.conn,params=(digest,))
    def close(self):self.conn.close()


def cleanup_workdir(folder: str) -> None:
    work=Path(folder)/".wpp_work"
    if not work.exists():return
    for p in work.glob("*.part.xlsx"):
        try:p.unlink()
        except OSError:pass
    try:
        if not any(work.iterdir()):work.rmdir()
    except OSError:pass

def candidate_set_id(domain: str, row) -> str:
    raw="|".join([domain,serial_key(row.get("Serial number","")),pkey(row.get("Configuration Item","")),lifecycle(row.get("Life Cycle Stage","")),clean(row.get("Action",""))])
    return hashlib.sha256(raw.encode("utf-8")).hexdigest()[:16]

def safe_template(path: str):
    import warnings
    with warnings.catch_warnings():
        warnings.filterwarnings("ignore",message="Workbook contains no default style, apply openpyxl's default",category=UserWarning,module="openpyxl.styles.stylesheet")
        return read_tabular(path,True)

def read_load_to_is_file(path: str, domain: str, progress=lambda m,p:None) -> pd.DataFrame:
    label="NW" if domain=="network" else "Server";progress(f"Reading {label} Load To IS input",5);ext=Path(path).suffix.lower()
    import warnings
    with warnings.catch_warnings():
        warnings.filterwarnings("ignore",message="Workbook contains no default style, apply openpyxl's default",category=UserWarning,module="openpyxl.styles.stylesheet")
        if ext in {".xlsx",".xlsm"}:
            book=pd.ExcelFile(path,engine="openpyxl");sheet=next((s for s in book.sheet_names if hkey(s)==hkey("Load To IS")),book.sheet_names[0]);df=pd.read_excel(path,sheet_name=sheet,dtype=str,keep_default_na=False,engine="openpyxl")
        elif ext==".xlsb":
            book=pd.ExcelFile(path,engine="pyxlsb");sheet=next((s for s in book.sheet_names if hkey(s)==hkey("Load To IS")),book.sheet_names[0]);df=pd.read_excel(path,sheet_name=sheet,dtype=str,keep_default_na=False,engine="pyxlsb")
        elif ext==".csv":df=pd.read_csv(path,dtype=str,keep_default_na=False,encoding_errors="replace")
        else:raise ValueError("Load To IS input must be CSV, XLSX, XLSM, or XLSB")
    df.columns=[clean(c) for c in df.columns]
    required=["Configuration Item","Life Cycle Stage","Serial number","Present in IS ?","Action","Required operating_system_provider","Required operating_system_name","Required operating_system_version"]
    missing=[c for c in required if find_col(df,c,required=False) is None]
    if missing:raise ValueError("Invalid Load To IS input. Missing columns: "+", ".join(missing))
    action=find_col(df,"Action");life=find_col(df,"Life Cycle Stage");present=find_col(df,"Present in IS ?");serial=find_col(df,"Serial number")
    mask=df[action].map(clean).str.contains("Load To IS",case=False,na=False)&df[life].map(lifecycle).eq("operational")&df[present].map(clean).str.lower().eq("no")&df[serial].map(clean).ne("")
    valid=df[mask].copy().reset_index(drop=True);excluded=int((~mask).sum());valid["Candidate Set ID"]=valid.apply(lambda r:candidate_set_id(domain,r),axis=1)
    progress(f"Validated {label} Load To IS candidates: {len(valid)}; excluded: {excluded}",15)
    if valid.empty:raise ValueError("The selected file contains no valid Operational Load To IS candidates")
    return valid

def category_decisions_from_load(domain: str, load_df: pd.DataFrame, category_path: str, progress=lambda m,p:None):
    label="NW" if domain=="network" else "Server";progress(f"Reading current {label} Category report",18);cat=read_tabular(category_path)
    cserial=find_col(cat,"network_serial_number" if domain=="network" else "serial_number","server_serial_number" if domain=="server" else "network_serial_number");opaque=find_col(cat,"network_opaque_id" if domain=="network" else "server_opaque_id",required=False);clife=find_col(cat,"network_lifecycle_status" if domain=="network" else "server_lifecycle_status","lifecycle_status",required=False)
    cmap=defaultdict(list)
    for i,r in cat.iterrows():
        if serial_key(r[cserial]) and (not clife or lifecycle(r[clife]) in {"production","installed","operational"}):cmap[serial_key(r[cserial])].append(i)
    decisions=[];total=max(len(load_df),1)
    for pos,(_,r) in enumerate(load_df.iterrows(),1):
        if pos==1 or pos==total or pos%500==0:progress(f"Resolving {label} parent relationships: {pos} of {len(load_df)}",20+int(pos/total*35))
        ids=cmap.get(serial_key(r["Serial number"]),[]);d={"CMDB Hostname":clean(r["Configuration Item"]),"CMDB Serial Number":clean(r["Serial number"]),"Candidate Set ID":r["Candidate Set ID"],"Reconciliation Action":r["Action"],"Catalog Match Status":clean(r.get("Catalogue Match?",""))}
        if len(ids)==1:
            cr=cat.iloc[ids[0]];d.update({"Recommended Action":"LOAD OS ONLY","Parent Dependency Status":"EXISTING PRODUCTION CATEGORY","Category Presence Status":"FOUND IN CATEGORY","os_parent_opaque_id":clean(cr[opaque]) if opaque else "","os_parent_serial_number":clean(cr[cserial])})
            for target,aliases in {"os_parent_manufacturer":["network_manufacturer","server_manufacturer","manufacturer"],"os_parent_hardware_type":["network_hardware_type","server_hardware_type","hardware_type"],"os_parent_hardware_model":["network_hardware_model","server_hardware_model","hardware_model"],"os_parent_category":["category","network_subcategory","server_subcategory","subcategory"]}.items():
                c=find_col(cat,*aliases,required=False);d[target]=clean(cr[c]) if c else ""
        elif not ids:d.update({"Recommended Action":"LOAD CATEGORY AND OS","Parent Dependency Status":"PARENT CATEGORY LOAD REQUIRED","Category Presence Status":"MISSING FROM CATEGORY","os_parent_opaque_id":"","os_parent_serial_number":clean(r["Serial number"])})
        else:d.update({"Recommended Action":"CATEGORY DATA QUALITY REVIEW","Parent Dependency Status":"BLOCKED - DUPLICATE CATEGORY SERIAL","Category Presence Status":"DUPLICATE NORMALIZED SERIAL"})
        decisions.append(d)
    return {"outputs":[],"decisions":pd.DataFrame(decisions),"category_load":pd.DataFrame(),"candidate_ids":set(load_df["Candidate Set ID"])}

# -------------------------------- bulk load engine --------------------------------
def generate_bulk_load(rec,source_type,is_report,template,outdir,fmt,askfq,askcat,progress=lambda m,p:None,governance=None):
    label="NW" if source_type=="network" else "Server";progress(f"Starting {label} OS Bulk Load workflow",2);candidates=rec[(rec["Present in IS ?"]=="No")&(rec["Life Cycle Stage"].map(lifecycle)=="operational")&rec["Action"].str.contains("Load To IS",na=False)].copy();candidates["Candidate Set ID"]=candidates.apply(lambda r:candidate_set_id(source_type,r),axis=1);progress(f"Using {label} reconciliation Load To IS population: {len(candidates)}",8)
    decisions=governance["decisions"] if governance else pd.DataFrame();decmap={x["Candidate Set ID"]:x for _,x in decisions.iterrows()};consistency=[];eligible=[]
    for _,r in candidates.iterrows():
        cid=r["Candidate Set ID"];d=decmap.get(cid);ok=d is not None and d["Recommended Action"] in {"LOAD OS ONLY","LOAD CATEGORY AND OS"};consistency.append({"Domain":label,"Candidate Set ID":cid,"CMDB Hostname":r["Configuration Item"],"CMDB Serial Number":r["Serial number"],"In Reconciliation Load To IS":"Y","In Hardware Governance":"Y" if d is not None else "N","In Category Load":"Y" if d is not None and d["Recommended Action"]=="LOAD CATEGORY AND OS" else "N","In OS Bulk Load":"Y" if ok else "N","Consistency Status":"PASS" if ok else "BLOCKED","Blocking Reason":"" if ok else "MISSING OR BLOCKED HARDWARE/CATEGORY DECISION"})
        if ok:eligible.append((r,d))
    progress("Candidate consistency validation complete",15);template_df,meta=safe_template(template);headers=list(template_df.columns);extras=["CMDB Hostname","CMDB Serial Number","CMDB Lifecycle Stage","Reconciliation Action","Candidate Set ID","Category Load Required","Parent Dependency Status","Hardware Catalog Match Status"];rows=[];ip_audit=[];total=max(len(eligible),1)
    for pos,(r,d) in enumerate(eligible,1):
        if pos==1 or pos==total or pos%250==0:progress(f"Generating {label} OS load rows: {pos} of {len(eligible)}",20+int(pos/total*45))
        fq=valid_fqdn(r["Fully qualified domain name"])
        if not fq and askfq(1):fq=normalize_hostname(r["Configuration Item"]).lower()+".nw.wpp.net"
        if not fq:continue
        ip,msg,more=normalize_ipv4(r["IP Address"]);values={"hostname":normalize_hostname(r["Configuration Item"]),"fully_qualified_hostname":fq,"os_parent_serial_number":d.get("os_parent_serial_number",r["Serial number"]),"os_parent_opaque_id":d.get("os_parent_opaque_id",""),"os_parent_manufacturer":d.get("os_parent_manufacturer",""),"os_parent_hardware_type":d.get("os_parent_hardware_type",""),"os_parent_hardware_model":d.get("os_parent_hardware_model",""),"os_parent_category":d.get("os_parent_category",""),"os_lifecycle_status":"Production","parent_system_type":"Network" if source_type=="network" else "OS","operating_system_provider":r["Required operating_system_provider"],"operating_system_name":r["Required operating_system_name"],"operating_system_version":r["Required operating_system_version"],"ip_address":ip};row={h:"" for h in headers};lookup={hkey(k):v for k,v in values.items()}
        for h in headers:
            if hkey(h) in lookup:row[h]=lookup[hkey(h)]
        refs=[r["Configuration Item"],r["Serial number"],r["Life Cycle Stage"],r["Action"],r["Candidate Set ID"],"Y" if d["Recommended Action"]=="LOAD CATEGORY AND OS" else "N",d["Parent Dependency Status"],d["Catalog Match Status"]];rows.append(dict(zip(extras,refs),**row));ip_audit.append({"CMDB Hostname":r["Configuration Item"],"CMDB Serial Number":r["Serial number"],"Original IP":r["IP Address"],"Selected IP":ip,"Additional IPs":more,"Message":msg})
    progress("Loading Flag Intelligence and anomaly evidence",68)
    learner=FlagLearning();report_hash=learner.ingest(is_report);learning_anomalies=learner.anomalies(report_hash);learner.close()
    metadata=["Reference"]*len(extras)+([clean(meta[h]) for h in headers] if meta is not None else [""]*len(headers));final=pd.concat([pd.DataFrame([metadata],columns=extras+headers),pd.DataFrame(rows,columns=extras+headers)],ignore_index=True);set_output_progress(progress);load_path=next_versioned_path(outdir,label+" Enriched OS Bulk Load",fmt)
    load_path=_write_formatted_load(final,load_path) if fmt=="xlsx" else str(Path(load_path));
    if fmt!="xlsx":final.to_csv(load_path,index=False)
    validation=write_workbook(next_versioned_path(outdir,label+" Enriched OS Bulk Load Validation","xlsx"),{"Workflow Summary":pd.DataFrame([("Reconciliation Load To IS candidates",len(candidates)),("Category Load required",sum(1 for _,d in eligible if d["Recommended Action"]=="LOAD CATEGORY AND OS")),("OS load rows generated",len(rows))],columns=["Metric","Value"]),"IP Validation":pd.DataFrame(ip_audit),"Parent Dependency Validation":decisions,"Candidate Consistency":pd.DataFrame(consistency),"IS Learning Anomalies":learning_anomalies})
    progress(f"{label} OS Bulk Load workflow complete",100);return [load_path,validation]


@dataclass(frozen=True)
class HardwareMatch:
    record: Optional[dict]; status: str; detail: str; score: int=0

def _hw_index(cat):
    cols={x:find_col(cat,x) for x in ["manufacturer","hardware_type","hardware_model","subcategory"]}; eol=find_col(cat,"end_of_life_date",required=False);eosl=find_col(cat,"end_of_support_date",required=False);oid=find_col(cat,"product_opaque_id",required=False)
    records=[];exact=defaultdict(list);composite=defaultdict(list);mm=defaultdict(list)
    for _,r in cat.iterrows():
        d={"manufacturer":clean(r[cols["manufacturer"]]),"hardware_type":clean(r[cols["hardware_type"]]),"hardware_model":clean(r[cols["hardware_model"]]),"subcategory":clean(r[cols["subcategory"]]),"eol":clean(r[eol]) if eol else "","eosl":clean(r[eosl]) if eosl else "","opaque_id":clean(r[oid]) if oid else ""};records.append(d);i=len(records)-1;exact[pkey(d["hardware_model"])].append(i);composite[pkey(d["hardware_type"]+d["hardware_model"])].append(i);mm[pkey(d["manufacturer"])+"|"+pkey(d["hardware_model"])].append(i)
    return records,exact,composite,mm

def hardware_match(domain,manufacturer,model,idx):
    records,exact,composite,mm=idx;q=pkey(model);mk=pkey(manufacturer)
    ids=mm.get(mk+"|"+q,[]) if domain=="server" else []
    if ids:return HardwareMatch(records[ids[0]],"EXACT MATCH","EXACT MANUFACTURER + MODEL ID",10000)
    ids=exact.get(q,[])
    if len(ids)==1:
        r=records[ids[0]]
        if domain=="server" and mk and pkey(r["manufacturer"])!=mk:return HardwareMatch(r,"PARTIAL MATCH","MODEL MATCH - MANUFACTURER REVIEW",9000)
        return HardwareMatch(r,"EXACT MATCH","EXACT HARDWARE MODEL",9800)
    ids=composite.get(q,[])
    if len(ids)==1:
        r=records[ids[0]]
        if domain=="server" and mk and pkey(r["manufacturer"])!=mk:return HardwareMatch(r,"PARTIAL MATCH","MODEL MATCH - MANUFACTURER REVIEW",9500)
        return HardwareMatch(r,"EXACT MATCH","EXACT TYPE + MODEL COMPOSITE",9500)
    scored=[]
    tokens=set(re.findall(r"[a-z]+|\d+",clean(model).lower()))
    for i,r in enumerate(records):
        rt=set(re.findall(r"[a-z]+|\d+",(r["hardware_type"]+" "+r["hardware_model"]).lower()));score=len(tokens&rt)*100
        if score:scored.append((score,i))
    if not scored:return HardwareMatch(None,"NO MATCH","NO INDEXED CANDIDATE",0)
    top=max(x[0] for x in scored);ids=[i for s,i in scored if s==top]
    if len({(pkey(records[i]["manufacturer"]),pkey(records[i]["hardware_type"]),pkey(records[i]["hardware_model"])) for i in ids})>1:return HardwareMatch(None,"NO MATCH","AMBIGUOUS PARTIAL MATCH",top)
    r=records[ids[0]];detail="MODEL MATCH - MANUFACTURER REVIEW" if domain=="server" and mk and pkey(r["manufacturer"])!=mk else "INDEXED HARDWARE TYPE/TOKEN MATCH";return HardwareMatch(r,"PARTIAL MATCH",detail,top)

def run_hardware_governance(domain,cmdb_path,category_path,catalog_path,outdir,progress=lambda m,p:None,os_reconciliation=None,is_os_path=None,write_outputs=True):
    label="NW" if domain=="network" else "Server";progress(f"Loading {label} hardware governance inputs",5);cm=read_tabular(cmdb_path);catdf=read_tabular(category_path);catalog=read_tabular(catalog_path);idx=_hw_index(catalog)
    host=find_col(cm,"Configuration Item" if domain=="network" else "Name","CI.Name","Fully qualified domain name");serial=find_col(cm,"Serial number");life=find_col(cm,"Life Cycle Stage");man=find_col(cm,"Manufacturer");model=find_col(cm,"Model number" if domain=="network" else "Model ID");mid=find_col(cm,"Model ID",required=False);lstat=find_col(cm,"Life Cycle Stage Status",required=False)
    cserial=find_col(catdf,"network_serial_number" if domain=="network" else "serial_number","server_serial_number" if domain=="server" else "network_serial_number");opaque=find_col(catdf,"network_opaque_id" if domain=="network" else "server_opaque_id",required=False);clife=find_col(catdf,"network_lifecycle_status" if domain=="network" else "server_lifecycle_status","lifecycle_status",required=False)
    cmap=defaultdict(list)
    for i,r in catdf.iterrows():
        if serial_key(r[cserial]):cmap[serial_key(r[cserial])].append(i)
    recon=os_reconciliation if os_reconciliation is not None else pd.DataFrame();loads=recon[recon["Action"].str.contains("Load To IS",na=False)].copy() if not recon.empty else pd.DataFrame();loadkeys={serial_key(x) for x in loads["Serial number"]};decisions=[];loadrows=[];matches=[];recs=[]
    total=max(len(cm),1)
    for pos,(_,r) in enumerate(cm.iterrows(),1):
        if pos==1 or pos==total or pos%500==0:progress(f"Evaluating {label} hardware records: {pos} of {len(cm)}",10+int(pos/total*40))
        sk=serial_key(r[serial]);rr=next((x for _,x in loads.iterrows() if serial_key(x["Serial number"])==sk),None);hm=hardware_match(domain,r[man],r[model],idx);ids=cmap.get(sk,[]);presence="FOUND IN CATEGORY" if len(ids)==1 else "FOUND - DUPLICATE NORMALIZED SERIAL" if len(ids)>1 else "MISSING FROM CATEGORY";recommended="NO LOAD ACTION";parent={}
        if rr is not None and lifecycle(r[life])=="operational":
            if len(ids)==1:
                cr=catdf.iloc[ids[0]];recommended="LOAD OS ONLY";parent={"os_parent_opaque_id":clean(cr[opaque]) if opaque else "","os_parent_serial_number":clean(cr[cserial])}
                for target,aliases in {"os_parent_manufacturer":["network_manufacturer","server_manufacturer","manufacturer"],"os_parent_hardware_type":["network_hardware_type","server_hardware_type","hardware_type"],"os_parent_hardware_model":["network_hardware_model","server_hardware_model","hardware_model"],"os_parent_category":["category","network_subcategory","server_subcategory","subcategory"]}.items():
                    c=find_col(catdf,*aliases,required=False);parent[target]=clean(cr[c]) if c else ""
            elif not ids:recommended="LOAD CATEGORY AND OS";parent={"os_parent_opaque_id":"","os_parent_serial_number":clean(r[serial])}
            else:recommended="CATEGORY DATA QUALITY REVIEW"
        cid=candidate_set_id(domain,rr) if rr is not None else "";d={"CMDB Hostname":clean(r[host]),"CMDB Serial Number":clean(r[serial]),"Normalized Serial Number":sk,"CMDB Manufacturer":clean(r[man]),"CMDB Model":clean(r[model]),"CMDB Lifecycle Stage":clean(r[life]),"Reconciliation Action":clean(rr["Action"]) if rr is not None else "","Candidate Set ID":cid,"Catalog Match Status":hm.status,"Catalog Match Detail":hm.detail,"Match Score":hm.score,"Category Presence Status":presence,"Selected Category Opaque ID":parent.get("os_parent_opaque_id",""),"Recommended Action":recommended,"Parent Dependency Status":"EXISTING PRODUCTION CATEGORY" if recommended=="LOAD OS ONLY" else "PARENT CATEGORY LOAD REQUIRED" if recommended=="LOAD CATEGORY AND OS" else "NOT APPLICABLE"};d.update(parent);decisions.append(d);matches.append(d)
        if recommended=="LOAD CATEGORY AND OS":
            vals={c:"" for c in catdf.columns};vals[cserial]=clean(r[serial]);pref="network_" if domain=="network" else "server_"
            if pref+"lifecycle_status" in vals:vals[pref+"lifecycle_status"]="Production"
            if hm.status=="EXACT MATCH" and hm.record:
                for f,k in [("manufacturer","manufacturer"),("hardware_type","hardware_type"),("hardware_model","hardware_model"),("subcategory","subcategory")]:
                    if pref+f in vals:vals[pref+f]=hm.record[k]
            ref=[clean(r[host]),clean(r[serial]),clean(r[mid]) if mid else ""]+([clean(r[model])] if domain=="network" else [])+[clean(r[life]),clean(r[lstat]) if lstat else "",clean(rr["Action"]),cid,hm.status,"PARENT CATEGORY LOAD REQUIRED"];loadrows.append(ref+[vals[c] for c in catdf.columns])
        if rr is not None and hm.status!="EXACT MATCH":recs.append(d)
    decisions_df=pd.DataFrame(decisions);prefix=["CMDB Hostname","CMDB Serial Number","CMDB Model ID"]+(["CMDB Model Number"] if domain=="network" else [])+["CMDB Lifecycle Stage","CMDB Lifecycle Stage Status","Reconciliation Action","Candidate Set ID","Catalog Match Status","Category Dependency Status"]
    outputs=[]
    if write_outputs:
        set_output_progress(progress);p=write_workbook(next_versioned_path(outdir,label+" Hardware Catalog Match","xlsx"),{"Catalog Match":pd.DataFrame(matches)});outputs.append(p)
        metadata=["Reference"]*len(prefix)+[""]*len(catdf.columns);p=_write_formatted_load(pd.DataFrame([metadata]+loadrows,columns=prefix+list(catdf.columns)),next_versioned_path(outdir,label+" Category Load","xlsx"));outputs.append(p)
        p=write_workbook(next_versioned_path(outdir,label+" Catalog Enhancement Recommendations","xlsx"),{"Recommendations":pd.DataFrame(recs)});outputs.append(p)
    progress(f"{label} hardware governance complete",60);return {"outputs":outputs,"decisions":decisions_df,"category_load":pd.DataFrame(loadrows,columns=prefix+list(catdf.columns)),"candidate_ids":set(decisions_df.get("Candidate Set ID",pd.Series(dtype=str)))}


class NetworkHardwareGovernanceAgent:
    def execute(self,**request):return run_hardware_governance("network",**request)
class ServerHardwareGovernanceAgent:
    def execute(self,**request):return run_hardware_governance("server",**request)

FEATURE_MANIFEST={"NW_RECONCILIATION":"reconcile_nw","SERVER_RECONCILIATION":"reconcile_server","NW_OS_BULK_LOAD":"generate_bulk_load","SERVER_OS_BULK_LOAD":"generate_bulk_load","NW_HARDWARE_GOVERNANCE":"NetworkHardwareGovernanceAgent","SERVER_HARDWARE_GOVERNANCE":"ServerHardwareGovernanceAgent"}
class ReconciliationAgent:
    def execute(self,source_type,cmdb,is_report,catalog,field_mapping=None):
        fi=FieldIntelligence();
        if field_mapping:fi.load_workbook(field_mapping)
        if source_type=="network":return reconcile_nw(read_tabular(cmdb),read_tabular(is_report),read_tabular(catalog),fi)
        return reconcile_server(read_tabular(cmdb),read_tabular(is_report),read_tabular(catalog),fi)[0]
class BulkLoadAgent:
    def execute(self,**request):return generate_bulk_load(**request)
class ApplicationService:
    def reconcile(self,request):return ReconciliationAgent().execute(**request)
    def bulk_load(self,request):return BulkLoadAgent().execute(**request)
    def hardware_governance(self,source_type,request):return (NetworkHardwareGovernanceAgent() if source_type=="network" else ServerHardwareGovernanceAgent()).execute(**request)

# ------------------------------------ Tkinter UI ------------------------------------
class ScrollFrame(ttk.Frame):
    def __init__(self,parent):
        super().__init__(parent);self.canvas=tk.Canvas(self,bg="#F4F6F8",highlightthickness=0);self.bar=ttk.Scrollbar(self,command=self.canvas.yview);self.inner=ttk.Frame(self.canvas,padding=18);self.win=self.canvas.create_window((0,0),window=self.inner,anchor="nw");self.inner.bind("<Configure>",lambda e:self.canvas.configure(scrollregion=self.canvas.bbox("all")));self.canvas.bind("<Configure>",lambda e:self.canvas.itemconfigure(self.win,width=e.width));self.canvas.configure(yscrollcommand=self.bar.set);self.canvas.pack(side="left",fill="both",expand=True);self.bar.pack(side="right",fill="y")

class App(tk.Tk):
    def __init__(self):
        super().__init__();self.title(APP_NAME);self.geometry("1500x900");self.minsize(1180,720);self.registry=ResourceRegistry();self.fi=FieldIntelligence();self.paths={x:"" for x in ALL_RESOURCES};self.status={};self.results={};self.busy=False;self.shutdown_event=threading.Event();cleanup_workdir(str(APP_DIR));self._build();self._restore();self.protocol("WM_DELETE_WINDOW",self.close_app)
    def _build(self):
        sf=ScrollFrame(self);sf.pack(fill="both",expand=True);root=sf.inner
        ttk.Label(root,text=APP_NAME,font=("Segoe UI",19,"bold"),foreground="#1F4E78").pack(anchor="w");ttk.Label(root,text="Network and Server reconciliation, catalogue validation, OS bulk load, and hardware category governance").pack(anchor="w",pady=(0,12))
        pane=ttk.PanedWindow(root,orient="horizontal");pane.pack(fill="both",expand=True);left=ttk.Frame(pane);right=ttk.Frame(pane);pane.add(left,weight=3);pane.add(right,weight=4)
        load=ttk.LabelFrame(left,text="File Loader",padding=10);load.pack(fill="x")
        for i,name in enumerate(ALL_RESOURCES):
            ttk.Label(load,text=name,width=35).grid(row=i,column=0,sticky="w",pady=3);ttk.Button(load,text="Browse",command=lambda n=name:self.browse(n)).grid(row=i,column=1,padx=3);ttk.Button(load,text="Clear",command=lambda n=name:self.clear(n)).grid(row=i,column=2,padx=3);lab=tk.Label(load,text="Not loaded",fg="#9C0006",anchor="w");lab.grid(row=i,column=3,sticky="ew");self.status[name]=lab
        load.columnconfigure(3,weight=1)
        actions=ttk.LabelFrame(left,text="Dedicated Actions",padding=10);actions.pack(fill="x",pady=10)
        descriptions={"Reconcile NW CMDB and IS":"Uses loaded NW CMDB, IS OS and OS Catalogue reports","Reconcile Server CMDB and IS":"Applies W-KYN scope, lifecycle priority and Server rules","Create NW - OS Bulk Load File":"Asks for NW Load To IS input and creates OS load from the template","Create Server - OS Bulk Load File":"Asks for Server Load To IS input and creates OS load from the template"}
        for text,cmd in [("Reconcile NW CMDB and IS",self.reconcile_nw_click),("Reconcile Server CMDB and IS",self.reconcile_server_click),("Create NW - OS Bulk Load File",lambda:self.bulk_click("network")),("Create Server - OS Bulk Load File",lambda:self.bulk_click("server"))]:
            row=ttk.Frame(actions);row.pack(fill="x",pady=3);ttk.Button(row,text=text,command=cmd,width=42).pack(side="left");ttk.Label(row,text=descriptions[text]).pack(side="left",padx=12)
        hw=ttk.LabelFrame(left,text="Hardware Governance Actions",padding=10);hw.pack(fill="x",pady=(0,10))
        for text,cmd,desc in [("Run NW HW Catalog and IS Category Match",lambda:self.hardware_click("network"),"Existing approved full NW governance action"),("Run Server HW Catalog and IS Category Match",lambda:self.hardware_click("server"),"Existing approved full Server governance action"),("Create NW Category Load File",lambda:self.category_load_click("network"),"Asks for NW Load To IS input; creates Category Load and validation"),("Create Server Category Load File",lambda:self.category_load_click("server"),"Asks for Server Load To IS input; creates Category Load and validation")]:
            row=ttk.Frame(hw);row.pack(fill="x",pady=3);ttk.Button(row,text=text,command=cmd,width=42).pack(side="left");ttk.Label(row,text=desc).pack(side="left",padx=12)
        controls=ttk.Frame(left);controls.pack(fill="x");ttk.Button(controls,text="Clear Session Files",command=self.clear_session).pack(side="left");ttk.Button(controls,text="Close Application",command=self.close_app).pack(side="right")
        self.pb=ttk.Progressbar(left,maximum=100);self.pb.pack(fill="x",pady=(12,3));self.pt=ttk.Label(left,text="Progress: 0% - Ready");self.pt.pack(anchor="w")
        logf=ttk.LabelFrame(right,text="Activity Log",padding=8);logf.pack(fill="both",expand=True);toolbar=ttk.Frame(logf);toolbar.pack(fill="x",pady=(0,4));self.jump_button=ttk.Button(toolbar,text="Jump to Latest",command=self.jump_to_latest);self.jump_button.pack(side="right");self.log=tk.Text(logf,wrap="word",font=("Consolas",10),undo=False);sb=ttk.Scrollbar(logf,command=self._log_scroll);self.log.configure(yscrollcommand=lambda a,b:(sb.set(a,b),self._log_view_changed(a,b)));self.log.pack(side="left",fill="both",expand=True);sb.pack(side="right",fill="y");self.log.bind("<MouseWheel>",self._log_mousewheel,add="+");self.log.bind("<Button-4>",self._log_mousewheel,add="+");self.log.bind("<Button-5>",self._log_mousewheel,add="+");self._log_auto_scroll=True;self._activity_name="";self.say("Application ready.")
    def _log_scroll(self,*args):
        self.log.yview(*args);self._log_auto_scroll=self.log.yview()[1]>=0.995
    def _log_view_changed(self,first,last):
        self._log_auto_scroll=float(last)>=0.995
    def _log_mousewheel(self,event):
        self.after_idle(lambda:setattr(self,"_log_auto_scroll",self.log.yview()[1]>=0.995))
    def jump_to_latest(self):
        self.log.see("end");self._log_auto_scroll=True
    def activity_start(self,name):
        self._activity_name=name;self.log.tag_remove("transient","1.0","end");self.say(f"{name} started.")
    def activity_complete(self,name,rows,outputs):
        set_output_progress(None)
        if self.log.tag_ranges("transient"):self.log.delete("transient.first","transient.last")
        self.say(f"{name} completed: {rows} rows processed, {outputs} output files created.")
        self.pb["value"]=100;self.pt.configure(text=f"Progress: 100% - {name} completed")
    def _restore(self):
        for display,rid in PERSISTENT_RESOURCES.items():
            item=self.registry.get(rid)
            if item:self.paths[display]=item["path"];loaded=datetime.fromisoformat(item["loaded_on"]).strftime("%d.%b.%Y").upper();self.status[display].configure(text=f"{item['file_name']} | loaded {loaded}",fg="#006100")
    def say(self,msg):
        if threading.current_thread() is not threading.main_thread():self.after(0,lambda:self.say(msg));return
        self.log.insert("end",f"[{datetime.now():%H:%M:%S}] {msg}\n")
        if self._log_auto_scroll:self.log.see("end")
    def progress(self,msg,pct):
        if threading.current_thread() is not threading.main_thread():self.after(0,lambda:self.progress(msg,pct));return
        self.pb["value"]=pct;self.pt.configure(text=f"Progress: {pct}% - {msg}")
        if self.log.tag_ranges("transient"):self.log.delete("transient.first","transient.last")
        self.log.insert("end",f"[{datetime.now():%H:%M:%S}] {msg}\n","transient")
        if self._log_auto_scroll:self.log.see("end")
    def browse(self,name):
        p=filedialog.askopenfilename(title="Select "+name,filetypes=FILE_TYPES)
        if not p:return
        self.paths[name]=p
        if name in PERSISTENT_RESOURCES:self.registry.register(PERSISTENT_RESOURCES[name],p);item=self.registry.get(PERSISTENT_RESOURCES[name]);loaded=datetime.fromisoformat(item["loaded_on"]).strftime("%d.%b.%Y").upper();text=f"{Path(p).name} | loaded {loaded}"
        else:text=Path(p).name
        self.status[name].configure(text=text,fg="#006100");self.say(f"Loaded {name}: {p}")
        if name=="Field List and Mapping":self.fi.load_workbook(p);self.say("Enterprise Field Intelligence mapping loaded.")
    def clear(self,name):
        self.paths[name]="";self.status[name].configure(text="Not loaded",fg="#9C0006")
    def clear_session(self):
        for x in SESSION_RESOURCES:self.clear(x)
        self.results.clear();self.say("Session files cleared. Persistent template/catalog metadata retained.")
    def require(self,*names):
        missing=[x for x in names if not self.paths.get(x)]
        if missing:messagebox.showerror("Missing resources","Load:\n"+"\n".join(missing));return False
        return True
    def execute(self,fn):
        if self.busy:return
        self.busy=True
        def worker():
            try:fn()
            except Exception as exc:
                err=exc;tb=traceback.format_exc();self.after(0,lambda e=err,t=tb:self.fail(e,t))
            finally:self.busy=False
        threading.Thread(target=worker,daemon=True).start()
    def fail(self,e,t):self.say(t);self.progress("Failed",0);messagebox.showerror("Operation failed",str(e))
    def sync_yes(self,title,text):
        result=[];event=threading.Event();self.after(0,lambda:(result.append(messagebox.askyesno(title,text)),event.set()));event.wait();return result[0]
    def output_dir(self):return filedialog.askdirectory(title="Select output folder")
    def _load_field_mapping(self):
        p=self.paths.get("Field List and Mapping")
        if p:self.fi.load_workbook(p)
    def close_app(self):
        self.shutdown_event.set();self.say("Closing application cleanly.");self.quit();self.destroy()
    def hardware_click(self,source):
        category="IS Network Category Report" if source=="network" else "IS Server Category Report";cmdb="NW CMDB Report" if source=="network" else "Server CMDB Report";catalog="IS Product Catalog Network" if source=="network" else "IS Product Catalog Server"
        if not self.require(cmdb,category,catalog,"IS Operating System Report","IS Product Catalog Operating System"):return
        out=self.output_dir()
        if not out:return
        def task():
            cleanup_workdir(out);self._load_field_mapping()
            if source not in self.results:
                self.progress(f"Running prerequisite {'NW' if source=='network' else 'Server'} reconciliation",3)
                self.results[source]=reconcile_nw(read_tabular(self.paths[cmdb]),read_tabular(self.paths["IS Operating System Report"]),read_tabular(self.paths["IS Product Catalog Operating System"]),self.fi,self.progress) if source=="network" else reconcile_server(read_tabular(self.paths[cmdb]),read_tabular(self.paths["IS Operating System Report"]),read_tabular(self.paths["IS Product Catalog Operating System"]),self.fi,self.progress)[0]
            state=run_hardware_governance(source,self.paths[cmdb],self.paths[category],self.paths[catalog],out,self.progress,self.results[source],self.paths["IS Operating System Report"],True);self.results[source+"_governance"]=state
            for p in state["outputs"]:self.say("Created: "+p)
            self.after(0,lambda outputs=list(state["outputs"]):messagebox.showinfo("Hardware Governance Complete","\n".join(outputs)))
        self.execute(task)
    def category_load_click(self,source):
        domain="NW" if source=="network" else "Server";cmdb="NW CMDB Report" if source=="network" else "Server CMDB Report";category="IS Network Category Report" if source=="network" else "IS Server Category Report";catalog="IS Product Catalog Network" if source=="network" else "IS Product Catalog Server"
        if not self.require(cmdb,category,catalog):return
        load_path=filedialog.askopenfilename(title=f"Select {domain} Load To IS File",filetypes=FILE_TYPES)
        if not load_path:return
        out=self.output_dir()
        if not out:return
        def task():
            name=f"{domain} Category Load";self.after(0,lambda:self.activity_start(name));cleanup_workdir(out);set_output_progress(self.progress)
            load_df=read_load_to_is_file(load_path,source,self.progress)
            state=run_hardware_governance(source,self.paths[cmdb],self.paths[category],self.paths[catalog],out,self.progress,load_df,self.paths.get("IS Operating System Report"),False)
            category_rows=state["category_load"].copy();validation=state["decisions"].copy();outputs=[];data_count=len(category_rows)
            if data_count:
                self.progress(f"Preparing {domain} Category Load metadata",72)
                category_template,category_metadata=read_tabular(self.paths[category],True)
                reference_count=len(category_rows.columns)-len(category_template.columns)
                if reference_count<0:raise ValueError("Category Load schema does not align with the loaded Category report")
                target_meta=[clean(category_metadata.get(c,"")) for c in category_template.columns] if category_metadata is not None else [""]*len(category_template.columns)
                metadata_row=["Reference"]*reference_count+target_meta
                category_output=pd.concat([pd.DataFrame([metadata_row],columns=category_rows.columns),category_rows],ignore_index=True)
                load_out=_write_formatted_load(category_output,next_versioned_path(out,domain+" Category Load","xlsx"));outputs.append(load_out)
            else:
                self.progress(f"No {domain} Category Load created: 0 eligible Category rows",72)
            blocked=int(validation["Recommended Action"].isin(["CATEGORY DATA QUALITY REVIEW","NO LOAD ACTION"]).sum()) if not validation.empty and "Recommended Action" in validation.columns else 0
            val_out=write_workbook(next_versioned_path(out,domain+" Category Load Validation","xlsx"),{"Workflow Summary":pd.DataFrame([("Load To IS candidates",len(load_df)),("Category Load rows",data_count),("Blocked or review rows",blocked),("Category Load workbook created","Yes" if data_count else "No")],columns=["Metric","Value"]),"Hardware and Category Review":validation});outputs.append(val_out)
            self.after(0,lambda:self.activity_complete(name,len(load_df),len(outputs)));self.after(0,lambda files=list(outputs):messagebox.showinfo(f"{domain} Category Load Complete","\n".join(files)))
        self.execute(task)
    def reconcile_nw_click(self):
        if not self.require("IS Operating System Report","NW CMDB Report","IS Product Catalog Operating System"):return
        out=self.output_dir();
        if not out:return
        def task():
            set_output_progress(self.progress)
            self._load_field_mapping();r=reconcile_nw(read_tabular(self.paths["NW CMDB Report"]),read_tabular(self.paths["IS Operating System Report"]),read_tabular(self.paths["IS Product Catalog Operating System"]),self.fi,self.progress);self.results["network"]=r;p=next_versioned_path(out,"NW CMDB to IS Reconciliation","xlsx");p=write_workbook(p,{"Summary":nw_summary(r),"Reconciliation":r,"Load To IS":r[r.Action.str.contains("Load To IS",na=False)],"Retire From IS":r[r.Action.str.contains("Retire From IS",na=False)],"Inventory OS Update Required":r[r.Action.str.contains("Inventory OS update required",na=False)],"Catalogue Update Required":r[r.Action.str.contains("Catalogue update required",na=False)],"Retired Devices - No Action":r[r.Action.eq("Retired Device - No Action Required")]});self.say("Created: "+p);self.after(0,lambda path=p:messagebox.showinfo("Complete",path))
        self.execute(task)
    def reconcile_server_click(self):
        if not self.require("IS Operating System Report","Server CMDB Report","IS Product Catalog Operating System"):return
        out=self.output_dir()
        if not out:return
        def task():
            set_output_progress(self.progress)
            self._load_field_mapping();r,outside,dups,historical=reconcile_server(read_tabular(self.paths["Server CMDB Report"]),read_tabular(self.paths["IS Operating System Report"]),read_tabular(self.paths["IS Product Catalog Operating System"]),self.fi,self.progress);self.results["server"]=r;p=next_versioned_path(out,"Server CMDB to IS Reconciliation","xlsx");p=write_workbook(p,{"Summary":server_summary(r,outside,dups,historical),"Reconciliation":r,"Load To IS":r[r.Action.str.contains("Load To IS",na=False)] if not r.empty else r,"Retire From IS":r[r.Action.str.contains("Retire From IS",na=False)] if not r.empty else r,"Out of W-KYN Scope":outside,"Historical EOL Suppressed":historical,"CMDB Duplicate Review":dups,"Retired Devices - No Action":r[r.Action.eq("Retired Device - No Action Required")]});self.say("Created: "+p);self.after(0,lambda path=p:messagebox.showinfo("Complete",path))
        self.execute(task)
    def _template_notice(self):
        item=self.registry.get("bulk_load_template")
        if not item:return False
        loaded=datetime.fromisoformat(item["loaded_on"]).strftime("%d.%b.%Y").upper()
        return self.sync_yes("Current Bulk Load Template",f"Current Load template was loaded on date {loaded}.\n\nThe bulk load will be created based on the loaded template. If there are changes in the bulk load template, please upload a new Bulk Load template.\n\nContinue with {item['file_name']}?")
    def bulk_click(self,source):
        domain="NW" if source=="network" else "Server";category="IS Network Category Report" if source=="network" else "IS Server Category Report"
        if not self.require("IS Operating System Report","Bulk Load Template",category):return
        if not self._template_notice():return
        load_path=filedialog.askopenfilename(title=f"Select {domain} Load To IS File",filetypes=FILE_TYPES)
        if not load_path:return
        out=self.output_dir()
        if not out:return
        fmt="xlsx" if messagebox.askyesno("Output format","Create XLSX? Select No for CSV.") else "csv"
        def task():
            name=f"{domain} OS Bulk Load";self.after(0,lambda:self.activity_start(name));cleanup_workdir(out);set_output_progress(self.progress);load_df=read_load_to_is_file(load_path,source,self.progress);state=category_decisions_from_load(source,load_df,self.paths[category],self.progress);paths=generate_bulk_load(load_df,source,self.paths["IS Operating System Report"],self.paths["Bulk Load Template"],out,fmt,lambda n:self.sync_yes("Missing FQDN",f"{n} devices do not have a valid FQDN. Generate .nw.wpp.net and include them?"),lambda n:True,self.progress,state)
            self.after(0,lambda:self.activity_complete(name,len(load_df),len(paths)));self.after(0,lambda outputs=list(paths):messagebox.showinfo(f"{domain} OS Bulk Load Complete","\n".join(outputs)))
        self.execute(task)

if __name__=="__main__":App().mainloop()
