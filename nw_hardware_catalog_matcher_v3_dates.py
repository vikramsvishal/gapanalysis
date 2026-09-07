#!/usr/bin/env python3
"""NW Hardware Catalog Matcher
Two GUI inputs only: IS NW Catalog and NW CMDB Report.
Supports .xlsx, .xlsm, .xlsb and .csv. Output is always .xlsx.
"""
from __future__ import annotations
import base64, csv, io, re, traceback
from collections import defaultdict
from dataclasses import dataclass
from datetime import date, datetime
from pathlib import Path
from tempfile import NamedTemporaryFile
from typing import Any, Iterable
from openpyxl import load_workbook
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.utils.datetime import from_excel

APP_TITLE="NW Hardware Catalog Matcher"
TEMPLATE_B64="UEsDBBQABgAIAAAAIQBi7p1oXgEAAJAEAAATAAgCW0NvbnRlbnRfVHlwZXNdLnhtbCCiBAIooAACAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAACslMtOwzAQRfdI/EPkLUrcskAINe2CxxIqUT7AxJPGqmNbnmlp/56J+xBCoRVqN7ESz9x7MvHNaLJubbaCiMa7UgyLgcjAVV4bNy/Fx+wlvxcZknJaWe+gFBtAMRlfX41mmwCYcbfDUjRE4UFKrBpoFRY+gOOd2sdWEd/GuQyqWqg5yNvB4E5W3hE4yqnTEOPRE9RqaSl7XvPjLUkEiyJ73BZ2XqVQIVhTKWJSuXL6l0u+cyi4M9VgYwLeMIaQvQ7dzt8Gu743Hk00GrKpivSqWsaQayu/fFx8er8ojov0UPq6NhVoXy1bnkCBIYLS2ABQa4u0Fq0ybs99xD8Vo0zL8MIg3fsl4RMcxN8bZLqej5BkThgibSzgpceeRE85NyqCfqfIybg4wE/tYxx8bqbRB+QERfj/FPYR6brzwEIQycAhJH2H7eDI6Tt77NDlW4Pu8ZbpfzL+BgAA//8DAFBLAwQUAAYACAAAACEAtVUwI/QAAABMAgAACwAIAl9yZWxzLy5yZWxzIKIEAiigAAIAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAKySTU/DMAyG70j8h8j31d2QEEJLd0FIuyFUfoBJ3A+1jaMkG92/JxwQVBqDA0d/vX78ytvdPI3qyCH24jSsixIUOyO2d62Gl/pxdQcqJnKWRnGs4cQRdtX11faZR0p5KHa9jyqruKihS8nfI0bT8USxEM8uVxoJE6UchhY9mYFaxk1Z3mL4rgHVQlPtrYawtzeg6pPPm3/XlqbpDT+IOUzs0pkVyHNiZ9mufMhsIfX5GlVTaDlpsGKecjoieV9kbMDzRJu/E/18LU6cyFIiNBL4Ms9HxyWg9X9atDTxy515xDcJw6vI8MmCix+o3gEAAP//AwBQSwMEFAAGAAgAAAAhAAMMuYa7AwAAfQkAAA8AAAB4bC93b3JrYm9vay54bWysVm1vozgQ/n7S/QfEd4rNSyCodMWrrlK7qtJseidVWrngFKuAc8Y0qar97zsmJGmb0ynXvSixsT08fsbzzDjnXzZNrT1T0THehjo+Q7pG24KXrH0M9W/z3PB1rZOkLUnNWxrqL7TTv1z8/tv5mounB86fNABou1CvpFwFptkVFW1Id8ZXtIWVJRcNkTAUj2a3EpSUXUWpbGrTQmhiNoS1+hYhEKdg8OWSFTTlRd/QVm5BBK2JBPpdxVbdDq0pToFriHjqV0bBmxVAPLCayZcBVNeaIrh8bLkgDzW4vcGuthHwncAPI2is3U6wdLRVwwrBO76UZwBtbkkf+Y+RifG7I9gcn8FpSI4p6DNTMdyzEpNPsprssSYHMIx+GQ2DtAatBHB4n0Rz99ws/eJ8yWq62EpXI6vVV9KoSNW6VpNOZiWTtAx1D4Z8Td9NiH4V96yGVRtZlq+bF3s53witpEvS13IOQt7Bh7qFLBshZQnCiGpJRUskTXgrQYejX7+quQE7qTgoXJvRv3smKCQW6At8hZYUAXnoboistF7UoZ4E9986cP9+wZ4EaRasq0h9n/J1W3PIs/s3AiXH2fAfJEoK5bcJjm/JbZ8/HgJwFMFOhjdSaPB8mV5BKG7JMwQGwl+OeXsJJ+9/f40TK7adNDJSK3ENB1uuMbV830DIzfLUsxI7z36AF2ISFJz0shqDrTBD3YHIHi1dk81uBaOgZ+Vh/1c0fgAfoQ/Nbu2H8lSVtQWj6+4gCzXUNnesLfk61A2sxPzyfrgeFu9YKSsopFPHApPt3B+UPVbAGCMlIVU7FLNQf02tGKcumhoon3qGkyS5EWE/MqIM+V4c506aTgZG5htKQwEFakOvtYPob1VRxVCpVa9OF55FoPYQlyUeord7rSB1ASJX3WA4xciaKgu6kVedHHrQFwN62EGRh6aOgTIb4uNPLcN3bMtInNTKXC9Ls9hV8VEXQPB/lMFB5sHuZlEsKyLkXJDiCe6jGV3GpAMlbR0Cvm/Jxq4fIxsoOjnOQUxTZMTxxDHcNLddD6dJ5uYHssr95SeLkG8Ob1Mie0hQlZvDOFBtPs7uJ5fbiTFO75IumKXq3Me3/83wFryv6YnG+eJEw+Tr9fz6RNurbP79Lj/VOLqO0+h0+2g2i/6aZ3/utjD/8UDNDwFPsTNFdhYZtp04huPlnuHnyDVsx3MS14kzjLxDwOt18fy5eFuOuVNk8vZ/wliMVPwVeDD+idI6KscluDWG1BuIK/pDfu3RLn4CAAD//wMAUEsDBBQABgAIAAAAIQCBPpSX8wAAALoCAAAaAAgBeGwvX3JlbHMvd29ya2Jvb2sueG1sLnJlbHMgogQBKKAAAQAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAACsUk1LxDAQvQv+hzB3m3YVEdl0LyLsVesPCMm0KdsmITN+9N8bKrpdWNZLLwNvhnnvzcd29zUO4gMT9cErqIoSBHoTbO87BW/N880DCGLtrR6CRwUTEuzq66vtCw6acxO5PpLILJ4UOOb4KCUZh6OmIkT0udKGNGrOMHUyanPQHcpNWd7LtOSA+oRT7K2CtLe3IJopZuX/uUPb9gafgnkf0fMZCUk8DXkA0ejUISv4wUX2CPK8/GZNec5rwaP6DOUcq0seqjU9fIZ0IIfIRx9/KZJz5aKZu1Xv4XRC+8opv9vyLMv072bkycfV3wAAAP//AwBQSwMEFAAGAAgAAAAhAMSmhy+PAgAA1wYAABgAAAB4bC93b3Jrc2hlZXRzL3NoZWV0MS54bWyclWtvmzAUhr9P2n+w/D3hEkISBFRp06xd161ru313jEmsAma2c6mm/fcdw0ILZFJVKUg2efyei88rwrNDnqEdk4qLIsLO0MaIFVQkvFhH+MfjcjDFSGlSJCQTBYvwM1P4LP74IdwL+aQ2jGkECoWK8EbrMrAsRTcsJ2ooSlbAP6mQOdGwlWtLlZKRpDqUZ5Zr276VE17gWiGQb9EQacopWwi6zVmhaxHJMqIhf7XhpTqq5fQtcjmRT9tyQEVegsSKZ1w/V6IY5TS4XhdCklUGdR8cj1B0kPBz4Rkdw1Tve5FyTqVQItVDULbqnPvlz6yZRWij1K//TTKOZ0m24+YCX6Tc96XkjBst90Vs9E4xvxEz7ZLBlicR/j11bM9ferPB+fJyPvDmF5PB3HfPB/bldOIuFwvfHtt/cBwmHG7YVIUkSyM8d4JHB1txWM3PT8726tUaabJ6YBmjmkEMByMznishngx4Da9sUFQVYBQJ1XzHLliWRfjGTPivKgYsIYDVRHi9PkZbVgN9J1HCUrLN9L3YXzG+3mgI6w3HUKiZlCB5XjBFYUQh9HA0bhJfEE3iUIo9guuGPFVJjHmcwAXnGZGR/3+ROKTm2Nycq04Dr6CwXWyH1g6ypf+I8z7htImLPuG2iUWfGLWJyz7htYllnxi3iU99wm8TV31i0iauawImtunHtE187hOzNnHTJ5xOU7+cQDpdvT2BdNr6tUbAEU2yTqev304gncbenUA6nf1eI97rQJ3W3p9AOr19OBGo01ywpRnHVkUv3bVg1o8+qoe/JGt2S+SaFwplLK0cMsFI1iayh7DWojS+mYCdVkJrkR93G/iAMLAAeAqjVAh93BjfNp+k+C8AAAD//wMAUEsDBBQABgAIAAAAIQD2YLRBuAcAABEiAAATAAAAeGwvdGhlbWUvdGhlbWUxLnhtbOxazY8btxW/B8j/QMxd1szoe2E50Kc39u564ZVd5EhJlIZeznBAUrsrFAEK59RLgQJp0UuB3nooigZogAa55I8xYCNN/4g8ckaa4YqKvf5AkmJ3LzPU7z3+5r3HxzePc/eTq5ihCyIk5UnXC+74HiLJjM9psux6TybjSttDUuFkjhlPSNdbE+l9cu/jj+7iAxWRmCCQT+QB7nqRUulBtSpnMIzlHZ6SBH5bcBFjBbdiWZ0LfAl6Y1YNfb9ZjTFNPJTgGNQ+WizojKCJVund2ygfMbhNlNQDMybOtGpiSRjs/DzQCLmWAybQBWZdD+aZ88sJuVIeYlgq+KHr+ebPq967W8UHuRBTe2RLcmPzl8vlAvPz0MwpltPtpP4obNeDrX4DYGoXN2rr/60+A8CzGTxpxqWsM2g0/XaYY0ug7NKhu9MKaja+pL+2wznoNPth3dJvQJn++u4zjjujYcPCG1CGb+zge37Y79QsvAFl+OYOvj7qtcKRhTegiNHkfBfdbLXbzRy9hSw4O3TCO82m3xrm8AIF0bCNLj3FgidqX6zF+BkXYwBoIMOKJkitU7LAM4jiXqq4REMqU4bXHkpxwiUM+2EQQOjV/XD7byyODwguSWtewETuDGk+SM4ETVXXewBavRLk5TffvHj+9Yvn/3nxxRcvnv8LHdFlpDJVltwhTpZluR/+/sf//fV36L///tsPX/7JjZdl/Kt//v7Vt9/9lHpYaoUpXv75q1dff/XyL3/4/h9fOrT3BJ6W4RMaE4lOyCV6zGN4QGMKmz+ZiptJTCJMLQkcgW6H6pGKLODJGjMXrk9sEz4VkGVcwPurZxbXs0isFHXM/DCKLeAx56zPhdMAD/VcJQtPVsnSPblYlXGPMb5wzT3AieXg0SqF9EpdKgcRsWieMpwovCQJUUj/xs8JcTzdZ5Radj2mM8ElXyj0GUV9TJ0mmdCpFUiF0CGNwS9rF0FwtWWb46eoz5nrqYfkwkbCssDMQX5CmGXG+3ilcOxSOcExKxv8CKvIRfJsLWZl3Egq8PSSMI5GcyKlS+aRgOctOf0hhsTmdPsxW8c2Uih67tJ5hDkvI4f8fBDhOHVypklUxn4qzyFEMTrlygU/5vYK0ffgB5zsdfdTSix3vz4RPIEEV6ZUBIj+ZSUcvrxPuL0e12yBiSvL9ERsZdeeoM7o6K+WVmgfEcLwJZ4Tgp586mDQ56ll84L0gwiyyiFxBdYDbMeqvk+IhDJJ1zW7KfKISitkz8iS7+FzvL6WeNY4ibHYp/kEvG6F7lTAYnRQeMRm52XgCYXyD+LFaZRHEnSUgnu0T+tphK29S99Ld7yuheW/N1ljsC6f3XRdggy5sQwk9je2zQQza4IiYCaYoiNXugURy/2FiN5XjdjKKbewF23hBiiMrHonpsnrip8TLAS//Hlqnw9W9bgVv0u9sy+vHF6rcvbhfoW1zRCvklMC28lu4rotbW5LG+//vrTZt5ZvC5rbgua2oHG9gn2QgqaoYaC8KVo9pvET7+37LChjZ2rNyJE0rR8JrzXzMQyanpRpTG77gGkEl/p5YAILtxTYyCDB1W+ois4inEJ/KDBdzKXMVS8lSrmEtpEZNv1Uck23aT6t4mM+z9qdpr/kZyaUWBXjfgMaT9k4tKpUhm628kHNb0PdsF2aVuuGgJa9CYnSZDaJmoNEazP4GhK6c/Z+WHQcLNpa/cZVO6YAaluvwHs3grf1rteoZ4ygIwc1+lz7KXP1xrvaOe/V0/uMycoRAK3FXU93NNe9j6efLgu1N/C0RcI4JQsrm4TxlSnwZARvw3l0lvvuPxVwN/V1p3CpRU+bYrMaChqt9ofwtU4i13IDS8qZgiXoEtZ4CIvOQzOcdr0F9I3hMk4heKR+98JsCYcvMyWyFf82qSUVUg2xjDKLm6yT+SemigjEaNz19PNvw4ElJolk5DqwdH+p5EK94H5p5MDrtpfJYkFmquz30oi2dHYLKT5LFs5fjfjbg7UkX4G7z6L5JZqylXiMIcQarUB7d04lHB8EmavnFM7DtpmsiL9rO1Oe/a1DriIfY5ZGON9Sytk8g5sNZUvH3G1tULrLnxkMumvC6VLvsO+87b5+r9aWK/bHTrFpWmlFb5vubPrhdvkSq2IXtVhluft6zu1skh0EqnObePe9v0StmMyiphnv5mGdtPNRm9p7rAhKu09zj922m4TTEm+79YPc9ajVO8SmsDSBbw7Oy2fbfPoMkscQThFXLDvtZgncmdIyPRXGt1M+X+eXTGaJJvO5LkqzVP6YLBCdX3W90FU55ofHeTXAEkCbmhdW2FbQWe3Zgnqzy0WzBbsVzsrYa/WqLbyV2ByzboVNa9FFW11tTtR1rW5m1g7LntqkYWMpuNq1IrTJBYbSOTvMzXIv5JkrlVfacIVWgna93/qNXn0QNgYVv90YVeq1ul9pN3q1Sq/RqAWjRuAP++HnQE9FcdDIvnwYw2kQW+ffP5jxnW8g4s2B150Zj6vcfONQNd4330AE4f5vIMCRQCscBfWwFw4qg2HQrNTDYbPSbtV6lUHYHIY92LSb497nHrow4KA/HI7HjbDSHACu7vcalV6/Nqg026N+OA5G9aEP4Hz7uYK3GJ1zc1vApeF170cAAAD//wMAUEsDBBQABgAIAAAAIQBKog6ycwMAAD8KAAANAAAAeGwvc3R5bGVzLnhtbMRW227bOBB9X6D/QPBd0cWSKxmSijiOgALdYIFkgX2lJcohyotA0Yncxf57h7rY8jbtuukWfRI5HB2emTlDMn3XCY6eqG6Zkhn2rzyMqCxVxeQuw38+FE6MUWuIrAhXkmb4QFv8Ln/zW9qaA6f3j5QaBBCyzfCjMc3KddvykQrSXqmGSliplRbEwFTv3LbRlFSt/UlwN/C8pSsIk3hAWInyEhBB9Md945RKNMSwLePMHHosjES5er+TSpMtB6qdH5ISdf5SB6jT0ya99Yt9BCu1alVtrgDXVXXNSvol3cRNXFKekAD5dUh+5HrBWeydfiVS6Gr6xGz5cJ7WSpoWlWovTYYXQNSmYPVRqmdZ2CWo8OiVp+0n9EQ4WHzs5mmpuNLIQOkgc71FEkEHj+vGqBbdEa3Vs/WtiWD8MKwF1tCXfHQWDApgja4lM1DK0631+lUbJj81wD7OFgJlnB9zH9o0gyFPQaSGalnABI3jh0MDSZbQT0Oeer//8N5pcvCD6PIfWsVZZVnsbvrS6t02w0WRBBsv8izMdlxgsqIdrTK8DHv0GWFbxEvI/XuvUUYhRoZZJXpXb5Mkif1lHMdJuPDDsJfNxOBCd/clZj1ByP5W6QpOsUn7VuaDKU85rQ3Eq9nu0X6Namz0yhjo9DytGNkpSbhV7PTHOADYknJ+b0+6v+oz7K5Gci8KYd5D4uDMtFqfhpCxcTjgDROLP0cbsGewEVD+fljU1Uf8r/0dAL+XSR3/RqRp+MGeEbb7h9k1Zzsp6GDKUzJN0bMmzQPtelcbVld/nflsbxjOE/LNvW27/M9MIKwpC3AyfjcTCPQsvz6oe6r6xUHe7cWW6qK/Ek/hnaf9B0LvBQaSmun2TLVH/SF7uGf4zhLhcE2MGkLbPePQsi8oFjCr7tQD/Qli7D3bd8dxF2iFitZkz83DcTHDp/HvtGJ7Afkavf5gT8r0EBk+jT/YVvWX9jgCmX1o4RqBL9prluG/b9dvk81tETixt46dcEEjJ4nWGycKb9abTZF4gXfzz+y2/4G7vn+cQN39cNVyeBHoMdiR/P3JluHZZKDfH6ZAe849CZbedeR7TrHwfCdcktiJl4vIKSI/2CzD9W1URDPu0SvfBJ7r+8PrwpKPVoYJypmcajVVaG6FIsH0G0G4UyXc08sv/wwAAP//AwBQSwMEFAAGAAgAAAAhAKbaQVZZAQAAHwMAABQAAAB4bC9zaGFyZWRTdHJpbmdzLnhtbHySTU/DMAyG70j8hyh3lrEDQqjthDoQEwyQBuJYea3bRmqSkjhA/z0p25BYO47x+zh+/RHNv1TDPtA6aXTMzydTzlDnppC6ivnry+3ZJWeOQBfQGI0x79DxeXJ6EjlHLORqF/OaqL0SwuU1KnAT06IOSmmsAgpPWwnXWoTC1YikGjGbTi+EAqk5y43XFPNZKOu1fPeY/gaSyMkkouTaOST2bKUC27F77CJBSSR6cQuky8kjKByEjS5l5S1QaI0tCdWAaMC5QRAIK2MHVdZ+kx+RHmSJLO3yBtmaoBo4OdR7ivyg8gq0LyEnb9EeulqZAhu2XIzGR9vfZmivNsPf1mglNGxcTIHCriv2n50wiR+mBlt8gsWMuna4gB3EfinVd3HYwt1bdqOL7KnM1r5tjaVsEeZ8nOqnOYrsjYcDLjJTZk0PFiN/7cEd53Zl/6IiHHjyDQAA//8DAFBLAwQUAAYACAAAACEAh4nlv0QBAABrAgAAEQAIAWRvY1Byb3BzL2NvcmUueG1sIKIEASigAAEAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAjJJRT8MgFIXfTfwPDe8ttDNTSdslavbkEhPnNL4RuNvICiWAdvv30nar1fngI5xzP865IZ/tVRV9gnWy1gVKE4Ii0LwWUm8K9LKcxzcocp5pwapaQ4EO4NCsvLzIuaG8tvBkawPWS3BRIGlHuSnQ1ntDMXZ8C4q5JDh0ENe1VcyHo91gw/iObQBnhEyxAs8E8wy3wNgMRHRECj4gzYetOoDgGCpQoL3DaZLib68Hq9yfA50ycirpDyZ0OsYdswXvxcG9d3IwNk2TNJMuRsif4rfF43NXNZa63RUHVOaCU26B+dqWK7mzTEUr6basyvFIabdYMecXYeFrCeLu8Nt8bgjkrkiPBxGFaLQvclJeJ/cPyzkqM5JNY3Ibk3RJrukVodnte/v+j/k2an+hjin+T0xpRkbEE6DM8dn3KL8AAAD//wMAUEsDBBQABgAIAAAAIQBsBndilQEAACADAAAQAAgBZG9jUHJvcHMvYXBwLnhtbCCiBAEooAABAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAJyS32/TMBDH35H4HyK/r04HmlDleIIONCQmKrUbz8a5NNYc27q7RS1/PU6idinwxNv9+Orrj+9O3R46X/SA5GKoxHJRigKCjbUL+0o87r5cfRAFsQm18TFAJY5A4la/faM2GBMgO6AiWwSqRMucVlKSbaEztMjtkDtNxM5wTnEvY9M4C3fRvnQQWF6X5Y2EA0Ooob5KZ0MxOa56/l/TOtqBj552x5SBtfqYknfWcP6lfnAWI8WGi88HC17JeVNlui3YF3R81KWS81RtrfGwzsa6MZ5AydeCugczDG1jHJJWPa96sByxIPcrj+1aFD8NwYBTid6gM4Ez1iCbkjH2iRj1j4jP1AIwKZkFU3EM59p57N7r5SjIwaVwMJhAcuMScefYA31vNgb5H8TLOfHIMPFOONuBb3pzzjd+Ob/0h/c6dsmEo/766aFYR0wRx00oeWqoby4802PaxTvDcJruZVFtW4NQ54Wcp38uqPs8WPSDybo1YQ/1SfN3Y7iFp+ng9fJmUb4r85pnNSVfT1v/BgAA//8DAFBLAQItABQABgAIAAAAIQBi7p1oXgEAAJAEAAATAAAAAAAAAAAAAAAAAAAAAABbQ29udGVudF9UeXBlc10ueG1sUEsBAi0AFAAGAAgAAAAhALVVMCP0AAAATAIAAAsAAAAAAAAAAAAAAAAAlwMAAF9yZWxzLy5yZWxzUEsBAi0AFAAGAAgAAAAhAAMMuYa7AwAAfQkAAA8AAAAAAAAAAAAAAAAAvAYAAHhsL3dvcmtib29rLnhtbFBLAQItABQABgAIAAAAIQCBPpSX8wAAALoCAAAaAAAAAAAAAAAAAAAAAKQKAAB4bC9fcmVscy93b3JrYm9vay54bWwucmVsc1BLAQItABQABgAIAAAAIQDEpocvjwIAANcGAAAYAAAAAAAAAAAAAAAAANcMAAB4bC93b3Jrc2hlZXRzL3NoZWV0MS54bWxQSwECLQAUAAYACAAAACEA9mC0QbgHAAARIgAAEwAAAAAAAAAAAAAAAACcDwAAeGwvdGhlbWUvdGhlbWUxLnhtbFBLAQItABQABgAIAAAAIQBKog6ycwMAAD8KAAANAAAAAAAAAAAAAAAAAIUXAAB4bC9zdHlsZXMueG1sUEsBAi0AFAAGAAgAAAAhAKbaQVZZAQAAHwMAABQAAAAAAAAAAAAAAAAAIxsAAHhsL3NoYXJlZFN0cmluZ3MueG1sUEsBAi0AFAAGAAgAAAAhAIeJ5b9EAQAAawIAABEAAAAAAAAAAAAAAAAArhwAAGRvY1Byb3BzL2NvcmUueG1sUEsBAi0AFAAGAAgAAAAhAGwGd2KVAQAAIAMAABAAAAAAAAAAAAAAAAAAKR8AAGRvY1Byb3BzL2FwcC54bWxQSwUGAAAAAAoACgCAAgAA9CEAAAAA"
SHORT_DATE_FORMAT="mm/dd/yyyy"

def txt(v): return "" if v is None else str(v).strip()
def hkey(v): return re.sub(r"[^a-z0-9]+","",txt(v).lower())
def compact(v): return re.sub(r"[^a-z0-9]+","",txt(v).lower())
def tokens(v):
    s=txt(v).lower(); out=[]
    for x in re.findall(r"[a-z0-9]+",s)+re.findall(r"[a-z]+|[0-9]+",s):
        if len(x)>=2 and x not in out: out.append(x)
    return out

@dataclass(frozen=True)
class Table:
    headers:list[str]; rows:list[list[Any]]; source_sheet:str
    @property
    def cols(self):
        d={}
        for i,h in enumerate(self.headers):
            if hkey(h) and hkey(h) not in d: d[hkey(h)]=i
        return d

@dataclass(frozen=True)
class CatalogRecord:
    hardware_type:str; hardware_model:str; manufacturer:str
    end_of_life_date:Any; end_of_support_date:Any; source_row:int
    type_key:str; model_key:str; type_tokens:tuple[str,...]; model_tokens:tuple[str,...]

@dataclass(frozen=True)
class Indexes:
    records:tuple[CatalogRecord,...]
    exact_model_index:dict[str,tuple[int,...]]
    hardware_type_index:dict[str,tuple[int,...]]
    partial_token_index:dict[str,tuple[int,...]]

@dataclass(frozen=True)
class Match:
    record:CatalogRecord|None; status:str; score:int=0; ambiguous:bool=False

def decode_csv(path):
    raw=path.read_bytes()
    for enc in ("utf-8-sig","utf-8","cp1252","latin-1"):
        try:return raw.decode(enc)
        except UnicodeDecodeError:pass
    raise ValueError(f"Unable to decode {path.name}")

def best_header(rows, expected, filename):
    exp={hkey(x) for x in expected}; best=(-1,0)
    for i,row in enumerate(rows[:20]):
        score=len({hkey(x) for x in row if txt(x)} & exp)
        if score>best[0]:best=(score,i)
    if best[0]<=0: raise ValueError(f"Expected columns were not found in {filename}.")
    return best[1]

def read_csv(path, expected):
    content=decode_csv(path)
    try:dialect=csv.Sniffer().sniff(content[:65536],delimiters=",;\t|")
    except csv.Error:dialect=csv.excel
    rows=list(csv.reader(io.StringIO(content),dialect))
    if not rows:raise ValueError(f"{path.name} is empty.")
    hr=best_header(rows,expected,path.name); heads=[txt(x) for x in rows[hr]]; w=len(heads)
    return Table(heads,[(r+[""]*w)[:w] for r in rows[hr+1:]],"CSV")

def read_excel(path,expected):
    wb=load_workbook(path,data_only=True,read_only=True)
    try:
        best=(-1,None,0)
        exp={hkey(x) for x in expected}
        for ws in wb.worksheets:
            for rn in range(1,min(ws.max_row,20)+1):
                vals=[c.value for c in ws[rn]]
                score=len({hkey(x) for x in vals if txt(x)} & exp)
                if score>best[0]:best=(score,ws,rn)
        if best[0]<=0:raise ValueError(f"Expected columns were not found in {path.name}.")
        ws,rn=best[1],best[2]; heads=[txt(c.value) for c in ws[rn]]; w=len(heads)
        return Table(heads,[list(r) for r in ws.iter_rows(min_row=rn+1,max_col=w,values_only=True)],ws.title)
    finally:wb.close()

def read_xlsb(path,expected):
    try:from pyxlsb import open_workbook
    except ImportError as e:raise RuntimeError("XLSB support requires pyxlsb. Run: pip install pyxlsb") from e
    best=(-1,None,None,"")
    exp={hkey(x) for x in expected}
    with open_workbook(str(path)) as wb:
        for name in wb.sheets:
            with wb.get_sheet(name) as sh:rows=[[c.v for c in row] for row in sh.rows()]
            for i,row in enumerate(rows[:20]):
                score=len({hkey(x) for x in row if txt(x)} & exp)
                if score>best[0]:best=(score,row,rows[i+1:],name)
    if best[0]<=0:raise ValueError(f"Expected columns were not found in {path.name}.")
    heads=[txt(x) for x in best[1]];w=len(heads)
    return Table(heads,[(r+[None]*w)[:w] for r in best[2]],best[3])

def read_table(path,expected):
    ext=path.suffix.lower()
    if ext==".csv":return read_csv(path,expected)
    if ext in (".xlsx",".xlsm"):return read_excel(path,expected)
    if ext==".xlsb":return read_xlsb(path,expected)
    raise ValueError(f"Unsupported type for {path.name}. Use .xlsx, .xlsm, .xlsb or .csv.")

def require(table,names,label):
    missing=[x for x in names if hkey(x) not in table.cols]
    if missing:raise ValueError(f"{label} is missing: {', '.join(missing)}")

def build_indexes(table):
    require(table,["hardware_type","hardware_model","manufacturer","end_of_life_date","end_of_support_date"],"IS NW Catalog")
    c=table.cols; recs=[]; exact=defaultdict(list); types=defaultdict(list); partial=defaultdict(list)
    for sr,row in enumerate(table.rows,2):
        ht=txt(row[c[hkey('hardware_type')]]); hm=txt(row[c[hkey('hardware_model')]])
        if not ht and not hm:continue
        rec=CatalogRecord(ht,hm,txt(row[c[hkey('manufacturer')]]),row[c[hkey('end_of_life_date')]],row[c[hkey('end_of_support_date')]],sr,compact(ht),compact(hm),tuple(tokens(ht)),tuple(tokens(hm)))
        rid=len(recs);recs.append(rec)
        if rec.model_key:exact[rec.model_key].append(rid)
        if rec.type_key:types[rec.type_key].append(rid)
        for token in set(rec.type_tokens+rec.model_tokens):partial[token].append(rid)
    if not recs:raise ValueError("IS NW Catalog has no usable records.")
    return Indexes(tuple(recs),{k:tuple(v) for k,v in exact.items()},{k:tuple(v) for k,v in types.items()},{k:tuple(v) for k,v in partial.items()})

def score(q,r):
    qk=compact(q);qt=tokens(q)
    if not qk or not r.type_key:return 0
    if r.type_key==qk:s=900
    elif r.type_key in qk:s=700+min(len(r.type_key),100)
    elif qk in r.type_key:s=550+min(len(qk),100)
    else:
        shared=set(qt)&set(r.type_tokens)
        if not shared:return 0
        s=120*len(shared)
    s+=250*len(set(qt)&set(r.model_tokens))
    if r.model_key and len(r.model_key)>=2 and r.model_key in qk:s+=350
    for t in qt:
        if t in r.model_key:s+=80+(100 if qk.endswith(t) and r.model_key.endswith(t) else 0)
    return s

def match_model(value,idx):
    raw=txt(value);key=compact(raw)
    if not key:return Match(None,"NO MATCH - BLANK MODEL")
    ids=idx.exact_model_index.get(key,())
    if ids:
        rs=[idx.records[i] for i in ids];outputs={(r.hardware_type.lower(),r.hardware_model.lower(),r.manufacturer.lower()) for r in rs}
        return Match(rs[0],"EXACT MATCH",10000) if len(outputs)==1 else Match(None,"NO MATCH - AMBIGUOUS EXACT",10000,True)
    cand=set(idx.hardware_type_index.get(key,()))
    for t in tokens(raw):cand.update(idx.hardware_type_index.get(compact(t),()));cand.update(idx.partial_token_index.get(t,()))
    for tk,record_ids in idx.hardware_type_index.items():
        if len(tk)>=3 and (tk in key or key in tk):cand.update(record_ids)
    ranked=[(score(raw,idx.records[i]),idx.records[i]) for i in cand];ranked=[x for x in ranked if x[0]>0]
    if not ranked:return Match(None,"NO MATCH")
    ranked.sort(key=lambda x:(-x[0],x[1].source_row));top=ranked[0][0];rs=[r for s,r in ranked if s==top]
    outputs={(r.hardware_type.lower(),r.hardware_model.lower(),r.manufacturer.lower()) for r in rs}
    return Match(rs[0],"PARTIAL MATCH",top) if len(outputs)==1 else Match(None,"NO MATCH - AMBIGUOUS PARTIAL",top,True)

def as_date(v):
    if v is None or txt(v)=="":return None
    if isinstance(v,datetime):return v.date()
    if isinstance(v,date):return v
    if isinstance(v,(int,float)):
        try:return from_excel(v).date()
        except Exception:return None
    s=txt(v)
    for fmt in ("%Y-%m-%d","%Y/%m/%d","%m/%d/%Y","%d/%m/%Y","%d-%m-%Y","%m-%d-%Y","%d-%b-%Y","%d %b %Y","%b %d, %Y"):
        try:return datetime.strptime(s,fmt).date()
        except ValueError:pass
    try:return datetime.fromisoformat(s.replace("Z","+00:00")).date()
    except ValueError:return None

def date_status(cmdb_eol,cmdb_eosl,cat_eol,cat_eosl,matched):
    if not matched:return "NOT APPLICABLE - NO CATALOG MATCH"
    ce,cs,ke,ks=map(as_date,(cmdb_eol,cmdb_eosl,cat_eol,cat_eosl))
    eol="MATCH" if ce and ke and ce==ke else ("BOTH BLANK" if not ce and not ke else "MISMATCH")
    eos="MATCH" if cs and ks and cs==ks else ("BOTH BLANK" if not cs and not ks else "MISMATCH")
    if eol=="MATCH" and eos=="MATCH":return "BOTH DATES MATCH"
    if eol=="BOTH BLANK" and eos=="BOTH BLANK":return "BOTH DATE PAIRS BLANK"
    return f"EOL: {eol}; EOSL: {eos}"

def write_template(path):path.write_bytes(base64.b64decode(TEMPLATE_B64))
def ensure_column(ws,name,outcols):
    key=hkey(name)
    if key in outcols:return outcols[key]
    col=ws.max_column+1;cell=ws.cell(1,col,name)
    if col>1:
        src=ws.cell(1,col-1);cell.font=src.font.copy();cell.fill=src.fill.copy();cell.border=src.border.copy();cell.alignment=src.alignment.copy()
    outcols[key]=col;return col

def create_output(catalog_path,cmdb_path,output_path):
    catalog=read_table(catalog_path,["hardware_type","hardware_model","manufacturer","end_of_life_date","end_of_support_date"])
    cmdb=read_table(cmdb_path,["Model number"]);require(cmdb,["Model number"],"NW CMDB Report")
    idx=build_indexes(catalog)
    with NamedTemporaryFile(suffix='.xlsx',delete=False) as f:temp=Path(f.name)
    try:
        write_template(temp);wb=load_workbook(temp);ws=wb.active
        if ws.max_row>1:ws.delete_rows(2,ws.max_row-1)
        outcols={hkey(c.value):c.column for c in ws[1] if hkey(c.value)}
        for name in ("Catalog Match Status","CMDB and Catalog Date Match Status"):ensure_column(ws,name,outcols)
        cc=cmdb.cols;shared=[(i,outcols[k]) for k,i in cc.items() if k in outcols]
        model_i=cc[hkey('Model number')]
        cmdb_eol_i=cc.get(hkey('HW_End_Of_Life_Date'));cmdb_eos_i=cc.get(hkey('HW_End_Of_Support_Date'))
        stats={"processed":0,"exact":0,"partial":0,"no_match":0,"date_match":0,"date_mismatch":0}
        for orow,row in enumerate((r for r in cmdb.rows if any(txt(v) for v in r)),2):
            stats['processed']+=1
            for si,di in shared:ws.cell(orow,di,row[si])
            m=match_model(row[model_i],idx);ws.cell(orow,outcols[hkey('Catalog Match Status')],m.status)
            ce=row[cmdb_eol_i] if cmdb_eol_i is not None else None;cs=row[cmdb_eos_i] if cmdb_eos_i is not None else None
            if m.record:
                r=m.record
                ws.cell(orow,outcols[hkey('Catalog Manufacturer')],r.manufacturer)
                ws.cell(orow,outcols[hkey('catalog hardware_type')],r.hardware_type)
                ws.cell(orow,outcols[hkey('Catalog hardware_model')],r.hardware_model)
                ws.cell(orow,outcols[hkey('Catalog end_of_life_date')],as_date(r.end_of_life_date))
                ws.cell(orow,outcols[hkey('Catalog end_of_support_date')],as_date(r.end_of_support_date))
                stats['exact' if m.status=='EXACT MATCH' else 'partial']+=1
            else:stats['no_match']+=1
            ds=date_status(ce,cs,m.record.end_of_life_date if m.record else None,m.record.end_of_support_date if m.record else None,bool(m.record))
            ws.cell(orow,outcols[hkey('CMDB and Catalog Date Match Status')],ds)
            if ds=='BOTH DATES MATCH':stats['date_match']+=1
            elif m.record:stats['date_mismatch']+=1
            # Apply actual Excel short-date values/format to all four output date columns.
            for name in ('HW_End_Of_Life_Date','HW_End_Of_Support_Date','Catalog end_of_life_date','Catalog end_of_support_date'):
                col=outcols.get(hkey(name))
                if col:
                    cell=ws.cell(orow,col);d=as_date(cell.value)
                    if d:cell.value=d;cell.number_format=SHORT_DATE_FORMAT
        ws.auto_filter.ref=f"A1:{ws.cell(1,ws.max_column).coordinate}";ws.freeze_panes='A2'
        output_path=output_path.with_suffix('.xlsx');wb.save(output_path);wb.close();return stats
    finally:temp.unlink(missing_ok=True)

def launch_gui():
    import tkinter as tk
    from tkinter import ttk,filedialog,messagebox
    class App(tk.Tk):
        def __init__(self):
            super().__init__();self.title(APP_TITLE);self.geometry('900x340');self.minsize(760,320)
            self.cat=tk.StringVar();self.cmdb=tk.StringVar();self.status=tk.StringVar(value='Select the two input files.')
            f=ttk.Frame(self,padding=20);f.pack(fill='both',expand=True);f.columnconfigure(1,weight=1)
            ttk.Label(f,text=APP_TITLE,font=('Segoe UI',16,'bold')).grid(row=0,column=0,columnspan=3,sticky='w',pady=(0,18))
            for rn,(label,var,cmd) in enumerate((('1. IS NW Catalog file',self.cat,self.pick_cat),('2. NW CMDB report',self.cmdb,self.pick_cmdb)),1):
                ttk.Label(f,text=label).grid(row=rn,column=0,sticky='w',padx=(0,12),pady=8);ttk.Entry(f,textvariable=var).grid(row=rn,column=1,sticky='ew');ttk.Button(f,text='Browse',command=cmd).grid(row=rn,column=2,padx=(12,0))
            self.pb=ttk.Progressbar(f,mode='indeterminate');self.pb.grid(row=3,column=0,columnspan=3,sticky='ew',pady=(20,8));ttk.Label(f,textvariable=self.status,wraplength=840).grid(row=4,column=0,columnspan=3,sticky='w');ttk.Button(f,text='Run Catalog Match',command=self.run,width=24).grid(row=5,column=0,columnspan=3,pady=20)
        def types(self):return [('Supported files','*.xlsx *.xlsm *.xlsb *.csv'),('Excel Binary','*.xlsb'),('Excel','*.xlsx *.xlsm'),('CSV','*.csv')]
        def pick_cat(self):
            p=filedialog.askopenfilename(title='Select IS NW Catalog',filetypes=self.types());self.cat.set(p or self.cat.get())
        def pick_cmdb(self):
            p=filedialog.askopenfilename(title='Select NW CMDB Report',filetypes=self.types());self.cmdb.set(p or self.cmdb.get())
        def run(self):
            cp,mp=Path(self.cat.get()),Path(self.cmdb.get())
            if not cp.is_file() or not mp.is_file():messagebox.showerror(APP_TITLE,'Select both valid input files.');return
            op=filedialog.asksaveasfilename(title='Save output',initialdir=str(mp.parent),initialfile=f'{mp.stem}_NW_Catalog_Match.xlsx',defaultextension='.xlsx',filetypes=[('Excel','.xlsx')])
            if not op:return
            self.pb.start(10);self.status.set('Matching records and comparing lifecycle dates...');self.update_idletasks()
            try:
                s=create_output(cp,mp,Path(op));msg=f"Completed | Processed: {s['processed']} | Exact: {s['exact']} | Partial: {s['partial']} | No match: {s['no_match']} | Both dates match: {s['date_match']} | Date review: {s['date_mismatch']}";self.status.set(msg);messagebox.showinfo(APP_TITLE,msg+'\n\n'+op)
            except Exception as e:self.status.set('Failed');messagebox.showerror(APP_TITLE,f'{e}\n\n{traceback.format_exc(limit=6)}')
            finally:self.pb.stop()
    App().mainloop()
if __name__=='__main__':launch_gui()
