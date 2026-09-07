"""WPP Enterprise CMDB and Inventory Services Gap Analysis V5.2

Standalone Tkinter application. Requires pandas and openpyxl. Optional: pyxlsb.
The application is deliberately layered: engines below have no Tkinter dependency and
can be imported by an agent, FastAPI service, Azure Function, or test runner.
"""
from __future__ import annotations
import os, re, json, sqlite3, hashlib, ipaddress, traceback, threading
from pathlib import Path
from datetime import datetime, date
from typing import Any, Dict, List, Optional, Tuple
import pandas as pd
from openpyxl import load_workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter
import tkinter as tk
from tkinter import ttk, filedialog, messagebox

APP_NAME = "WPP Enterprise CMDB and Inventory Services Gap Analysis V5.2"
APP_DIR = Path(__file__).resolve().parent
REGISTRY_PATH = APP_DIR / "v5_2_resource_registry.json"
LEARNING_DB = APP_DIR / "v5_2_flag_intelligence.sqlite"
DATE_LABEL = datetime.now().strftime("%d-%b-%Y")
DISPLAY_DATE = datetime.now().strftime("%d.%b.%Y").upper()
FILE_TYPES = [("Supported files", "*.csv *.xlsx *.xlsb"), ("CSV", "*.csv"), ("Excel", "*.xlsx"), ("Excel Binary", "*.xlsb"), ("All files", "*.*")]

PERSISTENT_RESOURCES = {
    "Bulk Load Template": "bulk_load_template",
    "IS Product Catalog Operating System": "catalog_os",
    "IS Product Catalog Network": "catalog_network",
    "IS Product Catalog Server": "catalog_server",
}
SESSION_RESOURCES = ["IS Operating System Report", "NW CMDB Report", "Server CMDB Report", "Field List and Mapping"]
ALL_RESOURCES = SESSION_RESOURCES + list(PERSISTENT_RESOURCES)

# Microsoft-published lifecycle dates captured from Windows Server release information.
# V5.2 treats versions as supported until extended support end. Refresh through code/config
# when Microsoft publishes changed lifecycle information.
WINDOWS_SERVER_SUPPORT = {
    "2025": ("10.0.26100", date(2034, 11, 14)),
    "2022": ("10.0.20348", date(2031, 10, 14)),
    "2019": ("10.0.17763", date(2029, 1, 9)),
    "2016": ("10.0.14393", date(2027, 1, 12)),
    "2012 r2": ("6.3.9600", date(2023, 10, 10)),
    "2012": ("6.2.9200", date(2023, 10, 10)),
    "2008 r2": ("6.1.7601", date(2020, 1, 14)),
}

# ---------------------------- reusable core utilities ----------------------------
def clean(v: Any) -> str:
    return "" if pd.isna(v) else re.sub(r"\s+", " ", str(v).strip())

def hkey(v: Any) -> str:
    return re.sub(r"[\s_.-]+", "", clean(v).lower())

def pkey(v: Any) -> str:
    return re.sub(r"[^a-z0-9]", "", clean(v).lower())

def vkey(v: Any) -> str:
    s = clean(v).lower()
    return re.sub(r"\d+", lambda m: str(int(m.group())), s) if s else ""

def lifecycle(v: Any) -> str:
    return {"operational":"operational", "endoflife":"end of life", "eol":"end of life",
            "production":"production", "installed":"installed", "deploy":"deploy", "design":"design"}.get(pkey(v), clean(v).lower())

def serial_key(v: Any) -> str:
    return re.sub(r"\s+", "", clean(v).lower())

def normalize_hostname(v: Any) -> str:
    return clean(v).replace("/", "_")

def normalize_fqdn(v: Any) -> str:
    return clean(v).replace("/", "_").lower().strip(".")

def valid_fqdn(v: Any) -> str:
    s = normalize_fqdn(v)
    invalid = {"", "domainnamenotconfigured", "domainnamenotapplicable", "notapplicable", "na", "workgroup", "workgroupserver"}
    return s if pkey(s) not in invalid and "." in s and " " not in s else ""

def normalize_ipv4(v: Any) -> Tuple[str, str, str]:
    raw = clean(v)
    if pkey(raw) in {"stack", "standby", "na", "notapplicable", "notconfigured", "none"}:
        return "", "Placeholder removed", ""
    found = []
    for token in re.findall(r"(?<!\d)(?:\d{1,3}\.){3}\d{1,3}(?!\d)", raw):
        try:
            ipaddress.IPv4Address(token); found.append(token)
        except ValueError:
            pass
    if not found:
        return "", ("No valid IPv4 found - review required" if raw else ""), ""
    if len(found) > 1:
        return found[0], "Multiple IP found, kept the first one, review required", ", ".join(found[1:])
    return found[0], ("IP normalized" if found[0] != raw else ""), ""

def split_product_version(v: Any) -> Tuple[str, str]:
    raw = clean(v)
    matches = list(re.finditer(r"(?<!\d)(\d+(?:\.\d+)+(?:\([^)]+\))*[A-Za-z0-9_.()\-]*)", raw))
    if not matches:
        return raw, ""
    m = matches[-1]
    name = clean(raw[:m.start()])
    return (name, clean(m.group(1))) if name else (raw, "")

def read_tabular(path: str, keep_metadata: bool=False):
    ext = Path(path).suffix.lower()
    if ext == ".csv":
        df = pd.read_csv(path, dtype=str, keep_default_na=False, encoding_errors="replace")
    elif ext == ".xlsx":
        df = pd.read_excel(path, dtype=str, keep_default_na=False, engine="openpyxl")
    elif ext == ".xlsb":
        df = pd.read_excel(path, dtype=str, keep_default_na=False, engine="pyxlsb")
    else:
        raise ValueError(f"Unsupported file type: {ext}")
    df.columns = [clean(c) for c in df.columns]
    metadata = None
    if len(df):
        terms = {pkey(x) for x in df.iloc[0].tolist() if clean(x)}
        if terms & {"mandatory", "optional", "reference", "required", "recommended"}:
            metadata = df.iloc[0].map(clean)
            df = df.iloc[1:].reset_index(drop=True)
    for c in df.columns:
        df[c] = df[c].map(clean)
    return (df, metadata) if keep_metadata else df

def find_col(df: pd.DataFrame, *aliases: str, required=True) -> Optional[str]:
    index = {hkey(c): c for c in df.columns}
    for alias in aliases:
        if hkey(alias) in index:
            return index[hkey(alias)]
    if required:
        raise ValueError("Required field not found. Expected one of: " + ", ".join(aliases))
    return None

def next_versioned_path(folder: str, stem: str, extension: str) -> str:
    safe_ext = extension.lstrip(".")
    pattern = re.compile(re.escape(stem) + r" - " + re.escape(DATE_LABEL) + r" - V(\d+)\." + re.escape(safe_ext) + r"$", re.I)
    versions = []
    for p in Path(folder).glob(f"{stem} - {DATE_LABEL} - V*.{safe_ext}"):
        m = pattern.match(p.name)
        if m:
            versions.append(int(m.group(1)))
    return str(Path(folder) / f"{stem} - {DATE_LABEL} - V{max(versions, default=0)+1}.{safe_ext}")

def write_workbook(path: str, sheets: Dict[str, pd.DataFrame]):
    with pd.ExcelWriter(path, engine="openpyxl") as writer:
        for name, df in sheets.items():
            data = df if df is not None and len(df) else pd.DataFrame(columns=["No records"])
            data.to_excel(writer, index=False, sheet_name=re.sub(r"[\[\]:*?/\\]", "-", name)[:31])
    wb = load_workbook(path)
    navy = PatternFill("solid", fgColor="1F4E78"); white = Font(color="FFFFFF", bold=True)
    thin = Side(style="thin", color="D9E2F3"); border = Border(left=thin,right=thin,top=thin,bottom=thin)
    for ws in wb.worksheets:
        ws.freeze_panes = "A2"; ws.sheet_view.showGridLines = False
        if ws.max_row: ws.auto_filter.ref = ws.dimensions
        for c in ws[1]:
            c.fill=navy; c.font=white; c.alignment=Alignment(horizontal="center",vertical="center",wrap_text=True); c.border=border
        for row in ws.iter_rows(min_row=2):
            for c in row:
                c.border=border; c.alignment=Alignment(vertical="top",wrap_text=True)
        for i,cells in enumerate(ws.columns,1):
            ws.column_dimensions[get_column_letter(i)].width=min(max(max(len("" if c.value is None else str(c.value)) for c in cells)+2,12),55)
    wb.save(path)

# -------------------------- persistent resource manager --------------------------
class ResourceRegistry:
    def __init__(self, path=REGISTRY_PATH):
        self.path = Path(path); self.data = self._load()
    def _load(self):
        if not self.path.exists(): return {"resources": {}}
        try: return json.loads(self.path.read_text(encoding="utf-8"))
        except Exception: return {"resources": {}}
    def save(self):
        self.path.write_text(json.dumps(self.data, indent=2), encoding="utf-8")
    def register(self, resource_id: str, path: str):
        digest = hashlib.sha256(Path(path).read_bytes()).hexdigest()
        now = datetime.now().isoformat(timespec="seconds")
        previous = self.data["resources"].get(resource_id, {})
        loaded_on = previous.get("loaded_on", now) if previous.get("sha256") == digest else now
        self.data["resources"][resource_id] = {"path":str(Path(path).resolve()), "file_name":Path(path).name,
            "loaded_on":loaded_on, "last_used":previous.get("last_used", ""), "sha256":digest}
        self.save()
    def get(self, resource_id: str):
        item = self.data["resources"].get(resource_id)
        if item and Path(item.get("path","")).exists(): return item
        return None
    def mark_used(self, resource_id: str):
        item=self.data["resources"].get(resource_id)
        if item: item["last_used"]=datetime.now().isoformat(timespec="seconds"); self.save()

# ------------------------- enterprise field intelligence -------------------------
class FieldIntelligence:
    SOURCES = ["IS Operating System Report","NW CMDB Report","Server CMDB Report","IS Server Report","IS Network Report","IS Product Catalog Operating System","IS Product Catalog Network","IS Product Catalog Server"]
    DEFAULT = {
        "fully_qualified_hostname":{"IS Operating System Report":"fully_qualified_hostname","NW CMDB Report":"Fully qualified domain name","Server CMDB Report":"Fully qualified domain name"},
        "hostname":{"IS Operating System Report":"hostname","NW CMDB Report":"Configuration Item","Server CMDB Report":"Name"},
        "domain_name":{"IS Operating System Report":"domain_name","Server CMDB Report":"Domain"},
        "os_lifecycle_status":{"IS Operating System Report":"os_lifecycle_status","NW CMDB Report":"Life Cycle Stage","Server CMDB Report":"Life Cycle Stage"},
        "os_parent_serial_number":{"IS Operating System Report":"os_parent_serial_number","NW CMDB Report":"Serial number","Server CMDB Report":"Serial number"},
        "os_parent_manufacturer":{"IS Operating System Report":"os_parent_manufacturer","NW CMDB Report":"Manufacturer","Server CMDB Report":"Manufacturer"},
        "os_parent_opaque_id":{"IS Operating System Report":"os_parent_opaque_id","IS Server Report":"server_opaque_id","IS Network Report":"network_opaque_id"},
        "operating_system_name":{"IS Operating System Report":"operating_system_name","Server CMDB Report":"Operating System","IS Product Catalog Operating System":"software_product_name"},
        "operating_system_version":{"IS Operating System Report":"operating_system_version","Server CMDB Report":"OS Version","IS Product Catalog Operating System":"software_version"},
        "operating_system_provider":{"IS Operating System Report":"operating_system_provider","IS Product Catalog Operating System":"software_provider"},
        "ip_address":{"IS Operating System Report":"ip_address","NW CMDB Report":"IP Address","Server CMDB Report":"IP Address"},
        "class":{"NW CMDB Report":"Class","Server CMDB Report":"Class"},
        "category":{"NW CMDB Report":"Category","Server CMDB Report":"Category"},
        "subcategory":{"NW CMDB Report":"Subcategory","Server CMDB Report":"Subcategory","IS Product Catalog Operating System":"subcategory"},
    }
    def __init__(self): self.mapping=json.loads(json.dumps(self.DEFAULT)); self.field_lists={}
    def load_workbook(self,path: str):
        book=pd.ExcelFile(path,engine="openpyxl")
        if "Field List" in book.sheet_names:
            raw=pd.read_excel(path,sheet_name="Field List",dtype=str,keep_default_na=False,engine="openpyxl")
            self.field_lists={c:[clean(x) for x in raw[c] if clean(x)] for c in raw.columns}
        if "Field Mapping" in book.sheet_names:
            raw=pd.read_excel(path,sheet_name="Field Mapping",dtype=str,keep_default_na=False,engine="openpyxl")
            cols=list(raw.columns)
            if cols:
                logical_source=cols[0]
                for _,r in raw.iterrows():
                    logical=clean(r.get(logical_source,""))
                    if not logical: continue
                    self.mapping.setdefault(logical,{})
                    for source in cols:
                        value=clean(r.get(source,""))
                        if value:self.mapping[logical][source]=value
    def field(self,logical: str,source: str,df: pd.DataFrame,required=True):
        aliases=[]
        mapped=self.mapping.get(logical,{}).get(source)
        if mapped:aliases.append(mapped)
        aliases.extend([logical,logical.replace("_"," ")])
        return find_col(df,*aliases,required=required)

# --------------------------- catalogue/reconciliation ---------------------------
def catalog_index(cat: pd.DataFrame, fi: FieldIntelligence):
    pc=fi.field("operating_system_name","IS Product Catalog Operating System",cat);vc=fi.field("operating_system_version","IS Product Catalog Operating System",cat);pr=fi.field("operating_system_provider","IS Product Catalog Operating System",cat,False)
    exact={};base={}
    for _,r in cat.iterrows():
        p={"name":clean(r[pc]),"version":clean(r[vc]),"provider":clean(r[pr]) if pr else ""}
        exact.setdefault(pkey(p["name"])+"|"+vkey(p["version"]),p)
        base.setdefault(pkey(p["name"])+"|"+re.sub(r"(?<=\d)[a-z]+$","",vkey(p["version"])),p)
    return exact,base

def reconcile_nw(cm,isr,cat,fi,progress=lambda m,p:None):
    progress("Resolving NW fields through Field Intelligence",10)
    logical=["hostname","class","os_lifecycle_status","ip_address","fully_qualified_hostname"]
    ci=fi.field("hostname","NW CMDB Report",cm); cls=fi.field("class","NW CMDB Report",cm); life_c=fi.field("os_lifecycle_status","NW CMDB Report",cm); ip=fi.field("ip_address","NW CMDB Report",cm); fq=fi.field("fully_qualified_hostname","NW CMDB Report",cm,False)
    serial=fi.field("os_parent_serial_number","NW CMDB Report",cm); manuf=fi.field("os_parent_manufacturer","NW CMDB Report",cm,False)
    firmware=find_col(cm,"Firmware version","Firmware Version")
    out=pd.DataFrame({"Configuration Item":cm[ci],"Class":cm[cls],"Life Cycle Stage":cm[life_c],"Life Cycle Stage Status":cm[find_col(cm,"Life Cycle Stage Status",required=False)] if find_col(cm,"Life Cycle Stage Status",required=False) else "","Manufacturer":cm[manuf] if manuf else "","Model ID":cm[find_col(cm,"Model ID",required=False)] if find_col(cm,"Model ID",required=False) else "","Model.Name":cm[find_col(cm,"Model.Name","Model Name",required=False)] if find_col(cm,"Model.Name","Model Name",required=False) else "","Model number":cm[find_col(cm,"Model number",required=False)] if find_col(cm,"Model number",required=False) else "","Serial number":cm[serial],"IP Address":cm[ip],"Firmware version":cm[firmware],"Fully qualified domain name":cm[fq] if fq else ""})
    parsed=out["Firmware version"].map(split_product_version).apply(pd.Series);parsed.columns=["Parsed Firmware Name","Parsed Firmware Version"];out=pd.concat([out,parsed],axis=1)
    exact,base=catalog_index(cat,fi);ih=fi.field("hostname","IS Operating System Report",isr);isn=fi.field("os_parent_serial_number","IS Operating System Report",isr);il=fi.field("os_lifecycle_status","IS Operating System Report",isr);io=find_col(isr,"os_opaque_id","OS Opaque ID",required=False);ion=fi.field("operating_system_name","IS Operating System Report",isr);iov=fi.field("operating_system_version","IS Operating System Report",isr);iop=fi.field("operating_system_provider","IS Operating System Report",isr,False)
    hm={};sm={}
    for i,r in isr.iterrows():
        if clean(r[ih]):hm.setdefault(clean(r[ih]).lower(),i)
        if serial_key(r[isn]):sm.setdefault(serial_key(r[isn]),i)
    rows=[];progress("Reconciling NW CMDB and IS",45)
    for _,r in out.iterrows():
        name=clean(r["Configuration Item"]);idx=next((hm[x] for x in [name.lower(),name.replace("/","_").lower(),name.replace("/","-").lower()] if x in hm),None);mode="Hostname Match" if idx is not None else ""
        if idx is None and serial_key(r["Serial number"]) in sm:idx=sm[serial_key(r["Serial number"])];mode="Serial Match Different Hostname"
        present=idx is not None;ir=isr.iloc[idx] if present else None;k=pkey(r["Parsed Firmware Name"])+"|"+vkey(r["Parsed Firmware Version"]);hit=exact.get(k);status="Match"
        if not hit:hit=base.get(pkey(r["Parsed Firmware Name"])+"|"+re.sub(r"(?<=\d)[a-z]+$","",vkey(r["Parsed Firmware Version"])));status="Closest Match Proposed" if hit else "No Match";hit=hit or {}
        osmatch=present and hit and pkey(ir[ion])==pkey(hit.get("name")) and vkey(ir[iov])==vkey(hit.get("version"));actions=[]
        if status=="No Match":actions.append("Catalogue update required")
        if present and hit and not osmatch:actions.append("Inventory OS update required")
        if lifecycle(r["Life Cycle Stage"])=="operational" and not present:actions.append("Load To IS")
        if lifecycle(r["Life Cycle Stage"])=="end of life" and present and lifecycle(ir[il]) in {"production","installed"}:actions.append("Retire From IS")
        d=r.to_dict();d.update({"Present in IS ?":"Yes - "+mode if present else "No","Host name in IS":clean(ir[ih]) if present else "","os_opaque_id":clean(ir[io]) if present and io else "","os_parent_serial_number":clean(ir[isn]) if present else "","Lifecycle Status in IS":clean(ir[il]) if present else "","Current operating_system_provider":clean(ir[iop]) if present and iop else "","Current operating_system_name":clean(ir[ion]) if present else "","Current operating_system_version":clean(ir[iov]) if present else "","Catalogue Match?":status,"Required operating_system_provider":hit.get("provider",""),"Required operating_system_name":hit.get("name",""),"Required operating_system_version":hit.get("version",""),"Action":"; ".join(actions) or "No Action"});rows.append(d)
    progress("NW reconciliation complete",100);return pd.DataFrame(rows)

def nw_summary(r):
    return pd.DataFrame([("NW CMDB records analyzed",len(r)),("Records present in IS",int(r["Present in IS ?"].astype(str).str.startswith("Yes").sum())),("Records not present in IS",int((r["Present in IS ?"]=="No").sum())),("Catalogue exact matches",int((r["Catalogue Match?"]=="Match").sum())),("Catalogue closest matches proposed",int((r["Catalogue Match?"]=="Closest Match Proposed").sum())),("Catalogue updates required",int(r["Action"].str.contains("Catalogue update required",na=False).sum())),("Inventory OS updates required",int(r["Action"].str.contains("Inventory OS update required",na=False).sum())),("Load To IS records",int(r["Action"].str.contains("Load To IS",na=False).sum())),("Retire From IS records",int(r["Action"].str.contains("Retire From IS",na=False).sum()))],columns=["Metric","Value"])

# ------------------------------- policy intelligence -------------------------------
def windows_supported(name,version):
    text=(clean(name)+" "+clean(version)).lower()
    for label,(build,end) in WINDOWS_SERVER_SUPPORT.items():
        if label in text or (build and build in text): return date.today() <= end,label,end
    return False,"Unknown",None

def is_server_context(context):
    text=" ".join(clean(context.get(k,"")) for k in ["source_type","subcategory","parent_system_type","image_purpose","os_name","class","category"]).lower()
    return any(x in text for x in ["server","windows","linux","rhel","ubuntu","aix","solaris","esx","hypervisor"])

def apply_policy(row: Dict[str,str],context: Dict[str,str]):
    applied=[]
    def put(field,value,rule): row[field]=str(value);applied.append({"Field":field,"Value":str(value),"Rule":rule})
    put("is_disaster_recovery_contracted","N","P001 Global")
    put("is_designated_for_compliance_reporting","Y","P002 Global")
    put("bac_id","WPP Europe","P003 Global")
    put("bam_id","BAM000415E","P004 Global")
    sub=clean(context.get("subcategory","")).lower();parent=clean(context.get("parent_system_type","")).lower();osname=clean(context.get("os_name","")).lower();source=clean(context.get("source_type","")).lower()
    supported,release,end=windows_supported(context.get("os_name",""),context.get("os_version",""))
    if source=="network" or parent in {"network","appliance"} or "esx" in osname:
        put("is_anti_virus_required","N","AV002/AV003/AV004")
    elif sub=="windows" and supported:
        put("is_anti_virus_required","Y","AV001 Supported Windows")
        put("is_patching_required","Y","PATCH001 Supported Windows")
        put("is_health_check_required","Y","HC001 Supported Windows")
    if row.get("is_health_check_required")=="Y":
        put("health_check_interval","12","HC002")
        put("health_check_unit","month","HC002")
    if is_server_context(context):
        put("is_employment_verification_required","Y","QEV001")
        put("employment_verification_interval","3","QEV001")
        put("employment_verification_unit","month","QEV001")
        put("is_continued_business_need_required","Y","CBN001")
        sox="sox" in clean(context.get("attribute_field_value","")).lower()
        put("continued_business_need_interval","3" if sox else "12","CBN002 SOX" if sox else "CBN001 Annual")
        put("continued_business_need_unit","month","CBN001/CBN002")
    return applied

# ----------------------------- learned flag intelligence -----------------------------
class FlagLearning:
    def __init__(self,path=LEARNING_DB):
        self.conn=sqlite3.connect(path);self.conn.executescript("CREATE TABLE IF NOT EXISTS reports(hash TEXT PRIMARY KEY,loaded TEXT,name TEXT);CREATE TABLE IF NOT EXISTS observations(field TEXT,value TEXT,priority TEXT,cohort TEXT);");self.conn.commit()
    def ingest(self,path):
        digest=hashlib.sha256(Path(path).read_bytes()).hexdigest()
        if self.conn.execute("SELECT 1 FROM reports WHERE hash=?",(digest,)).fetchone():return
        df,meta=read_tabular(path,True);meta=meta if meta is not None else pd.Series("",index=df.columns)
        flags=[]
        for c in df.columns:
            values={clean(x).upper() for x in df[c] if clean(x)}
            if values and values.issubset({"Y","N"}):flags.append((c,"Priority 1" if clean(meta.get(c,"")).lower()=="mandatory" else "Priority 2"))
        for _,r in df.iterrows():
            cohort="|".join([pkey(r.get(x,"")) for x in ["operating_system_subcategory","parent_system_type","image_purpose","os_lifecycle_status","operating_system_name","business_criticality_device_grouping","os_inventory_item_owned_by","os_inventory_item_managed_by"]])
            for field,priority in flags:
                value=clean(r[field]).upper()
                if value in {"Y","N"}:self.conn.execute("INSERT INTO observations VALUES(?,?,?,?)",(hkey(field),value,priority,cohort))
        self.conn.execute("INSERT INTO reports VALUES(?,?,?)",(digest,datetime.now().isoformat(timespec="seconds"),Path(path).name));self.conn.commit()
    def predict(self,field,priority,cohort):
        parts=cohort.split("|")
        for n in range(len(parts),0,-1):
            prefix="|".join(parts[:n]);vals=[x[0] for x in self.conn.execute("SELECT value FROM observations WHERE field=? AND priority=? AND cohort LIKE ?",(hkey(field),priority,prefix+"%")).fetchall()]
            if vals:
                y=vals.count("Y");z=vals.count("N");support=y+z;agreement=max(y,z)/support;value="Y" if y>=z else "N"
                status="Auto-prefill" if agreement>=.90 and support>=20 else "Suggest for review" if agreement>=.75 else "No proposal"
                return value,agreement,support,status,n
        return "",0,0,"No evidence",0
    def close(self):self.conn.close()

# -------------------------------- bulk load engine --------------------------------
def generate_bulk_load(rec,source_type,is_report,template,outdir,fmt,askfq,askcat,progress=lambda m,p:None):
    candidates=rec[(rec["Present in IS ?"]=="No")&(rec["Life Cycle Stage"].map(lifecycle)=="operational")&rec["Action"].str.contains("Load To IS",na=False)].copy()
    dedupe=candidates["Configuration Item"].map(pkey)+"|"+candidates["Life Cycle Stage"].map(lifecycle)+"|"+candidates["Serial number"].map(serial_key);bad=dedupe.duplicated(False);duplicate=candidates[bad];candidates=candidates[~bad]
    unmatched=candidates[candidates["Catalogue Match?"]=="No Match"];catalog_path=""
    if len(unmatched):
        if askcat(len(unmatched)):
            catalog_path=next_versioned_path(outdir,"Catalogue Update Required","xlsx");write_workbook(catalog_path,{"Catalogue Update Required":unmatched})
        candidates=candidates[candidates["Catalogue Match?"]!="No Match"]
    fq=candidates["Fully qualified domain name"].map(valid_fqdn);missing=fq.eq("");fq_excluded=pd.DataFrame()
    if missing.any():
        if askfq(int(missing.sum())):fq.loc[missing]=candidates.loc[missing,"Configuration Item"].map(lambda x:normalize_hostname(x).lower()+".nw.wpp.net")
        else:fq_excluded=candidates[missing];candidates=candidates[~missing];fq=fq[~missing]
    template_df,meta=read_tabular(template,True);headers=list(template_df.columns);extras=["CMDB Hostname","CMDB Serial Number","CMDB Lifecycle Stage"];rows=[];policy_audit=[];learning_audit=[];ip_audit=[]
    learner=FlagLearning();learner.ingest(is_report)
    flag_info=[]
    if meta is not None:
        isdf,_=read_tabular(is_report,True)
        for h in headers:
            source_col=next((c for c in isdf.columns if hkey(c)==hkey(h)),None)
            if source_col:
                values={clean(x).upper() for x in isdf[source_col] if clean(x)}
                if values and values.issubset({"Y","N"}):flag_info.append((h,"Priority 1" if clean(meta.get(h,"")).lower()=="mandatory" else "Priority 2"))
    for ix,r in candidates.iterrows():
        ip,msg,more=normalize_ipv4(r["IP Address"]);parent="Appliance" if pkey(r["Class"]) in {"wirelessaccesspoint","interfacecard"} else ("Network" if source_type=="network" else "OS")
        values={"hostname":normalize_hostname(r["Configuration Item"]),"fully_qualified_hostname":fq.loc[ix].lower(),"os_parent_serial_number":r["Serial number"],"os_lifecycle_status":"Production","parent_system_type":parent,"operating_system_provider":r["Required operating_system_provider"],"operating_system_name":r["Required operating_system_name"],"operating_system_version":r["Required operating_system_version"],"ip_address":ip}
        row={h:"" for h in headers};lookup={hkey(k):v for k,v in values.items()}
        for h in headers:
            if hkey(h) in lookup:row[h]=lookup[hkey(h)]
        context={"source_type":source_type,"subcategory":"Network" if source_type=="network" else ("Windows" if "windows" in clean(r["Required operating_system_name"]).lower() else "Server"),"parent_system_type":parent,"image_purpose":r["Class"],"os_name":r["Required operating_system_name"],"os_version":r["Required operating_system_version"],"class":r["Class"],"category":r.get("category",""),"attribute_field_value":r.get("Attribute Field Value",r.get("Attributes",""))}
        # Apply canonical policy fields, then resolve them to the actual template headers.
        policy_values={}
        applied=apply_policy(policy_values,context)
        header_lookup={hkey(h):h for h in headers}
        for canonical,value in policy_values.items():
            target=header_lookup.get(hkey(canonical))
            if target:
                row[target]=value
        for a in applied:
            a["Template Field Found"]="Y" if hkey(a["Field"]) in header_lookup else "N"
            policy_audit.append({"CMDB Hostname":r["Configuration Item"],**a})
        cohort="|".join([pkey(context["subcategory"]),pkey(context["parent_system_type"]),pkey(context["image_purpose"]),pkey("Production"),pkey(context["os_name"]),"","",""])
        for field,priority in flag_info:
            if clean(row.get(field,"")).upper() in {"Y","N"}:continue
            value,agreement,support,status,level=learner.predict(field,priority,cohort)
            applied_value="N"
            if status=="Auto-prefill" and value:row[field]=value;applied_value="Y"
            if value:learning_audit.append({"CMDB Hostname":r["Configuration Item"],"Field":field,"Priority":priority,"Proposed Value":value,"Agreement":agreement,"Support":support,"Decision":status,"Cohort Level":level,"Applied":applied_value})
        rows.append({extras[0]:r["Configuration Item"],extras[1]:r["Serial number"],extras[2]:r["Life Cycle Stage"],**row});ip_audit.append({"CMDB Hostname":r["Configuration Item"],"Original IP":r["IP Address"],"Selected IP":ip,"Additional IPs":more,"Message":msg})
    learner.close();metadata=["Reference"]*3+([clean(meta[h]) for h in headers] if meta is not None else [""]*len(headers));final=pd.concat([pd.DataFrame([metadata],columns=extras+headers),pd.DataFrame(rows,columns=extras+headers)],ignore_index=True)
    label="NW - OS Bulk Load" if source_type=="network" else "Server - OS Bulk Load";load_path=next_versioned_path(outdir,label,fmt);final.to_excel(load_path,index=False) if fmt=="xlsx" else final.to_csv(load_path,index=False)
    validation=next_versioned_path(outdir,label+" Validation","xlsx");write_workbook(validation,{"Policy Intelligence":pd.DataFrame(policy_audit),"Flag Intelligence":pd.DataFrame(learning_audit),"IP Validation":pd.DataFrame(ip_audit),"CMDB Correction Required":duplicate,"Missing FQDN Excluded":fq_excluded})
    return [x for x in [load_path,validation,catalog_path] if x]

# ------------------------------------ Tkinter UI ------------------------------------
class ScrollFrame(ttk.Frame):
    def __init__(self,parent):
        super().__init__(parent);self.canvas=tk.Canvas(self,bg="#F4F6F8",highlightthickness=0);self.bar=ttk.Scrollbar(self,command=self.canvas.yview);self.inner=ttk.Frame(self.canvas,padding=18);self.win=self.canvas.create_window((0,0),window=self.inner,anchor="nw");self.inner.bind("<Configure>",lambda e:self.canvas.configure(scrollregion=self.canvas.bbox("all")));self.canvas.bind("<Configure>",lambda e:self.canvas.itemconfigure(self.win,width=e.width));self.canvas.configure(yscrollcommand=self.bar.set);self.canvas.pack(side="left",fill="both",expand=True);self.bar.pack(side="right",fill="y")

class App(tk.Tk):
    def __init__(self):
        super().__init__();self.title(APP_NAME);self.geometry("1500x900");self.minsize(1180,720);self.registry=ResourceRegistry();self.fi=FieldIntelligence();self.paths={x:"" for x in ALL_RESOURCES};self.status={};self.results={};self.busy=False;self._build();self._restore()
    def _build(self):
        sf=ScrollFrame(self);sf.pack(fill="both",expand=True);root=sf.inner
        ttk.Label(root,text=APP_NAME,font=("Segoe UI",19,"bold"),foreground="#1F4E78").pack(anchor="w");ttk.Label(root,text="Persistent templates/catalogs, enterprise field intelligence, reconciliation, policy and learned flag intelligence").pack(anchor="w",pady=(0,12))
        pane=ttk.PanedWindow(root,orient="horizontal");pane.pack(fill="both",expand=True);left=ttk.Frame(pane);right=ttk.Frame(pane);pane.add(left,weight=3);pane.add(right,weight=4)
        load=ttk.LabelFrame(left,text="Resource Loader",padding=10);load.pack(fill="x")
        for i,name in enumerate(ALL_RESOURCES):
            ttk.Label(load,text=name,width=35).grid(row=i,column=0,sticky="w",pady=3);ttk.Button(load,text="Browse",command=lambda n=name:self.browse(n)).grid(row=i,column=1,padx=3);ttk.Button(load,text="Clear",command=lambda n=name:self.clear(n)).grid(row=i,column=2,padx=3);lab=tk.Label(load,text="Not loaded",fg="#9C0006",anchor="w");lab.grid(row=i,column=3,sticky="ew");self.status[name]=lab
        load.columnconfigure(3,weight=1)
        actions=ttk.LabelFrame(left,text="Dedicated Features",padding=10);actions.pack(fill="x",pady=10)
        for text,cmd in [("Reconcile NW CMDB and IS",self.reconcile_nw_click),("Reconcile Server CMDB and IS",self.reconcile_server_click),("Create NW - OS Bulk Load File",lambda:self.bulk_click("network")),("Create Server - OS Bulk Load File",lambda:self.bulk_click("server"))]:ttk.Button(actions,text=text,command=cmd,width=42).pack(anchor="w",pady=4)
        controls=ttk.Frame(left);controls.pack(fill="x");ttk.Button(controls,text="Clear Session Files",command=self.clear_session).pack(side="left");ttk.Button(controls,text="Close App",command=self.destroy).pack(side="right")
        self.pb=ttk.Progressbar(left,maximum=100);self.pb.pack(fill="x",pady=(12,3));self.pt=ttk.Label(left,text="Progress: 0% - Ready");self.pt.pack(anchor="w")
        logf=ttk.LabelFrame(right,text="Activity Log",padding=8);logf.pack(fill="both",expand=True);self.log=tk.Text(logf,wrap="word",font=("Consolas",10));sb=ttk.Scrollbar(logf,command=self.log.yview);self.log.configure(yscrollcommand=sb.set);self.log.pack(side="left",fill="both",expand=True);sb.pack(side="right",fill="y");self.say("Application ready.")
    def _restore(self):
        for display,rid in PERSISTENT_RESOURCES.items():
            item=self.registry.get(rid)
            if item:self.paths[display]=item["path"];loaded=datetime.fromisoformat(item["loaded_on"]).strftime("%d.%b.%Y").upper();self.status[display].configure(text=f"{item['file_name']} | loaded {loaded}",fg="#006100")
    def say(self,msg):
        if threading.current_thread() is not threading.main_thread():self.after(0,lambda:self.say(msg));return
        self.log.insert("end",f"[{datetime.now():%H:%M:%S}] {msg}\n");self.log.see("end")
    def progress(self,msg,pct):
        if threading.current_thread() is not threading.main_thread():self.after(0,lambda:self.progress(msg,pct));return
        self.pb["value"]=pct;self.pt.configure(text=f"Progress: {pct}% - {msg}");self.say(msg)
    def browse(self,name):
        p=filedialog.askopenfilename(title="Select "+name,filetypes=FILE_TYPES)
        if not p:return
        self.paths[name]=p
        if name in PERSISTENT_RESOURCES:self.registry.register(PERSISTENT_RESOURCES[name],p);item=self.registry.get(PERSISTENT_RESOURCES[name]);loaded=datetime.fromisoformat(item["loaded_on"]).strftime("%d.%b.%Y").upper();text=f"{Path(p).name} | loaded {loaded}"
        else:text=Path(p).name
        self.status[name].configure(text=text,fg="#006100");self.say(f"Loaded {name}: {p}")
        if name=="Field List and Mapping":self.fi.load_workbook(p);self.say("Enterprise Field Intelligence mapping loaded.")
    def clear(self,name):
        self.paths[name]="";self.status[name].configure(text="Not loaded",fg="#9C0006")
    def clear_session(self):
        for x in SESSION_RESOURCES:self.clear(x)
        self.results.clear();self.say("Session files cleared. Persistent template/catalog metadata retained.")
    def require(self,*names):
        missing=[x for x in names if not self.paths.get(x)]
        if missing:messagebox.showerror("Missing resources","Load:\n"+"\n".join(missing));return False
        return True
    def execute(self,fn):
        if self.busy:return
        self.busy=True
        def worker():
            try:fn()
            except Exception as exc:
                err=exc;tb=traceback.format_exc();self.after(0,lambda e=err,t=tb:self.fail(e,t))
            finally:self.busy=False
        threading.Thread(target=worker,daemon=True).start()
    def fail(self,e,t):self.say(t);self.progress("Failed",0);messagebox.showerror("Operation failed",str(e))
    def sync_yes(self,title,text):
        result=[];event=threading.Event();self.after(0,lambda:(result.append(messagebox.askyesno(title,text)),event.set()));event.wait();return result[0]
    def output_dir(self):return filedialog.askdirectory(title="Select output folder")
    def _load_field_mapping(self):
        p=self.paths.get("Field List and Mapping")
        if p:self.fi.load_workbook(p)
    def reconcile_nw_click(self):
        if not self.require("IS Operating System Report","NW CMDB Report","IS Product Catalog Operating System"):return
        out=self.output_dir();
        if not out:return
        def task():
            self._load_field_mapping();r=reconcile_nw(read_tabular(self.paths["NW CMDB Report"]),read_tabular(self.paths["IS Operating System Report"]),read_tabular(self.paths["IS Product Catalog Operating System"]),self.fi,self.progress);self.results["network"]=r;p=next_versioned_path(out,"NW CMDB to IS Reconciliation V5.2","xlsx");write_workbook(p,{"Summary":nw_summary(r),"Reconciliation":r,"Load To IS":r[r.Action.str.contains("Load To IS",na=False)],"Retire From IS":r[r.Action.str.contains("Retire From IS",na=False)],"Inventory OS Update Required":r[r.Action.str.contains("Inventory OS update required",na=False)],"Catalogue Update Required":r[r.Action.str.contains("Catalogue update required",na=False)]});self.say("Created: "+p);self.after(0,lambda:messagebox.showinfo("Complete",p))
        self.execute(task)
    def reconcile_server_click(self):
        if not self.require("IS Operating System Report","Server CMDB Report","IS Product Catalog Operating System"):return
        out=self.output_dir();
        if not out:return
        def task():
            self._load_field_mapping();cm=read_tabular(self.paths["Server CMDB Report"]);sg=find_col(cm,"Support Group",required=False);mg=find_col(cm,"Managed by Group",required=False);scope=pd.Series(False,index=cm.index)
            if sg:scope|=cm[sg].map(clean).str.upper().str.startswith("W-KYN")
            if mg:scope|=cm[mg].map(clean).str.upper().str.startswith("W-KYN")
            scoped=cm[scope].copy();outside=cm[~scope].copy();name=self.fi.field("hostname","Server CMDB Report",scoped);lc=self.fi.field("os_lifecycle_status","Server CMDB Report",scoped);ip=self.fi.field("ip_address","Server CMDB Report",scoped);serial=self.fi.field("os_parent_serial_number","Server CMDB Report",scoped,False);osc=self.fi.field("operating_system_name","Server CMDB Report",scoped);osv=self.fi.field("operating_system_version","Server CMDB Report",scoped);fq=self.fi.field("fully_qualified_hostname","Server CMDB Report",scoped,False);cl=self.fi.field("class","Server CMDB Report",scoped,False)
            base=pd.DataFrame({"Configuration Item":scoped[name],"Class":scoped[cl] if cl else "Server","Life Cycle Stage":scoped[lc],"Life Cycle Stage Status":"","Manufacturer":"","Model ID":"","Model.Name":"","Model number":"","Serial number":scoped[serial] if serial else "","IP Address":scoped[ip],"Firmware version":scoped[osc]+" "+scoped[osv],"Fully qualified domain name":scoped[fq] if fq else scoped[name]});base["_name"]=base["Configuration Item"].map(pkey);base["_life"]=base["Life Cycle Stage"].map(lifecycle);dups=base[base.duplicated(["_name","_life"],False)].copy();base=base.drop_duplicates(["_name","_life"]);active=set(base.loc[base["_life"]=="operational","_name"]);hist=base[(base["_life"]=="end of life")&base["_name"].isin(active)].copy();base=base[~((base["_life"]=="end of life")&base["_name"].isin(active))].drop(columns=["_name","_life"])
            r=reconcile_nw(base,read_tabular(self.paths["IS Operating System Report"]),read_tabular(self.paths["IS Product Catalog Operating System"]),self.fi,self.progress);r["Attribute Field Value"]=scoped.get("Attributes","").reset_index(drop=True) if "Attributes" in scoped else "";self.results["server"]=r;p=next_versioned_path(out,"Server CMDB to IS Reconciliation V5.2","xlsx");write_workbook(p,{"Reconciliation":r,"Load To IS":r[r.Action.str.contains("Load To IS",na=False)],"Retire From IS":r[r.Action.str.contains("Retire From IS",na=False)],"Out of Kyndryl Scope":outside,"Historical EOL Suppressed":hist,"CMDB Duplicate Review":dups});self.say("Created: "+p);self.after(0,lambda:messagebox.showinfo("Complete",p))
        self.execute(task)
    def _template_notice(self):
        item=self.registry.get("bulk_load_template")
        if not item:return False
        loaded=datetime.fromisoformat(item["loaded_on"]).strftime("%d.%b.%Y").upper()
        return self.sync_yes("Current Bulk Load Template",f"Current Load template was loaded on date {loaded}.\n\nThe bulk load will be created based on the loaded template. If there are changes in the bulk load template, please upload a new Bulk Load template.\n\nContinue with {item['file_name']}?")
    def bulk_click(self,source):
        required=["IS Operating System Report","Bulk Load Template","IS Product Catalog Operating System", "NW CMDB Report" if source=="network" else "Server CMDB Report"]
        if not self.require(*required):return
        if not self._template_notice():return
        out=self.output_dir();
        if not out:return
        fmt="xlsx" if messagebox.askyesno("Output format","Create XLSX? Select No for CSV.") else "csv"
        def task():
            if source not in self.results:
                self.progress("Running required reconciliation",5)
                if source=="network":r=reconcile_nw(read_tabular(self.paths["NW CMDB Report"]),read_tabular(self.paths["IS Operating System Report"]),read_tabular(self.paths["IS Product Catalog Operating System"]),self.fi,self.progress)
                else:
                    messagebox.showwarning("Run Server Reconciliation","Run Reconcile Server CMDB and IS once before creating the Server OS bulk load.");return
                self.results[source]=r
            paths=generate_bulk_load(self.results[source],source,self.paths["IS Operating System Report"],self.paths["Bulk Load Template"],out,fmt,lambda n:self.sync_yes("Missing FQDN",f"{n} devices do not have a valid FQDN. Generate <hostname>.nw.wpp.net and include them? Select No to exclude and continue."),lambda n:self.sync_yes("Catalogue Match",f"{n} devices do not have an OS catalogue match. Create a separate Catalogue Update Required file?"),self.progress)
            self.registry.mark_used("bulk_load_template");self.registry.mark_used("catalog_os")
            for p in paths:self.say("Created: "+p)
            self.after(0,lambda:messagebox.showinfo("Complete","\n".join(paths)))
        self.execute(task)

if __name__=="__main__":App().mainloop()
