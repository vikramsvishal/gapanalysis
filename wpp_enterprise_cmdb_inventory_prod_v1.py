#!/usr/bin/env python3
"""WPP Enterprise CMDB and Inventory Services Prod V1

Inputs per module:
  Network: Product Catalog, NW CMDB Report, IS Network Category Report
  Server : Product Catalog, Server CMDB Report, IS Server Category Report

Supported input formats: .csv, .xlsx, .xlsm, .xlsb
Outputs are .xlsx workbooks. Category-load headers are taken at runtime from the
corresponding IS Category report and kept in the same order, with CMDB reference
columns inserted first.

Dependencies:
  pip install openpyxl pyxlsb
"""
from __future__ import annotations

import csv
import io
import json
import re
import traceback
import sys
import subprocess
import importlib
import threading
from collections import defaultdict
from copy import copy
from dataclasses import dataclass
from datetime import date, datetime
from pathlib import Path
from typing import Any, Iterable

def _ensure_dependency(import_name: str, pip_name: str):
    try:return importlib.import_module(import_name)
    except ImportError:
        subprocess.check_call([sys.executable,"-m","pip","install","--user",pip_name])
        return importlib.import_module(import_name)
_ensure_dependency("openpyxl","openpyxl")
from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from openpyxl.utils.datetime import from_excel

APP_TITLE = "WPP Enterprise CMDB and Inventory Services Prod V1"
DATE_FORMAT = "mm/dd/yyyy"
SUPPORTED_EXTENSIONS = {".csv", ".xlsx", ".xlsm", ".xlsb"}
DATE_NOT_PUBLISHED = date(2099, 12, 31)
LEGACY_PLACEHOLDER = date(1999, 12, 31)

# -----------------------------------------------------------------------------
# Normalization and safe comparison
# -----------------------------------------------------------------------------
def text(value: Any) -> str:
    return "" if value is None else str(value).strip()


def header_key(value: Any) -> str:
    return re.sub(r"[^a-z0-9]+", "", text(value).lower())


def value_key(value: Any) -> str:
    return re.sub(r"[^a-z0-9]+", "", text(value).lower())


def serial_key(value: Any) -> str:
    # Covers case, '-', '_', '/', '\\', '.', and spaces by removing all
    # non-alphanumeric characters.
    return re.sub(r"[^a-z0-9]+", "", text(value).lower())


def model_tokens(value: Any) -> list[str]:
    raw = text(value).lower()
    values = re.findall(r"[a-z0-9]+", raw) + re.findall(r"[a-z]+|[0-9]+", raw)
    result: list[str] = []
    for token in values:
        if len(token) >= 2 and token not in result:
            result.append(token)
    return result


def equal_value(left: Any, right: Any) -> bool:
    return value_key(left) == value_key(right)


def is_operational(value: Any) -> bool:
    return value_key(value) == "operational"


# -----------------------------------------------------------------------------
# Date handling
# -----------------------------------------------------------------------------
def parse_date(value: Any) -> date | None:
    if value is None or text(value) == "":
        return None
    if isinstance(value, datetime):
        return value.date()
    if isinstance(value, date):
        return value
    if isinstance(value, (int, float)):
        try:
            return from_excel(value).date()
        except Exception:
            return None
    raw = text(value)
    formats = (
        "%Y-%m-%d", "%d-%m-%Y", "%m/%d/%Y", "%d/%m/%Y", "%Y/%m/%d",
        "%m-%d-%Y", "%d-%b-%Y", "%d %b %Y", "%b %d, %Y",
        "%Y-%m-%d %H:%M:%S", "%d-%m-%Y %H:%M:%S",
    )
    for fmt in formats:
        try:
            return datetime.strptime(raw, fmt).date()
        except ValueError:
            pass
    try:
        return datetime.fromisoformat(raw.replace("Z", "+00:00")).date()
    except ValueError:
        return None


def lifecycle_date_state(value: Any) -> tuple[date | None, str]:
    parsed = parse_date(value)
    if parsed is None:
        return None, "BLANK_OR_UNPARSEABLE"
    if parsed == DATE_NOT_PUBLISHED:
        return parsed, "DATE_NOT_PUBLISHED"
    if parsed == LEGACY_PLACEHOLDER:
        return parsed, "LEGACY_PLACEHOLDER_REVIEW"
    return parsed, "PUBLISHED_DATE"


def compare_date_pair(cmdb_value: Any, catalog_value: Any, label: str) -> str:
    cmdb_date, cmdb_state = lifecycle_date_state(cmdb_value)
    catalog_date, catalog_state = lifecycle_date_state(catalog_value)
    if catalog_state == "DATE_NOT_PUBLISHED":
        return f"{label}: CATALOG DATE NOT PUBLISHED"
    if catalog_state == "LEGACY_PLACEHOLDER_REVIEW":
        return f"{label}: CATALOG LEGACY PLACEHOLDER REVIEW"
    if cmdb_date is None and catalog_date is None:
        return f"{label}: BOTH BLANK"
    if cmdb_date is None:
        return f"{label}: CMDB DATE MISSING"
    if catalog_date is None:
        return f"{label}: CATALOG DATE MISSING"
    return f"{label}: MATCH" if cmdb_date == catalog_date else f"{label}: MISMATCH"


# -----------------------------------------------------------------------------
# Generic tabular readers
# -----------------------------------------------------------------------------
@dataclass(frozen=True)
class TableData:
    headers: list[str]
    rows: list[list[Any]]
    source_sheet: str
    header_row: int
    metadata_row: list[Any] | None = None

    @property
    def columns(self) -> dict[str, int]:
        result: dict[str, int] = {}
        for index, name in enumerate(self.headers):
            key = header_key(name)
            if key and key not in result:
                result[key] = index
        return result


def decode_csv(path: Path) -> str:
    raw = path.read_bytes()
    for encoding in ("utf-8-sig", "utf-8", "cp1252", "latin-1"):
        try:
            return raw.decode(encoding)
        except UnicodeDecodeError:
            pass
    raise ValueError(f"Unable to decode CSV file: {path.name}")


def identify_header(rows: list[list[Any]], expected: Iterable[str], filename: str) -> int:
    expected_keys = {header_key(item) for item in expected}
    best_score, best_index = -1, 0
    for index, row in enumerate(rows[:25]):
        score = len({header_key(item) for item in row if text(item)} & expected_keys)
        if score > best_score:
            best_score, best_index = score, index
    if best_score <= 0:
        raise ValueError(f"Expected columns were not found in {filename}.")
    return best_index


def looks_like_requiredness_row(row: list[Any]) -> bool:
    values = [value_key(v) for v in row if text(v)]
    if not values:
        return False
    recognized = sum(v in {"mandatory", "optional", "required", "reference"} for v in values)
    return recognized >= max(1, len(values) // 2)


def read_csv_table(path: Path, expected: Iterable[str]) -> TableData:
    content = decode_csv(path)
    try:
        dialect = csv.Sniffer().sniff(content[:65536], delimiters=",;\t|")
    except csv.Error:
        dialect = csv.excel
    all_rows = list(csv.reader(io.StringIO(content), dialect))
    if not all_rows:
        raise ValueError(f"{path.name} is empty.")
    header_index = identify_header(all_rows, expected, path.name)
    headers = [text(v) for v in all_rows[header_index]]
    width = len(headers)
    metadata = None
    data_start = header_index + 1
    if data_start < len(all_rows) and looks_like_requiredness_row(all_rows[data_start]):
        metadata = (all_rows[data_start] + [None] * width)[:width]
        data_start += 1
    rows = [(row + [None] * width)[:width] for row in all_rows[data_start:]]
    return TableData(headers, rows, "CSV", header_index + 1, metadata)


def read_excel_table(path: Path, expected: Iterable[str]) -> TableData:
    workbook = load_workbook(path, data_only=True, read_only=True)
    try:
        expected_keys = {header_key(item) for item in expected}
        best_score, best_sheet, best_row = -1, None, 1
        for worksheet in workbook.worksheets:
            for row_number in range(1, min(worksheet.max_row, 25) + 1):
                values = [cell.value for cell in worksheet[row_number]]
                score = len({header_key(v) for v in values if text(v)} & expected_keys)
                if score > best_score:
                    best_score, best_sheet, best_row = score, worksheet, row_number
        if best_sheet is None or best_score <= 0:
            raise ValueError(f"Expected columns were not found in {path.name}.")
        headers = [text(cell.value) for cell in best_sheet[best_row]]
        width = len(headers)
        start = best_row + 1
        metadata = None
        if start <= best_sheet.max_row:
            candidate = [best_sheet.cell(start, col).value for col in range(1, width + 1)]
            if looks_like_requiredness_row(candidate):
                metadata = candidate
                start += 1
        rows = [list(values) for values in best_sheet.iter_rows(
            min_row=start, max_col=width, values_only=True
        )]
        return TableData(headers, rows, best_sheet.title, best_row, metadata)
    finally:
        workbook.close()


def read_xlsb_table(path: Path, expected: Iterable[str]) -> TableData:
    try:
        from pyxlsb import open_workbook
    except ImportError as exc:
        raise RuntimeError("XLSB support requires pyxlsb. Install: pip install pyxlsb") from exc
    expected_keys = {header_key(item) for item in expected}
    best_score, best_headers, best_rows, best_sheet, best_header_index = -1, None, None, "", 0
    with open_workbook(str(path)) as workbook:
        for sheet_name in workbook.sheets:
            with workbook.get_sheet(sheet_name) as sheet:
                rows = [[cell.v for cell in row] for row in sheet.rows()]
            for index, row in enumerate(rows[:25]):
                score = len({header_key(v) for v in row if text(v)} & expected_keys)
                if score > best_score:
                    best_score = score
                    best_headers = row
                    best_rows = rows[index + 1:]
                    best_sheet = sheet_name
                    best_header_index = index
    if best_headers is None or best_rows is None or best_score <= 0:
        raise ValueError(f"Expected columns were not found in {path.name}.")
    headers = [text(v) for v in best_headers]
    width = len(headers)
    metadata = None
    if best_rows and looks_like_requiredness_row(best_rows[0]):
        metadata = (best_rows[0] + [None] * width)[:width]
        best_rows = best_rows[1:]
    rows = [(row + [None] * width)[:width] for row in best_rows]
    return TableData(headers, rows, best_sheet, best_header_index + 1, metadata)


def read_table(path: Path, expected: Iterable[str]) -> TableData:
    if not path.is_file():
        raise FileNotFoundError(path)
    extension = path.suffix.lower()
    if extension == ".csv":
        return read_csv_table(path, expected)
    if extension in {".xlsx", ".xlsm"}:
        return read_excel_table(path, expected)
    if extension == ".xlsb":
        return read_xlsb_table(path, expected)
    raise ValueError(f"Unsupported file type: {extension}")


def required_column(table: TableData, aliases: Iterable[str], label: str) -> int:
    for alias in aliases:
        if header_key(alias) in table.columns:
            return table.columns[header_key(alias)]
    raise ValueError(f"{label} column was not found. Expected one of: {', '.join(aliases)}")


def optional_column(table: TableData, aliases: Iterable[str]) -> int | None:
    for alias in aliases:
        if header_key(alias) in table.columns:
            return table.columns[header_key(alias)]
    return None


def row_value(row: list[Any], index: int | None) -> Any:
    return None if index is None or index >= len(row) else row[index]


# -----------------------------------------------------------------------------
# Catalog records and indexes
# -----------------------------------------------------------------------------
@dataclass(frozen=True)
class CatalogRecord:
    manufacturer: str
    hardware_type: str
    hardware_model: str
    subcategory: str
    eol: Any
    eosl: Any
    product_opaque_id: str
    source_row: int
    manufacturer_key: str
    type_key: str
    model_key: str
    type_tokens: tuple[str, ...]
    model_tokens: tuple[str, ...]


@dataclass(frozen=True)
class NetworkCatalogIndexes:
    records: tuple[CatalogRecord, ...]
    exact_model_index: dict[str, tuple[int, ...]]
    hardware_type_index: dict[str, tuple[int, ...]]
    partial_token_index: dict[str, tuple[int, ...]]
    composite_index: dict[str, tuple[int, ...]]


@dataclass(frozen=True)
class ServerCatalogIndexes:
    records: tuple[CatalogRecord, ...]
    manufacturer_model_index: dict[str, tuple[int, ...]]
    model_only_index: dict[str, tuple[int, ...]]
    hardware_type_index: dict[str, tuple[int, ...]]


@dataclass(frozen=True)
class CatalogMatch:
    record: CatalogRecord | None
    status: str
    detail: str
    score: int = 0
    ambiguous: bool = False


def build_catalog_records(table: TableData) -> list[CatalogRecord]:
    cols = {
        "manufacturer": required_column(table, ["manufacturer"], "Catalog manufacturer"),
        "hardware_type": required_column(table, ["hardware_type"], "Catalog hardware_type"),
        "hardware_model": required_column(table, ["hardware_model"], "Catalog hardware_model"),
        "subcategory": required_column(table, ["subcategory"], "Catalog subcategory"),
        "eol": optional_column(table, ["end_of_life_date"]),
        "eosl": optional_column(table, ["end_of_support_date"]),
        "opaque": optional_column(table, ["product_opaque_id"]),
    }
    records: list[CatalogRecord] = []
    for source_row, row in enumerate(table.rows, start=table.header_row + 1 + (1 if table.metadata_row else 0)):
        manufacturer = text(row_value(row, cols["manufacturer"]))
        hardware_type = text(row_value(row, cols["hardware_type"]))
        hardware_model = text(row_value(row, cols["hardware_model"]))
        if not manufacturer and not hardware_type and not hardware_model:
            continue
        records.append(CatalogRecord(
            manufacturer=manufacturer,
            hardware_type=hardware_type,
            hardware_model=hardware_model,
            subcategory=text(row_value(row, cols["subcategory"])),
            eol=row_value(row, cols["eol"]),
            eosl=row_value(row, cols["eosl"]),
            product_opaque_id=text(row_value(row, cols["opaque"])),
            source_row=source_row,
            manufacturer_key=value_key(manufacturer),
            type_key=value_key(hardware_type),
            model_key=value_key(hardware_model),
            type_tokens=tuple(model_tokens(hardware_type)),
            model_tokens=tuple(model_tokens(hardware_model)),
        ))
    if not records:
        raise ValueError("Catalog contains no usable records.")
    return records


def build_network_catalog_indexes(table: TableData) -> NetworkCatalogIndexes:
    records = build_catalog_records(table)
    exact, types, tokens_index, composites = defaultdict(list), defaultdict(list), defaultdict(list), defaultdict(list)
    for record_id, record in enumerate(records):
        if record.model_key:
            exact[record.model_key].append(record_id)
        if record.type_key:
            types[record.type_key].append(record_id)
        if record.type_key and record.model_key:
            composites[record.type_key + record.model_key].append(record_id)
        for token in set(record.type_tokens + record.model_tokens):
            tokens_index[token].append(record_id)
    freeze = lambda source: {key: tuple(values) for key, values in source.items()}
    return NetworkCatalogIndexes(tuple(records), freeze(exact), freeze(types), freeze(tokens_index), freeze(composites))


def build_server_catalog_indexes(table: TableData) -> ServerCatalogIndexes:
    records = build_catalog_records(table)
    manufacturer_model, model_only, types = defaultdict(list), defaultdict(list), defaultdict(list)
    for record_id, record in enumerate(records):
        if record.manufacturer_key and record.model_key:
            manufacturer_model[f"{record.manufacturer_key}|{record.model_key}"].append(record_id)
        if record.model_key:
            model_only[record.model_key].append(record_id)
        if record.type_key:
            types[record.type_key].append(record_id)
    freeze = lambda source: {key: tuple(values) for key, values in source.items()}
    return ServerCatalogIndexes(tuple(records), freeze(manufacturer_model), freeze(model_only), freeze(types))


def choose_unique(records: list[CatalogRecord], status: str, detail: str, score: int) -> CatalogMatch:
    outcomes = {
        (r.manufacturer_key, r.type_key, r.model_key, value_key(r.subcategory), parse_date(r.eol), parse_date(r.eosl))
        for r in records
    }
    if len(outcomes) == 1:
        return CatalogMatch(records[0], status, detail, score)
    return CatalogMatch(None, "NO MATCH", f"AMBIGUOUS {detail}", score, True)


def network_fallback_score(query: str, record: CatalogRecord) -> int:
    query_key = value_key(query)
    query_tokens = model_tokens(query)
    if not query_key or not record.type_key:
        return 0
    if record.type_key == query_key:
        score = 900
    elif record.type_key in query_key:
        score = 700 + len(record.type_key)
    elif query_key in record.type_key:
        score = 550 + len(query_key)
    else:
        shared = set(query_tokens) & set(record.type_tokens)
        if not shared:
            return 0
        score = 120 * len(shared)
    score += 250 * len(set(query_tokens) & set(record.model_tokens))
    if record.model_key and record.model_key in query_key:
        score += 350
    for token in query_tokens:
        if token in record.model_key:
            score += 80
    return score


def match_network_model(model_number: Any, indexes: NetworkCatalogIndexes) -> CatalogMatch:
    raw = text(model_number)
    query = value_key(raw)
    if not query:
        return CatalogMatch(None, "NO MATCH", "BLANK MODEL NUMBER")
    ids = indexes.exact_model_index.get(query, ())
    if ids:
        return choose_unique([indexes.records[i] for i in ids], "EXACT MATCH", "EXACT HARDWARE_MODEL", 10000)
    ids = indexes.composite_index.get(query, ())
    if ids:
        return choose_unique([indexes.records[i] for i in ids], "EXACT MATCH", "EXACT TYPE + MODEL COMPOSITE", 9500)

    candidates: set[int] = set(indexes.hardware_type_index.get(query, ()))
    for token in model_tokens(raw):
        candidates.update(indexes.hardware_type_index.get(value_key(token), ()))
        candidates.update(indexes.partial_token_index.get(token, ()))
    # Indexed type-prefix/containment pass supports ASA5506-K9 -> asa5506.
    for type_key, type_ids in indexes.hardware_type_index.items():
        if len(type_key) >= 3 and (type_key in query or query in type_key):
            candidates.update(type_ids)
    ranked = [(network_fallback_score(raw, indexes.records[i]), indexes.records[i]) for i in candidates]
    ranked = [(score, record) for score, record in ranked if score > 0]
    if not ranked:
        return CatalogMatch(None, "NO MATCH", "NO INDEXED CANDIDATE")
    ranked.sort(key=lambda item: (-item[0], item[1].source_row))
    top_score = ranked[0][0]
    top_records = [record for score, record in ranked if score == top_score]
    return choose_unique(top_records, "PARTIAL MATCH", "INDEXED HARDWARE_TYPE/TOKEN MATCH", top_score)


def match_server_model(manufacturer: Any, model_id: Any, indexes: ServerCatalogIndexes) -> CatalogMatch:
    manufacturer_value, model_value=text(manufacturer),text(model_id);mk,q=value_key(manufacturer_value),value_key(model_value)
    if not q:return CatalogMatch(None,"NO MATCH","BLANK MODEL ID")
    ids=indexes.manufacturer_model_index.get(f"{mk}|{q}",()) if mk else ()
    if ids:return choose_unique([indexes.records[i] for i in ids],"EXACT MATCH","EXACT MANUFACTURER + MODEL ID",10000)
    ids=indexes.model_only_index.get(q,())
    if ids:
        records=[indexes.records[i] for i in ids]
        chosen=choose_unique(records,"EXACT MATCH","EXACT MODEL ID",9800)
        if chosen.record and mk and chosen.record.manufacturer_key!=mk:return CatalogMatch(chosen.record,"PARTIAL MATCH","MODEL MATCH - MANUFACTURER REVIEW",9800)
        return chosen
    # Option C: hardware_type + hardware_model, then indexed partial/token fallback.
    candidates=[]
    for r in indexes.records:
        composite=r.type_key+r.model_key
        if composite==q:candidates.append((9500,r,"EXACT TYPE + MODEL COMPOSITE"))
    if candidates:
        rs=[x[1] for x in candidates];chosen=choose_unique(rs,"EXACT MATCH","EXACT TYPE + MODEL COMPOSITE",9500)
        if chosen.record and mk and chosen.record.manufacturer_key!=mk:return CatalogMatch(chosen.record,"PARTIAL MATCH","MODEL MATCH - MANUFACTURER REVIEW",9500)
        return chosen
    qt=set(model_tokens(model_value));ranked=[]
    for r in indexes.records:
        score=0
        if r.type_key and r.type_key in q:score+=700+len(r.type_key)
        score+=250*len(qt&set(r.model_tokens))+120*len(qt&set(r.type_tokens))
        if r.model_key and r.model_key in q:score+=350
        if mk and r.manufacturer_key==mk:score+=300
        if score:ranked.append((score,r))
    if not ranked:return CatalogMatch(None,"NO MATCH","NO INDEXED SERVER CANDIDATE")
    top=max(x[0] for x in ranked);rs=[r for sc,r in ranked if sc==top];chosen=choose_unique(rs,"PARTIAL MATCH","INDEXED HARDWARE TYPE/TOKEN MATCH",top)
    if chosen.record and mk and chosen.record.manufacturer_key!=mk:return CatalogMatch(chosen.record,"PARTIAL MATCH","MODEL MATCH - MANUFACTURER REVIEW",top)
    return chosen


# -----------------------------------------------------------------------------
# Category indexes and output helpers
# -----------------------------------------------------------------------------
def build_serial_index(table: TableData, serial_aliases: Iterable[str]) -> tuple[dict[str, list[int]], int]:
    serial_col = required_column(table, serial_aliases, "Category serial number")
    index: dict[str, list[int]] = defaultdict(list)
    for row_id, row in enumerate(table.rows):
        normalized = serial_key(row_value(row, serial_col))
        if normalized:
            index[normalized].append(row_id)
    return dict(index), serial_col


def category_row_for_serial(table: TableData, serial_index: dict[str, list[int]], serial: Any) -> tuple[list[Any] | None, str]:
    key = serial_key(serial)
    if not key:
        return None, "BLANK SERIAL NUMBER"
    ids = serial_index.get(key, [])
    if not ids:
        return None, "MISSING FROM CATEGORY"
    if len(ids) > 1:
        return table.rows[ids[0]], "FOUND - DUPLICATE NORMALIZED SERIAL"
    return table.rows[ids[0]], "FOUND IN CATEGORY"


def standard_headers(title: str) -> list[str]:
    return [
        "CMDB Hostname", "CMDB Serial Number", "CMDB Manufacturer", "CMDB Model ID",
        "CMDB Model Number", "CMDB Lifecycle Stage", "CMDB Lifecycle Stage Status",
        "CMDB Class", "CMDB Category", "CMDB Subcategory",
        "Catalog Manufacturer", "Catalog Hardware Type", "Catalog Hardware Model",
        "Catalog Subcategory", "Catalog Product Opaque ID", "Catalog End of Life Date",
        "Catalog End of Support Date", "Catalog Match Status", "Catalog Match Detail",
        "Lifecycle Date Validation", "Category Presence Status", "Module"
    ]


def write_table_workbook(path: Path, sheet_name: str, headers: list[str], rows: list[list[Any]], date_headers: set[str] | None = None) -> None:
    workbook = Workbook()
    worksheet = workbook.active
    worksheet.title = sheet_name[:31]
    header_fill = PatternFill("solid", fgColor="1F4E78")
    header_font = Font(color="FFFFFF", bold=True)
    for col, header in enumerate(headers, 1):
        cell = worksheet.cell(1, col, header)
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    for row_number, row in enumerate(rows, 2):
        for col_number, value in enumerate(row, 1):
            cell = worksheet.cell(row_number, col_number, value)
            if date_headers and headers[col_number - 1] in date_headers and parse_date(value):
                cell.value = parse_date(value)
                cell.number_format = DATE_FORMAT
    worksheet.freeze_panes = "A2"
    if headers:
        worksheet.auto_filter.ref = f"A1:{get_column_letter(len(headers))}{max(1, len(rows)+1)}"
    for col_number, header in enumerate(headers, 1):
        width = max(12, min(36, len(text(header)) + 3))
        worksheet.column_dimensions[get_column_letter(col_number)].width = width
    workbook.save(path)


def make_load_workbook(path: Path, prefix_headers: list[str], category: TableData, rows: list[tuple[list[Any], list[Any]]]) -> None:
    headers = prefix_headers + category.headers
    workbook = Workbook()
    worksheet = workbook.active
    worksheet.title = "Category Load"
    header_fill = PatternFill("solid", fgColor="1F4E78")
    for col, header in enumerate(headers, 1):
        cell = worksheet.cell(1, col, header)
        cell.fill = header_fill
        cell.font = Font(color="FFFFFF", bold=True)
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    row_number = 2
    if category.metadata_row is not None:
        prefix_metadata = ["Reference"] * len(prefix_headers)
        for col, value in enumerate(prefix_metadata + category.metadata_row, 1):
            worksheet.cell(row_number, col, value)
        row_number += 1
    for prefix, category_values in rows:
        values = prefix + category_values
        for col, value in enumerate(values, 1):
            worksheet.cell(row_number, col, value)
        row_number += 1
    worksheet.freeze_panes = f"A{3 if category.metadata_row else 2}"
    worksheet.auto_filter.ref = f"A1:{get_column_letter(len(headers))}{max(1,row_number-1)}"
    workbook.save(path)


def blank_category_row(category: TableData, mappings: dict[str, Any]) -> list[Any]:
    row = [None] * len(category.headers)
    for logical_name, value in mappings.items():
        target = category.columns.get(header_key(logical_name))
        if target is not None:
            row[target] = value
    return row


def drift_action(flags: list[str]) -> str:
    changes = [flag for flag in flags if flag]
    if not changes:
        return "NO UPDATE REQUIRED"
    if len(changes) == 1:
        return f"UPDATE {changes[0]}"
    return "UPDATE MULTIPLE FIELDS: " + ", ".join(changes)


# -----------------------------------------------------------------------------
# Network reconciliation
# -----------------------------------------------------------------------------
def next_output_path(folder: Path, stem: str) -> Path:
    folder=Path(folder);folder.mkdir(parents=True,exist_ok=True);day=datetime.now().strftime("%d-%b-%Y")
    pattern=re.compile(re.escape(stem)+r" - "+re.escape(day)+r" - V(\d+)\.xlsx$",re.I);versions=[]
    for f in folder.glob(f"{stem} - {day} - V*.xlsx"):
        match=pattern.match(f.name)
        if match:versions.append(int(match.group(1)))
    return folder/f"{stem} - {day} - V{max(versions,default=0)+1}.xlsx"

class NetworkHardwareGovernanceAgent:
    def execute(self,catalog_path,cmdb_path,category_path,output_dir):return run_network(Path(catalog_path),Path(cmdb_path),Path(category_path),Path(output_dir))
class ServerHardwareGovernanceAgent:
    def execute(self,catalog_path,cmdb_path,category_path,output_dir):return run_server(Path(catalog_path),Path(cmdb_path),Path(category_path),Path(output_dir))
class HardwareGovernanceService:
    def run_network(self,request):return NetworkHardwareGovernanceAgent().execute(**request)
    def run_server(self,request):return ServerHardwareGovernanceAgent().execute(**request)

def run_network(catalog_path: Path, cmdb_path: Path, category_path: Path, output_dir: Path) -> dict[str, Any]:
    catalog_table = read_table(catalog_path, ["manufacturer", "hardware_type", "hardware_model", "subcategory"])
    cmdb = read_table(cmdb_path, ["Model number", "Serial number", "Life Cycle Stage"])
    category = read_table(category_path, ["network_serial_number", "network_manufacturer", "network_hardware_model"])
    indexes = build_network_catalog_indexes(catalog_table)
    serial_index, _ = build_serial_index(category, ["network_serial_number"])

    c = {
        "hostname": required_column(cmdb, ["Configuration Item", "CI.Name", "Fully qualified domain name"], "NW CMDB hostname"),
        "serial": required_column(cmdb, ["Serial number"], "NW CMDB serial"),
        "manufacturer": optional_column(cmdb, ["Manufacturer"]),
        "model_id": optional_column(cmdb, ["Model ID"]),
        "model_number": required_column(cmdb, ["Model number"], "NW CMDB Model number"),
        "stage": required_column(cmdb, ["Life Cycle Stage"], "NW CMDB lifecycle stage"),
        "stage_status": optional_column(cmdb, ["Life Cycle Stage Status"]),
        "class": optional_column(cmdb, ["Class"]),
        "category": optional_column(cmdb, ["Category"]),
        "subcategory": optional_column(cmdb, ["Subcategory"]),
        "eol": optional_column(cmdb, ["HW_End_Of_Life_Date", "HW End Of Life Date"]),
        "eosl": optional_column(cmdb, ["HW_End_Of_Support_Date", "HW End Of Support Date"]),
    }
    ic = {
        "opaque": optional_column(category, ["network_opaque_id", "id"]),
        "serial": required_column(category, ["network_serial_number"], "IS Network serial"),
        "manufacturer": required_column(category, ["network_manufacturer"], "IS Network manufacturer"),
        "type": required_column(category, ["network_hardware_type"], "IS Network hardware type"),
        "model": required_column(category, ["network_hardware_model"], "IS Network hardware model"),
        "subcategory": required_column(category, ["network_subcategory"], "IS Network subcategory"),
    }

    match_headers = standard_headers("Network")
    match_rows, load_rows, drift_rows, recommendation_rows = [], [], [], []
    drift_headers = [
        "CMDB CI Name", "CMDB Serial Number", "CMDB Model Number", "network_opaque_id",
        "network_serial_number", "Current network_manufacturer", "Current network_hardware_type",
        "Current network_hardware_model", "Current network_subcategory",
        "Required network_manufacturer", "Required network_hardware_type",
        "Required network_hardware_model", "Required network_subcategory",
        "Manufacturer Match?", "Hardware Type Match?", "Hardware Model Match?",
        "Subcategory Match?", "Catalog Match Status", "Catalog Match Detail", "Required Action"
    ]
    rec_headers = [
        "CMDB CI Name", "CMDB Serial Number", "CMDB Model Number", "CMDB Manufacturer",
        "CMDB Class", "CMDB Category", "CMDB Subcategory", "Recommendation Reason"
    ]
    stats = defaultdict(int)
    prefix_headers = [
        "CMDB Hostname", "CMDB Serial Number", "CMDB Model ID", "CMDB Model Number",
        "CMDB Lifecycle Stage", "CMDB Lifecycle Stage Status"
    ]

    for row in cmdb.rows:
        if not any(text(value) for value in row):
            continue
        stats["total_cmdb"] += 1
        hostname = row_value(row, c["hostname"])
        serial = row_value(row, c["serial"])
        model_number = row_value(row, c["model_number"])
        stage = row_value(row, c["stage"])
        result = match_network_model(model_number, indexes)
        category_row, presence = category_row_for_serial(category, serial_index, serial)
        if result.record:
            stats["catalog_matched"] += 1
            stats["exact_matches" if result.status == "EXACT MATCH" else "partial_matches"] += 1
            lifecycle = "; ".join([
                compare_date_pair(row_value(row, c["eol"]), result.record.eol, "EOL"),
                compare_date_pair(row_value(row, c["eosl"]), result.record.eosl, "EOSL"),
            ])
        else:
            stats["catalog_unmatched"] += 1
            lifecycle = "NOT APPLICABLE - NO CATALOG MATCH"
            recommendation_rows.append([
                hostname, serial, model_number, row_value(row, c["manufacturer"]),
                row_value(row, c["class"]), row_value(row, c["category"]), row_value(row, c["subcategory"]),
                result.detail,
            ])

        rec = result.record
        match_rows.append([
            hostname, serial, row_value(row, c["manufacturer"]), row_value(row, c["model_id"]),
            model_number, stage, row_value(row, c["stage_status"]), row_value(row, c["class"]),
            row_value(row, c["category"]), row_value(row, c["subcategory"]),
            rec.manufacturer if rec else None, rec.hardware_type if rec else None,
            rec.hardware_model if rec else None, rec.subcategory if rec else None,
            rec.product_opaque_id if rec else None, parse_date(rec.eol) if rec else None,
            parse_date(rec.eosl) if rec else None, result.status, result.detail, lifecycle, presence, "Network"
        ])

        if is_operational(stage):
            stats["operational_cmdb"] += 1
            if category_row is None:
                if serial_key(serial):
                    stats["missing_from_category"] += 1
                    cat_values = blank_category_row(category, {
                        "network_serial_number": serial,
                        "network_manufacturer": rec.manufacturer if rec else None,
                        "network_hardware_type": rec.hardware_type if rec else None,
                        "network_hardware_model": rec.hardware_model if rec else None,
                        "network_subcategory": rec.subcategory if rec else None,
                        "network_lifecycle_status": "PRODUCTION",
                    })
                    prefix = [hostname, serial, row_value(row, c["model_id"]), model_number, stage, row_value(row, c["stage_status"])]
                    load_rows.append((prefix, cat_values))
                else:
                    stats["operational_blank_serial"] += 1
            else:
                stats["found_in_category"] += 1

        if category_row is not None and rec is not None:
            current = {
                "manufacturer": row_value(category_row, ic["manufacturer"]),
                "type": row_value(category_row, ic["type"]),
                "model": row_value(category_row, ic["model"]),
                "subcategory": row_value(category_row, ic["subcategory"]),
            }
            flags = [
                "MANUFACTURER" if not equal_value(current["manufacturer"], rec.manufacturer) else "",
                "HARDWARE TYPE" if not equal_value(current["type"], rec.hardware_type) else "",
                "HARDWARE MODEL" if not equal_value(current["model"], rec.hardware_model) else "",
                "SUBCATEGORY" if not equal_value(current["subcategory"], rec.subcategory) else "",
            ]
            action = drift_action(flags)
            if action != "NO UPDATE REQUIRED":
                stats["category_drift"] += 1
                drift_rows.append([
                    hostname, serial, model_number, row_value(category_row, ic["opaque"]),
                    row_value(category_row, ic["serial"]), current["manufacturer"], current["type"],
                    current["model"], current["subcategory"], rec.manufacturer, rec.hardware_type,
                    rec.hardware_model, rec.subcategory,
                    "YES" if not flags[0] else "NO", "YES" if not flags[1] else "NO",
                    "YES" if not flags[2] else "NO", "YES" if not flags[3] else "NO",
                    result.status, result.detail, action,
                ])

    output_dir.mkdir(parents=True, exist_ok=True)
    date_headers = {"Catalog End of Life Date", "Catalog End of Support Date"}
    write_table_workbook(next_output_path(output_dir,"NW Hardware Catalog Match"), "NW Catalog Match", match_headers, match_rows, date_headers)
    make_load_workbook(next_output_path(output_dir,"NW Category Load"), prefix_headers, category, load_rows)
    write_table_workbook(next_output_path(output_dir,"NW Category Catalog Drift"), "NW Category Drift", drift_headers, drift_rows)
    write_table_workbook(next_output_path(output_dir,"NW Catalog Enhancement Recommendations"), "NW Catalog Recommendations", rec_headers, recommendation_rows)
    return dict(stats)


# -----------------------------------------------------------------------------
# Server reconciliation
# -----------------------------------------------------------------------------
def run_server(catalog_path: Path, cmdb_path: Path, category_path: Path, output_dir: Path) -> dict[str, Any]:
    catalog_table = read_table(catalog_path, ["manufacturer", "hardware_type", "hardware_model", "subcategory"])
    cmdb = read_table(cmdb_path, ["Manufacturer", "Model ID", "Serial number", "Life Cycle Stage", "Due"])
    category = read_table(category_path, ["serial_number", "server_manufacturer", "server_hardware_model"])
    indexes = build_server_catalog_indexes(catalog_table)
    serial_index, _ = build_serial_index(category, ["serial_number", "server_serial_number"])

    c = {
        "hostname": required_column(cmdb, ["Name", "Fully qualified domain name"], "Server CMDB hostname"),
        "serial": required_column(cmdb, ["Serial number"], "Server CMDB serial"),
        "manufacturer": required_column(cmdb, ["Manufacturer"], "Server CMDB manufacturer"),
        "model_id": required_column(cmdb, ["Model ID"], "Server CMDB Model ID"),
        "stage": required_column(cmdb, ["Life Cycle Stage"], "Server CMDB lifecycle stage"),
        "stage_status": optional_column(cmdb, ["Life Cycle Stage Status"]),
        "class": optional_column(cmdb, ["Class"]),
        "category": optional_column(cmdb, ["Category"]),
        "subcategory": optional_column(cmdb, ["Subcategory"]),
        "due": optional_column(cmdb, ["Due"]),
    }
    ic = {
        "opaque": optional_column(category, ["server_opaque_id", "id"]),
        "serial": required_column(category, ["serial_number", "server_serial_number"], "IS Server serial"),
        "manufacturer": required_column(category, ["server_manufacturer"], "IS Server manufacturer"),
        "type": required_column(category, ["server_hardware_type"], "IS Server hardware type"),
        "model": required_column(category, ["server_hardware_model"], "IS Server hardware model"),
        "subcategory": required_column(category, ["server_subcategory"], "IS Server subcategory"),
    }

    match_headers = standard_headers("Server")
    match_rows, load_rows, drift_rows, recommendation_rows = [], [], [], []
    drift_headers = [
        "CMDB Hostname", "CMDB Serial Number", "CMDB Manufacturer", "CMDB Model ID",
        "server_opaque_id", "server_serial_number", "Current server_manufacturer",
        "Current server_hardware_type", "Current server_hardware_model", "Current server_subcategory",
        "Required server_manufacturer", "Required server_hardware_type",
        "Required server_hardware_model", "Required server_subcategory",
        "Manufacturer Match?", "Hardware Type Match?", "Hardware Model Match?",
        "Subcategory Match?", "Catalog Match Status", "Catalog Match Detail", "Required Action"
    ]
    rec_headers = [
        "CMDB Hostname", "CMDB Serial Number", "CMDB Manufacturer", "CMDB Model ID",
        "CMDB Class", "CMDB Category", "CMDB Subcategory", "Is Virtual", "Recommendation Reason"
    ]
    prefix_headers = [
        "CMDB Hostname", "CMDB Serial Number", "CMDB Model ID",
        "CMDB Lifecycle Stage", "CMDB Lifecycle Stage Status"
    ]
    stats = defaultdict(int)

    for row in cmdb.rows:
        if not any(text(value) for value in row):
            continue
        stats["total_cmdb"] += 1
        hostname = row_value(row, c["hostname"])
        serial = row_value(row, c["serial"])
        manufacturer = row_value(row, c["manufacturer"])
        model_id = row_value(row, c["model_id"])
        stage = row_value(row, c["stage"])
        result = match_server_model(manufacturer, model_id, indexes)
        category_row, presence = category_row_for_serial(category, serial_index, serial)
        rec = result.record
        if rec:
            stats["catalog_matched"] += 1
            stats["exact_matches" if result.status == "EXACT MATCH" else "partial_matches"] += 1
            lifecycle = compare_date_pair(row_value(row, c["due"]), rec.eol, "EOL")
        else:
            stats["catalog_unmatched"] += 1
            lifecycle = "NOT APPLICABLE - NO CATALOG MATCH"
            recommendation_rows.append([
                hostname, serial, manufacturer, model_id, row_value(row, c["class"]),
                row_value(row, c["category"]), row_value(row, c["subcategory"]),
                row_value(row, optional_column(cmdb, ["Is Virtual"])), result.detail,
            ])

        match_rows.append([
            hostname, serial, manufacturer, model_id, None, stage, row_value(row, c["stage_status"]),
            row_value(row, c["class"]), row_value(row, c["category"]), row_value(row, c["subcategory"]),
            rec.manufacturer if rec else None, rec.hardware_type if rec else None,
            rec.hardware_model if rec else None, rec.subcategory if rec else None,
            rec.product_opaque_id if rec else None, parse_date(rec.eol) if rec else None,
            parse_date(rec.eosl) if rec else None, result.status, result.detail, lifecycle, presence, "Server"
        ])

        if is_operational(stage):
            stats["operational_cmdb"] += 1
            if category_row is None:
                if serial_key(serial):
                    stats["missing_from_category"] += 1
                    cat_values = blank_category_row(category, {
                        "serial_number": serial,
                        "server_serial_number": serial,
                        "server_manufacturer": rec.manufacturer if rec else None,
                        "server_hardware_type": rec.hardware_type if rec else None,
                        "server_hardware_model": rec.hardware_model if rec else None,
                        "server_subcategory": rec.subcategory if rec else None,
                        "server_lifecycle_status": "PRODUCTION",
                    })
                    prefix = [hostname, serial, model_id, stage, row_value(row, c["stage_status"])]
                    load_rows.append((prefix, cat_values))
                else:
                    stats["operational_blank_serial"] += 1
            else:
                stats["found_in_category"] += 1

        if category_row is not None and rec is not None:
            current = {
                "manufacturer": row_value(category_row, ic["manufacturer"]),
                "type": row_value(category_row, ic["type"]),
                "model": row_value(category_row, ic["model"]),
                "subcategory": row_value(category_row, ic["subcategory"]),
            }
            flags = [
                "MANUFACTURER" if not equal_value(current["manufacturer"], rec.manufacturer) else "",
                "HARDWARE TYPE" if not equal_value(current["type"], rec.hardware_type) else "",
                "HARDWARE MODEL" if not equal_value(current["model"], rec.hardware_model) else "",
                "SUBCATEGORY" if not equal_value(current["subcategory"], rec.subcategory) else "",
            ]
            action = drift_action(flags)
            if action != "NO UPDATE REQUIRED":
                stats["category_drift"] += 1
                drift_rows.append([
                    hostname, serial, manufacturer, model_id, row_value(category_row, ic["opaque"]),
                    row_value(category_row, ic["serial"]), current["manufacturer"], current["type"],
                    current["model"], current["subcategory"], rec.manufacturer, rec.hardware_type,
                    rec.hardware_model, rec.subcategory,
                    "YES" if not flags[0] else "NO", "YES" if not flags[1] else "NO",
                    "YES" if not flags[2] else "NO", "YES" if not flags[3] else "NO",
                    result.status, result.detail, action,
                ])

    output_dir.mkdir(parents=True, exist_ok=True)
    date_headers = {"Catalog End of Life Date", "Catalog End of Support Date"}
    write_table_workbook(next_output_path(output_dir,"Server Hardware Catalog Match"), "Server Catalog Match", match_headers, match_rows, date_headers)
    make_load_workbook(next_output_path(output_dir,"Server Category Load"), prefix_headers, category, load_rows)
    write_table_workbook(next_output_path(output_dir,"Server Category Catalog Drift"), "Server Category Drift", drift_headers, drift_rows)
    write_table_workbook(next_output_path(output_dir,"Server Catalog Enhancement Recommendations"), "Server Catalog Recommendations", rec_headers, recommendation_rows)
    return dict(stats)


# -----------------------------------------------------------------------------
# Governance summary and GUI
# -----------------------------------------------------------------------------
def write_governance_summary(path: Path, network_stats: dict[str, Any] | None, server_stats: dict[str, Any] | None) -> None:
    headers = ["Module", "Metric", "Value"]
    rows: list[list[Any]] = []
    for module, stats in (("Network", network_stats), ("Server", server_stats)):
        if stats:
            for metric, value in sorted(stats.items()):
                rows.append([module, metric.replace("_", " ").title(), value])
            operational = stats.get("operational_cmdb", 0)
            found = stats.get("found_in_category", 0)
            coverage = found / operational if operational else None
            rows.append([module, "Operational Category Coverage", coverage])
    write_table_workbook(path, "Governance Summary", headers, rows)
    workbook = load_workbook(path)
    worksheet = workbook.active
    for row in range(2, worksheet.max_row + 1):
        if worksheet.cell(row, 2).value == "Operational Category Coverage" and worksheet.cell(row, 3).value is not None:
            worksheet.cell(row, 3).number_format = "0.0%"
    workbook.save(path)
    workbook.close()


def validate_inputs(paths: Iterable[Path]) -> None:
    for path in paths:
        if not path.is_file():
            raise FileNotFoundError(f"File not found: {path}")
        if path.suffix.lower() not in SUPPORTED_EXTENSIONS:
            raise ValueError(f"Unsupported input type: {path.name}")


FEATURE_MANIFEST={
 "NW_HARDWARE_GOVERNANCE":{"agent":"NetworkHardwareGovernanceAgent","service":"HardwareGovernanceService.run_network"},
 "SERVER_HARDWARE_GOVERNANCE":{"agent":"ServerHardwareGovernanceAgent","service":"HardwareGovernanceService.run_server"}}

def launch_gui() -> None:
    import tkinter as tk
    from tkinter import filedialog, messagebox, ttk

    class GovernanceApp(tk.Tk):
        def __init__(self) -> None:
            super().__init__()
            self.title(APP_TITLE)
            self.geometry("1120x720")
            self.minsize(940, 620)
            self.vars = {name: tk.StringVar() for name in (
                "nw_catalog", "nw_cmdb", "nw_category",
                "server_catalog", "server_cmdb", "server_category", "output_dir"
            )}
            self.status = tk.StringVar(value="Select input files and an output folder.")
            self.worker=None;self.shutdown_event=threading.Event();self._build();self.protocol("WM_DELETE_WINDOW",self.close_app)

        def _build(self) -> None:
            root = ttk.Frame(self, padding=18)
            root.pack(fill="both", expand=True)
            root.columnconfigure(1, weight=1)
            ttk.Label(root, text=APP_TITLE, font=("Segoe UI", 17, "bold")).grid(
                row=0, column=0, columnspan=3, sticky="w", pady=(0, 14)
            )
            ttk.Label(root, text="Network Module", font=("Segoe UI", 12, "bold")).grid(
                row=1, column=0, columnspan=3, sticky="w", pady=(4, 4)
            )
            row = 2
            row = self._file_row(root, row, "NW Product Catalog", "nw_catalog")
            row = self._file_row(root, row, "NW CMDB Report", "nw_cmdb")
            row = self._file_row(root, row, "IS Network Category Report", "nw_category")
            ttk.Separator(root).grid(row=row, column=0, columnspan=3, sticky="ew", pady=12); row += 1
            ttk.Label(root, text="Server Module", font=("Segoe UI", 12, "bold")).grid(
                row=row, column=0, columnspan=3, sticky="w", pady=(4, 4)
            ); row += 1
            row = self._file_row(root, row, "Server Product Catalog", "server_catalog")
            row = self._file_row(root, row, "Server CMDB Report", "server_cmdb")
            row = self._file_row(root, row, "IS Server Category Report", "server_category")
            ttk.Separator(root).grid(row=row, column=0, columnspan=3, sticky="ew", pady=12); row += 1
            ttk.Label(root, text="Output Folder").grid(row=row, column=0, sticky="w", padx=(0, 10), pady=6)
            ttk.Entry(root, textvariable=self.vars["output_dir"]).grid(row=row, column=1, sticky="ew", pady=6)
            ttk.Button(root, text="Browse", command=self.pick_output).grid(row=row, column=2, padx=(10, 0)); row += 1
            self.progress = ttk.Progressbar(root, mode="indeterminate")
            self.progress.grid(row=row, column=0, columnspan=3, sticky="ew", pady=(18, 8)); row += 1
            ttk.Label(root, textvariable=self.status, wraplength=1040).grid(row=row, column=0, columnspan=3, sticky="w"); row += 1
            buttons = ttk.Frame(root)
            buttons.grid(row=row, column=0, columnspan=3, pady=18)
            ttk.Button(buttons, text="Run Network", command=lambda: self.execute("network"), width=20).pack(side="left", padx=8)
            ttk.Button(buttons, text="Run Server", command=lambda: self.execute("server"), width=20).pack(side="left", padx=8)
            ttk.Button(buttons, text="Run Network + Server", command=lambda: self.execute("both"), width=24).pack(side="left", padx=8)
            ttk.Button(buttons, text="Open Output Folder", command=self.open_output, width=20).pack(side="left", padx=8)

        def _file_row(self, parent, row: int, label: str, variable_name: str) -> int:
            ttk.Label(parent, text=label).grid(row=row, column=0, sticky="w", padx=(0, 10), pady=6)
            ttk.Entry(parent, textvariable=self.vars[variable_name]).grid(row=row, column=1, sticky="ew", pady=6)
            ttk.Button(parent, text="Browse", command=lambda n=variable_name: self.pick_file(n)).grid(
                row=row, column=2, padx=(10, 0)
            )
            return row + 1

        @staticmethod
        def file_types():
            return [
                ("Supported files", "*.csv *.xlsx *.xlsm *.xlsb"),
                ("CSV", "*.csv"), ("Excel", "*.xlsx *.xlsm"),
                ("Excel Binary", "*.xlsb"), ("All files", "*.*")
            ]

        def pick_file(self, variable_name: str) -> None:
            selected = filedialog.askopenfilename(title="Select input file", filetypes=self.file_types())
            if selected:
                self.vars[variable_name].set(selected)

        def pick_output(self) -> None:
            selected = filedialog.askdirectory(title="Select output folder")
            if selected:
                self.vars["output_dir"].set(selected)

        def open_output(self) -> None:
            path = Path(self.vars["output_dir"].get())
            if not path.is_dir():
                messagebox.showerror(APP_TITLE, "Select a valid output folder first.")
                return
            import os, sys, subprocess
            if sys.platform.startswith("win"):
                os.startfile(path)  # type: ignore[attr-defined]
            elif sys.platform == "darwin":
                subprocess.run(["open", str(path)], check=False)
            else:
                subprocess.run(["xdg-open", str(path)], check=False)

        def execute(self, mode: str) -> None:
            output = Path(self.vars["output_dir"].get().strip())
            if not text(output):
                messagebox.showerror(APP_TITLE, "Select an output folder.")
                return
            output.mkdir(parents=True, exist_ok=True)
            network_stats = server_stats = None
            self.progress.start(10)
            self.status.set(f"Running {mode} reconciliation...")
            self.update_idletasks()
            try:
                if mode in {"network", "both"}:
                    paths = [Path(self.vars[name].get().strip()) for name in ("nw_catalog", "nw_cmdb", "nw_category")]
                    validate_inputs(paths)
                    network_stats = run_network(paths[0], paths[1], paths[2], output)
                if mode in {"server", "both"}:
                    paths = [Path(self.vars[name].get().strip()) for name in ("server_catalog", "server_cmdb", "server_category")]
                    validate_inputs(paths)
                    server_stats = run_server(paths[0], paths[1], paths[2], output)
                write_governance_summary(next_output_path(output,"Hardware Governance Summary"), network_stats, server_stats)
                parts = []
                if network_stats:
                    parts.append(
                        f"Network: {network_stats.get('operational_cmdb', 0)} operational, "
                        f"{network_stats.get('missing_from_category', 0)} missing, "
                        f"{network_stats.get('category_drift', 0)} drift"
                    )
                if server_stats:
                    parts.append(
                        f"Server: {server_stats.get('operational_cmdb', 0)} operational, "
                        f"{server_stats.get('missing_from_category', 0)} missing, "
                        f"{server_stats.get('category_drift', 0)} drift"
                    )
                summary = "Completed. " + " | ".join(parts)
                self.status.set(summary)
                messagebox.showinfo(APP_TITLE, summary + f"\n\nOutput folder:\n{output}")
            except Exception as exc:
                self.status.set("Reconciliation failed. Review the error details.")
                messagebox.showerror(APP_TITLE, f"{exc}\n\n{traceback.format_exc(limit=8)}")
            finally:
                self.progress.stop()

    GovernanceApp().mainloop()


if __name__ == "__main__":
    launch_gui()
